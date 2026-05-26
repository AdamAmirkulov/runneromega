from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\User\Downloads\Карызсыз когам.xlsx")
ws = wb.active

for cell in ws[3]:  # 3-я строка
    if cell.value:
        print(repr(cell.value), "| колонка:", cell.column)