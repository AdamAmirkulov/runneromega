# -*- coding: utf-8 -*-
r"""
Полный скрипт подачи исков.
"""

# ========= АРГУМЕНТЫ КОМАНДНОЙ СТРОКИ =========
import argparse
import os
import sys

def parse_args():
    parser = argparse.ArgumentParser(description="Скрипт подачи исков в office.sud.kz")
    parser.add_argument('--workdir',    type=str, default=None)
    parser.add_argument('--excel_path', type=str, default=None)
    parser.add_argument('--company_id', type=str, default=None)
    return parser.parse_args()

args = parse_args()

if not args.company_id or not args.company_id.strip():
    print("❌ ОШИБКА: не передан --company_id — компания не определена, запуск остановлен.")
    sys.exit(1)

# ДО импорта config — чтобы он увидел правильный COMPANY_ID
os.environ['COMPANY_ID'] = args.company_id.strip()

if args.workdir and args.workdir.strip():
    os.chdir(args.workdir.strip())

if args.excel_path and args.excel_path.strip():
    os.environ['OVERRIDE_EXCEL_PATH'] = args.excel_path.strip()

# ========= ОСТАЛЬНЫЕ ИМПОРТЫ =========

from datetime import datetime
import os
import time
import logging
import re
from openpyxl import load_workbook
from docx import Document
from config import MAIN_EXCEL, CREDENTIALS, ROOT
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select, WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from selenium.common.exceptions import (
    NoSuchElementException,
    TimeoutException,
    StaleElementReferenceException,
    WebDriverException,
)
from selenium.webdriver.firefox.service import Service
from webdriver_manager.firefox import GeckoDriverManager
from selenium.webdriver.firefox.options import Options as FirefoxOptions

# ========= КОНФИГ =========

# ========= КОНФИГ =========

USER_AUTH     = CREDENTIALS['sk_login']
USER_PASSWORD = CREDENTIALS['sk_password']

EXCEL_RESP_FILE  = MAIN_EXCEL
EXCEL_RESP_SHEET = "Отмены"

EXCEL_FIRST_ROW  = 2
PRODUCT_COL      = "B"
EXCEL_RESP_COL   = "D"
EXCEL_REGION_COL = "P"
EXCEL_COURT_COL  = "Q"
EXCEL_FIO_COL    = "C"
EXCEL_IIN_COL    = "D"
EXCEL_SUM_COL    = "J"
EXCEL_DUTY_COL   = "K"

CASES_BASE_DIR = os.path.join(ROOT, "Документы для подачи Исков")


REP_IIN_DEFAULT  = CREDENTIALS['rep_iin']
ORG_BIN_DEFAULT  = CREDENTIALS['org_bin']
ORG_BANK_DEFAULT = CREDENTIALS['org_bank']

BASE         = "https://office.sud.kz"
LOGIN_URL    = f"{BASE}/index.xhtml"
HOME_URL     = f"{BASE}/form/proceedings/services.xhtml"
SEND_DOCS_URL = f"{BASE}/form/send/index.xhtml"


WAIT = 1.0

# Максимальное количество перезапусков браузера на одну строку
MAX_BROWSER_RESTARTS = 3

REGION_MAP = {
    "ГОРОД АСТАНА": "город Астана",
    "АСТАНА": "город Астана",
    "ГОРОД АЛМАТЫ": "город Алматы",
    "АЛМАТЫ": "город Алматы",
    "ГОРОД ШЫМКЕНТ": "город Шымкент",
    "ШЫМКЕНТ": "город Шымкент",
    "АКМОЛИНСКАЯ ОБЛАСТЬ": "Акмолинская область",
    "АКТЮБИНСКАЯ ОБЛАСТЬ": "Актюбинская область",
    "АЛМАТИНСКАЯ ОБЛАСТЬ": "Алматинская область",
    "АТЫРАУСКАЯ ОБЛАСТЬ": "Атырауская область",
    "ВОСТОЧНО-КАЗАХСТАНСКАЯ ОБЛАСТЬ": "Восточно-Казахстанская область",
    "ЖАМБЫЛСКАЯ ОБЛАСТЬ": "Жамбылская область",
    "ЗАПАДНО-КАЗАХСТАНСКАЯ ОБЛАСТЬ": "Западно-Казахстанская область",
    "КАРАГАНДИНСКАЯ ОБЛАСТЬ": "Карагандинская область",
    "КОСТАНАЙСКАЯ ОБЛАСТЬ": "Костанайская область",
    "КЫЗЫЛОРДИНСКАЯ ОБЛАСТЬ": "Кызылординская область",
    "МАНГИСТАУСКАЯ ОБЛАСТЬ": "Мангистауская область",
    "ПАВЛОДАРСКАЯ ОБЛАСТЬ": "Павлодарская область",
    "СЕВЕРО-КАЗАХСТАНСКАЯ ОБЛАСТЬ": "Северо-Казахстанская область",
    "ТУРКЕСТАНСКАЯ ОБЛАСТЬ": "Туркестанская область",
    "ОБЛАСТЬ ҰЛЫТАУ": "Область Ұлытау",
    "ОБЛАСТЬ АБАЙ": "Область Абай",
    "ОБЛАСТЬ ЖЕТІСУ": "Область Жетісу",
    "ВОЕННЫЙ СУД РЕСПУБЛИКИ КАЗАХСТАН": "Военный суд Республики Казахстан",
    "УПРАЗДНЁННЫЕ СУДЫ (С 2022 ГОДА)": "Упразднённые суды (c 2022 года)",
}

# ========= ЛОГИ =========

logging.basicConfig(
    filename="sud_script.log",
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    encoding="utf-8",
)


def log_step(msg):
    print(msg)
    logging.info(msg)


def log_warn(msg):
    print("WARN:", msg)
    logging.warning(msg)


def log_err(msg):
    print("ERROR:", msg)
    logging.error(msg)


def _format_duration(seconds):
    mins, secs = divmod(seconds, 60)
    if mins >= 1:
        return f"{int(mins)} мин {secs:.1f} сек"
    return f"{secs:.1f} сек"


def timed_step(label):
    def decorator(func):
        def wrapper(*args, **kwargs):
            t0 = time.time()
            log_step(f"[ШАГ] {label} — старт")
            try:
                return func(*args, **kwargs)
            finally:
                dt = time.time() - t0
                log_step(f"[ШАГ] {label} — завершён за {_format_duration(dt)}")
        return wrapper
    return decorator


# ========= DRIVER =========

def init_driver():
    opts = FirefoxOptions()
    opts.set_preference("permissions.default.image", 2)
    opts.set_preference("dom.ipc.plugins.enabled.libflashplayer.so", False)
    opts.set_preference("extensions.enabledScopes", 0)
    opts.set_preference("browser.tabs.animate", False)
    opts.add_argument("--start-maximized")
    opts.set_preference("dom.content_processes.count", 1)
    opts.add_argument("--disable-gpu")

    service = Service(GeckoDriverManager().install())
    driver = webdriver.Firefox(service=service, options=opts)
    driver.set_page_load_timeout(60)
    return driver


def safe_quit_driver(driver):
    try:
        if driver:
            driver.quit()
    except Exception as e:
        log_warn(f"Ошибка при закрытии браузера: {e}")


def restart_driver(driver):
    log_step("🔄 Перезапускаю браузер...")
    safe_quit_driver(driver)
    time.sleep(3)
    new_driver = init_driver()
    log_step("✅ Браузер перезапущен.")
    return new_driver


# ========= ВСПОМОГАТЕЛЬНЫЕ =========

def find_with_fallbacks(driver, variants, desc, tries=8, delay=0.4):
    last = None
    for _ in range(tries):
        for by, sel in variants:
            try:
                el = driver.find_element(by, sel)
                logging.info(f"Нашёл {desc} по {by}={sel}")
                return el
            except Exception as e:
                last = e
        time.sleep(delay)
    raise NoSuchElementException(f"Не удалось найти {desc}. Последняя ошибка: {last}")


def _css_from_locator(locator):
    by, sel = locator
    if by == By.CSS_SELECTOR:
        return sel
    if by == By.ID:
        css_id = sel.replace(":", r"\:")
        return f"#{css_id}"
    return sel


def _wait(drv, cond, t=25):
    return WebDriverWait(drv, t, ignored_exceptions=(StaleElementReferenceException,)).until(cond)


def _select_by_value_robust(driver, locator, value, desc, timeout=20, tries=3):
    w = WebDriverWait(driver, timeout)
    w.until(EC.presence_of_element_located(locator))

    def _opt_present(d):
        try:
            el = d.find_element(*locator)
            return any((o.get_attribute("value") or "") == value
                       for o in el.find_elements(By.TAG_NAME, "option"))
        except StaleElementReferenceException:
            return False

    w.until(_opt_present)

    last_err = None
    for _ in range(tries):
        try:
            el = driver.find_element(*locator)
            Select(el).select_by_value(value)
            w.until(lambda d: d.find_element(*locator).get_attribute("value") == value)
            print(f"✔ {desc} → {value}")
            return
        except (StaleElementReferenceException, TimeoutException) as e:
            last_err = e

        ok = driver.execute_script(
            """
            var el = document.querySelector(arguments[0]);
            if(!el) return false;
            var idx = Array.from(el.options).findIndex(o => (o.value||'') === arguments[1]);
            if(idx < 0) return false;
            el.selectedIndex = idx;
            el.dispatchEvent(new Event('change', {bubbles:true}));
            return true;
            """,
            _css_from_locator(locator),
            value,
        )
        if not ok:
            raise TimeoutException(
                f"Не удалось выбрать '{desc}' со значением {value}. Последняя ошибка: {last_err}"
            )
        WebDriverWait(driver, 10).until(
            lambda d: d.find_element(*locator).get_attribute("value") == value
        )
        print(f"✔ {desc} → {value} (через JS)")


def _click_simple(drv, locator, desc, t=25):
    btn = _wait(drv, EC.element_to_be_clickable(locator), t)
    drv.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
    try:
        btn.click()
    except Exception:
        drv.execute_script("arguments[0].click();", btn)
    print(f"✔ Нажал {desc}")


def wait_ajax_idle(driver, timeout=25):
    end = time.time() + timeout
    overlays = [
        (By.CSS_SELECTOR, ".ui-widget-overlay"),
        (By.CSS_SELECTOR, ".ui-dialog-mask"),
        (By.CSS_SELECTOR, ".blockUI"),
        (By.CSS_SELECTOR, ".ui-blockui"),
        (By.CSS_SELECTOR, "[class*='loading']"),
        (By.CSS_SELECTOR, "[class*='spinner']"),
    ]
    while time.time() < end:
        try:
            rs = driver.execute_script("return document.readyState")
            if rs not in ("complete", "interactive"):
                time.sleep(0.2)
                continue
            jq_active = driver.execute_script(
                "return (window.jQuery && window.jQuery.active) ? window.jQuery.active : 0;"
            )
            if jq_active and int(jq_active) > 0:
                time.sleep(0.2)
                continue
            pf_busy = driver.execute_script(
                "return (window.PrimeFaces && PrimeFaces.ajax && PrimeFaces.ajax.Queue)"
                " ? PrimeFaces.ajax.Queue.isEmpty() : true;"
            )
            if pf_busy is False:
                time.sleep(0.2)
                continue
            any_visible = False
            for loc in overlays:
                for el in driver.find_elements(*loc):
                    try:
                        if el.is_displayed():
                            any_visible = True
                            break
                    except Exception:
                        pass
                if any_visible:
                    break
            if any_visible:
                time.sleep(0.2)
                continue
            return
        except Exception:
            time.sleep(0.2)


def click_next_and_wait_payment(driver, timeout_click=30, timeout_payment=90, retries=5):
    last_err = None
    for attempt in range(1, retries + 1):
        try:
            # Сначала закрываем все возможные открытые модалки
            for js_func in ['hideFizModalDialog', 'hideJurModalDialog', 'hideSelectSideModalDialog']:
                try:
                    driver.execute_script(f"if(typeof {js_func}==='function') {js_func}();")
                except Exception:
                    pass
            
            wait_ajax_idle(driver, timeout=timeout_click)
            time.sleep(0.5)
            
            # Кликаем кнопку Далее (button-orange)
            btn_loc = (By.XPATH, "//a[contains(@class,'button-orange') and normalize-space()='Далее']")
            btn = WebDriverWait(driver, timeout_click).until(EC.element_to_be_clickable(btn_loc))
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
            try:
                btn.click()
            except Exception:
                driver.execute_script("arguments[0].click();", btn)
            
            # Если клик не сработал — вызываем JS напрямую
            try:
                WebDriverWait(driver, 5).until(
                    EC.presence_of_element_located((By.CSS_SELECTOR, "select[id$='selectKbk']"))
                )
                print("✔ Успешно нажал «Далее» и перешёл на вкладку «Оплата»")
                return
            except TimeoutException:
                print("⚠ Клик не помог, пробую JS...")
                driver.execute_script("fillHideFields(); goNext();")
            
            WebDriverWait(driver, timeout_payment).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "select[id$='selectKbk']"))
            )
            print("✔ Успешно нажал «Далее» и перешёл на вкладку «Оплата»")
            return
            
        except Exception as e:
            last_err = e
            print(f"⚠ «Далее» не сработала (попытка {attempt}/{retries}): {e}")
            time.sleep(1.5)
    
    raise RuntimeError(f"Не удалось перейти на вкладку «Оплата». Последняя ошибка: {last_err}")

# ========= МОДАЛКА ВЫБОРА УЧАСТНИКА =========

def select_person_type_modal(driver, side_text, person_type, resident="Да", timeout=30, retries=5):
    def _wait_visible_modal():
        return WebDriverWait(driver, timeout).until(
            EC.visibility_of_element_located((By.CSS_SELECTOR, "select[id$=':pp-type']"))
        )

    def _pick_type_like_human(sel_el, want_true):
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", sel_el)
        sel_el.click()
        time.sleep(0.1)
        target_val = "true" if want_true else "false"
        try:
            Select(sel_el).select_by_value(target_val)
            return
        except Exception:
            pass
        for _ in range(10):
            try:
                cur_val = Select(sel_el).first_selected_option.get_attribute("value")
            except Exception:
                cur_val = None
            if cur_val == target_val:
                sel_el.send_keys(Keys.ENTER)
                return
            sel_el.send_keys(Keys.ARROW_DOWN)
            time.sleep(0.05)
        driver.execute_script(
            "var el=arguments[0]; el.value=arguments[1];"
            " el.dispatchEvent(new Event('change',{bubbles:true}));",
            sel_el, target_val
        )

    def _safe_select_by_visible(locator, text):
        sel = WebDriverWait(driver, timeout).until(EC.presence_of_element_located(locator))
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", sel)
        sel.click()
        Select(sel).select_by_visible_text(text)

    want_jur = (person_type.strip().lower() == "юр лицо")
    last_err = None

    for attempt in range(1, retries + 1):
        try:
            _wait_visible_modal()
            type_sel = WebDriverWait(driver, timeout).until(
                EC.element_to_be_clickable((By.CSS_SELECTOR, "select[id$=':pp-type']"))
            )
            _pick_type_like_human(type_sel, want_true=want_jur)

            try:
                _safe_select_by_visible((By.CSS_SELECTOR, "select[id$=':pp-side']"), side_text)
            except (TimeoutException, NoSuchElementException):
                pass

            try:
                _safe_select_by_visible((By.CSS_SELECTOR, "select[id$=':pp-resident']"), resident)
            except (TimeoutException, NoSuchElementException):
                pass

            btn_next = WebDriverWait(driver, timeout).until(
                EC.element_to_be_clickable((
                    By.XPATH,
                    "//button[normalize-space()='Далее'] | //input[@type='button' and @value='Далее']"
                ))
            )
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn_next)
            try:
                btn_next.click()
            except Exception:
                driver.execute_script("arguments[0].click();", btn_next)

            print(f"✔ Модалка: {'Юр лицо' if want_jur else 'Физ лицо'}")
            return
        except StaleElementReferenceException as e:
            last_err = e
            print(f"⚠ stale на модалке (попытка {attempt}/{retries}) — повторяю...")
            time.sleep(0.5)
        except Exception as e:
            last_err = e
            print(f"⚠ ошибка на модалке (попытка {attempt}/{retries}): {e}")
            time.sleep(0.5)

    raise RuntimeError(f"Не удалось обработать модалку. Последняя ошибка: {last_err}")


# ========= LOGIN =========

@timed_step("Логин в office.sud.kz")
def login(driver, timeout=20, retries=5):
    for attempt in range(1, retries + 1):
        try:
            log_step("Открываю страницу логина…")
            driver.get(LOGIN_URL)
            time.sleep(WAIT)

            cur = (driver.current_url or "").lower()
            if "index.xhtml" not in cur:
                log_step(f"ℹ Уже авторизован ({driver.current_url}). Перехожу на HOME.")
                driver.get(HOME_URL)
                time.sleep(WAIT)
                return

            if not driver.find_elements(
                By.XPATH,
                "//input[contains(@id,':auth:xin') or contains(@placeholder,'ИИН') "
                "or contains(@placeholder,'ИИН/БИН')]"
            ):
                driver.get(HOME_URL)
                time.sleep(WAIT)
                log_step("ℹ Поля логина нет — открываю HOME.")
                return

            try:
                ru = driver.find_elements(
                    By.XPATH, "//a[contains(.,'РУС') and not(contains(@class,'active'))]"
                )
                if ru:
                    ru[0].click()
                    log_step("Переключил язык на РУС")
                    time.sleep(WAIT)
            except Exception:
                log_warn("Не удалось переключить язык на РУС")

            login_el = find_with_fallbacks(
                driver,
                variants=[
                    (By.ID, "j_idt78:auth:xin"),
                    (By.XPATH, "//input[contains(@id,':auth:xin')]"),
                    (By.NAME, "j_idt78:auth:xin"),
                    (By.XPATH, "//input[contains(@placeholder,'ИИН') or contains(@placeholder,'ЖСН')]"),
                ],
                desc="поле ИИН/БИН", tries=10, delay=0.4,
            )
            pass_el = find_with_fallbacks(
                driver,
                variants=[
                    (By.ID, "j_idt78:auth:password"),
                    (By.XPATH, "//input[contains(@id,':auth:password')]"),
                    (By.XPATH, "//input[@type='password']"),
                    (By.XPATH, "//input[contains(@placeholder,'Пароль')]"),
                ],
                desc="поле Пароль", tries=10, delay=0.4,
            )
            submit = find_with_fallbacks(
                driver,
                variants=[
                    (By.CSS_SELECTOR, "input.button-primary[type='submit']"),
                    (By.XPATH, "//input[@type='submit' and contains(@value,'Войти')]"),
                    (By.XPATH, "//button[contains(.,'Войти')]"),
                ],
                desc="кнопка Войти", tries=10, delay=0.4,
            )

            login_el.clear()
            login_el.send_keys(USER_AUTH)
            pass_el.clear()
            pass_el.send_keys(USER_PASSWORD)
            log_step("Ввёл ИИН/БИН и Пароль")
            submit.click()
            log_step("Нажал 'Войти'")
            time.sleep(WAIT * 2)
            driver.get(HOME_URL)
            time.sleep(WAIT)
            return
        except Exception as e:
            log_warn(f"Логин не удался (попытка {attempt}/{retries}): {e}")
            time.sleep(2.0)
    raise RuntimeError("Не удалось выполнить логин после нескольких попыток.")


# ========= ПЕРЕХОД В «Подача документов» =========

@timed_step("Переход к 'Подача документов'")
def go_to_send_docs(driver):
    log_step("Открываю главную страницу (плитки)…")
    driver.get(HOME_URL)
    time.sleep(WAIT)

    link = None
    try:
        link = WebDriverWait(driver, 5).until(
            EC.element_to_be_clickable(
                (By.XPATH, "//a[contains(@href, '/form/send/index.xhtml')]")
            )
        )
    except TimeoutException:
        link = None

    if not link:
        try:
            link = driver.find_element(
                By.XPATH, "//p[normalize-space()='Подача документов']/ancestor::a"
            )
        except Exception:
            try:
                link = driver.find_element(
                    By.XPATH, "//a[.//p[contains(normalize-space(),'Подача документ')]]"
                )
            except Exception:
                link = None

    if link:
        try:
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", link)
        except Exception:
            pass
        try:
            link.click()
        except Exception:
            driver.execute_script("arguments[0].click();", link)
        log_step("Кликнул по плитке 'Подача документов'")
    else:
        log_warn("Плитка не найдена, открываю по прямому URL")
        driver.get(SEND_DOCS_URL)

    try:
        WebDriverWait(driver, 10).until(EC.url_contains("/form/send"))
    except TimeoutException:
        log_warn("Не удалось подтвердить загрузку страницы 'Подача документов'")


# ========= БЛОК 2: CIVIL/FIRSTINSTANCE/Иск =========

@timed_step("Выбор CIVIL/FIRSTINSTANCE/Иск")
def send_claim(driver, timeout=20):
    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$=':case-type']"), "CIVIL", "Тип производства", timeout
    )
    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$=':instance']"), "FIRSTINSTANCE", "Инстанция", timeout
    )
    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$=':request']"), "3", "Тип документа", timeout
    )

    btn = None
    w = WebDriverWait(driver, timeout)
    for loc in [
        (By.XPATH, "//input[@type='submit' and @value='Отправить']"),
        (By.XPATH, "//input[contains(@id,':j_idt') and @type='submit']"),
        (By.XPATH, "//button[normalize-space()='Отправить']"),
    ]:
        try:
            btn = w.until(EC.element_to_be_clickable(loc))
            break
        except TimeoutException:
            continue

    if not btn:
        raise TimeoutException("Кнопка «Отправить» не найдена.")
    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)
    print("✔ Нажал «Отправить».")


# ========= ХЕЛПЕР: характер заявления + упрощённое =========

def ensure_character_and_simple_process(driver, timeout=12):
    w = WebDriverWait(driver, timeout, ignored_exceptions=(StaleElementReferenceException,))
    try:
        w.until(EC.presence_of_element_located((By.CSS_SELECTOR, "select[id$=':edit-character']")))
        _select_by_value_robust(
            driver, (By.CSS_SELECTOR, "select[id$=':edit-character']"),
            "1", "Характер заявления", timeout
        )
    except TimeoutException:
        print("⚠ Не нашёл 'Характер заявления' — продолжаю.")
    except Exception as e:
        print(f"⚠ Ошибка при 'Характер заявления': {e}")

    try:
        cb = w.until(EC.presence_of_element_located(
            (By.CSS_SELECTOR, "input[id$=':edit-simpleProcess'][type='checkbox']")
        ))
        if not cb.is_selected():
            driver.execute_script("arguments[0].click();", cb)
        print("✔ Галочка «Дело упрощенного производства» установлена")
    except TimeoutException:
        print("⚠ Не нашёл чекбокс 'Дело упрощенного производства' — продолжаю.")
    except Exception as e:
        print(f"⚠ Ошибка при галочке: {e}")


# ========= ФОРМА ОРГАНИЗАЦИИ =========

def _fill_org_form(driver, timeout=30):
    w = WebDriverWait(driver, timeout)
    ORG_BIN = (By.XPATH, "//input[contains(@id,':org-bin')]")
    ORG_JUR = (By.XPATH, "//input[contains(@id,':org-jurAddress')]")
    ORG_FACT = (By.XPATH, "//input[contains(@id,':org-factAddress')]")
    ORG_BANK = (By.XPATH, "//input[contains(@id,':org-bankDetails')]")

    bin_el = w.until(EC.presence_of_element_located(ORG_BIN))
    bin_el.clear()
    bin_el.send_keys(ORG_BIN_DEFAULT)

    try:
        driver.execute_script(
            "if (typeof fillOrgData==='function'){ fillOrgData(arguments[0]); }",
            bin_el.get_attribute("id"),
        )
        print("✔ вызвал fillOrgData()")
    except Exception:
        try:
            lupa = driver.find_element(
                By.XPATH,
                "//input[contains(@id,':org-bin')]/following-sibling::span[contains(@class,'gbdSearch')]",
            )
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", lupa)
            lupa.click()
            print("✔ клик по лупе возле БИН")
        except Exception:
            print("⚠ не удалось вызвать поиск организации")

    try:
        jur_el = w.until(EC.presence_of_element_located(ORG_JUR))
        w.until(lambda d: d.find_element(*ORG_JUR).get_attribute("value").strip() != "")
    except TimeoutException:
        print("⚠ юр. адрес не подтянулся, продолжаю")
        jur_el = driver.find_element(*ORG_JUR)

    jur_addr = jur_el.get_attribute("value") or ""
    fact_el = driver.find_element(*ORG_FACT)
    fact_el.clear()
    fact_el.send_keys(jur_addr)
    print("✔ Фактический адрес скопирован из юридического")

    bank_el = driver.find_element(*ORG_BANK)
    bank_el.clear()
    bank_el.send_keys(ORG_BANK_DEFAULT)
    print("✔ Банковские реквизиты установлены")

    for loc in [
        (By.XPATH, "//input[@type='button' and @value='Сохранить']"),
        (By.CSS_SELECTOR, "input[id$=':j_idt273'][value='Сохранить']"),
    ]:
        try:
            _click_simple(driver, loc, "«Сохранить» (организация)", t=timeout)
            break
        except TimeoutException:
            continue
    else:
        raise TimeoutException("Кнопка «Сохранить» в форме организации не найдена")

    try:
        WebDriverWait(driver, 10).until_not(EC.presence_of_element_located(ORG_BIN))
    except TimeoutException:
        print("⚠ Не дождался закрытия формы организации")
    print("✅ Форма организации сохранена")


# ========= БЛОК 3: истец (юрлицо) =========

@timed_step("Блок 3: истец (юрлицо)")
def block3_fill_case_and_org_participant(driver, timeout=40):
    w = WebDriverWait(driver, timeout)
    try:
        w.until(
            lambda d: "/form/requestType2/createRequest.xhtml" in d.current_url
            or d.find_element(By.CSS_SELECTOR, "select[id$=':edit-category']")
        )
    except Exception:
        raise TimeoutException("Страница createRequest.xhtml не открыта")

    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$=':edit-categoryGroup']"),
        "2", "Вид производства по делу", timeout
    )
    _wait(driver, EC.presence_of_element_located(
        (By.CSS_SELECTOR, "select[id$=':edit-category'] option[value='27']")
    ), 30)
    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$=':edit-category']"),
        "27", "Категория дела", timeout
    )
    ensure_character_and_simple_process(driver, timeout=timeout)

    for loc in [
        (By.XPATH, "//button[contains(.,'Добавить участника процесса')]"),
        (By.CSS_SELECTOR, "button[id*='addPerson']"),
    ]:
        try:
            _click_simple(driver, loc, "«Добавить участника процесса»", t=timeout)
            break
        except TimeoutException:
            continue
    else:
        raise TimeoutException("Кнопка «Добавить участника процесса» не найдена")

    select_person_type_modal(driver, side_text="истец", person_type="Юр лицо",
                             resident="Да", timeout=timeout)
    _fill_org_form(driver, timeout=timeout)


# ========= СОХРАНИТЬ ФИЗЛИЦО =========

def click_save_person(driver, timeout=20):
    w = WebDriverWait(driver, timeout)
    
    save_locators = [
        (By.XPATH, "//input[contains(@onclick,'fillFizHideFields')]"),
        (By.XPATH, "//input[@type='submit' and @value='Сохранить' and contains(@onclick,'Fiz')]"),
        (By.XPATH, "//input[@value='Сохранить' and contains(@class,'button-primary')]"),
    ]
    
    btn = None
    for loc in save_locators:
        try:
            btn = w.until(EC.element_to_be_clickable(loc))
            break
        except TimeoutException:
            continue
    
    if not btn:
        raise TimeoutException("Кнопка «Сохранить» (физлицо) не найдена")
    
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)
    print("✔ Нажал «Сохранить» (физлицо)")
    
    try:
        WebDriverWait(driver, timeout).until(
            EC.invisibility_of_element_located(
                (By.XPATH, "//input[contains(@onclick,'hideFizModalDialog')]")
            )
        )
        print("✅ Модалка физлица закрыта")
    except TimeoutException:
        print("⚠ Модалка не закрылась сама, закрываю через JS...")
        driver.execute_script("if(typeof hideFizModalDialog==='function') hideFizModalDialog();")
        time.sleep(1.0)
    
    wait_ajax_idle(driver, timeout=20)
    
    # ← ДОБАВЬ ЭТО: ждём пока кнопка "Добавить участника" снова станет кликабельной
    try:
        WebDriverWait(driver, 15).until(
            EC.element_to_be_clickable(
                (By.XPATH, "//button[contains(.,'Добавить участника процесса')]")
            )
        )
        print("✅ Кнопка 'Добавить участника' снова активна")
    except TimeoutException:
        print("⚠ Кнопка не стала активной за 15с — всё равно продолжаю")
    
    time.sleep(1.5)  # ← увеличь с 0.5 до 1.5

# ========= БЛОК 4: ответчик — открыть форму =========

@timed_step("Блок 4: открыть форму ответчика")
def block4_open_respondent_form(driver, timeout=60):
    print("\n========== БЛОК 4: ОТВЕТЧИК-ФИЗЛИЦО ==========")
    w = WebDriverWait(driver, timeout)

    wait_ajax_idle(driver, timeout=30)
    time.sleep(1.0)

    last_err = None
    clicked = False
    for attempt in range(1, 6):
        for loc in [
            (By.XPATH, "//button[contains(.,'Добавить участника процесса')]"),
            (By.CSS_SELECTOR, "button[id*='addPerson']"),
        ]:
            try:
                _click_simple(driver, loc, "«Добавить участника процесса» (ответчик)", t=timeout)
                clicked = True
                break
            except TimeoutException:
                continue
        if clicked:
            break
        last_err = f"Попытка {attempt}/5"
        print(f"⚠ {last_err}, жду...")
        wait_ajax_idle(driver, timeout=20)
        time.sleep(2.0)

    if not clicked:
        raise TimeoutException(f"Не удалось нажать «Добавить участника» для ответчика. {last_err}")

    select_person_type_modal(driver, side_text="ответчик", person_type="Физ лицо",
                             resident="Да", timeout=timeout)

    IIN_LOC = (By.XPATH,
               "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
               "or contains(@placeholder,'ЖСН')]")
    w.until(EC.presence_of_element_located(IIN_LOC))
    print("✅ Открыта форма ответчика")

    side_val = ""
    try:
        side_text_input = driver.find_element(
            By.XPATH,
            "//label[contains(normalize-space(),'Сторона процесса')]"
            "/following::input[@type='text'][1]",
        )
        side_val = (side_text_input.get_attribute("value") or "").strip()
    except Exception as e:
        print(f"⚠ Не удалось прочитать 'Сторона процесса': {e}")

    print(f"🔎 'Сторона процесса' = {repr(side_val)}")
    if side_val and side_val.lower() != "ответчик":
        raise RuntimeError(f"Сторона процесса = {side_val!r}, должна быть 'ответчик'")
    print("=== Блок 4 завершён ===")


# ========= БЛОК 5: ИИН из Excel (ответчик) =========

@timed_step("Блок 5: ответчик — ИИН из Excel")
def block5_fill_person_from_excel_and_save(driver, row, timeout=40):
    print("\n========== БЛОК 5: ОТВЕТЧИК — ИИН ИЗ EXCEL ==========")
    w = WebDriverWait(driver, timeout)

    try:
        wb = load_workbook(EXCEL_RESP_FILE, data_only=True)
        ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active
        raw_val = ws[f"{EXCEL_RESP_COL}{row}"].value
        iin = str(raw_val).strip() if raw_val is not None else ""
        print(f"[EXCEL] {EXCEL_RESP_COL}{row}: {raw_val!r} → ИИН: {iin}")
    except Exception as e:
        raise RuntimeError(f"[EXCEL] Не удалось прочитать ИИН: {e}")

    if not iin.isdigit():
        raise RuntimeError(f"[EXCEL] Некорректный ИИН: {iin!r}")

    IIN_LOC = (By.XPATH,
               "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
               "or contains(@placeholder,'ЖСН')]")
    iin_el = w.until(EC.presence_of_element_located(IIN_LOC))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", iin_el)
    iin_el.clear()
    iin_el.send_keys(iin)
    print(f"✔ Ввел ИИН ответчика: {iin}")

    lupa_xpath = (
        "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
        "or contains(@placeholder,'ЖСН')]"
        "/following-sibling::span[contains(@class,'gbdSearch')]"
    )
    lupa = w.until(EC.element_to_be_clickable((By.XPATH, lupa_xpath)))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", lupa)
    try:
        lupa.click()
    except Exception:
        driver.execute_script("arguments[0].click();", lupa)
    print("✔ Нажал лупу")

    time.sleep(0.3)
    try:
        fio_el = driver.find_element(
            By.XPATH,
            "//input[contains(@id,':person-fio') or contains(@id,':person-lastName') "
            "or contains(@placeholder,'Фамилия')]",
        )
        print(f"ℹ Фамилия: {(fio_el.get_attribute('value') or '').strip()!r}")
    except Exception:
        print("⚠ Не удалось прочитать ФИО")

    click_save_person(driver, timeout=timeout)
    print("=== Блок 5 завершён ===")


# ========= БЛОК 6: представитель — открыть форму =========

@timed_step("Блок 6: открыть форму представителя")
def block6_open_representative_form(driver, timeout=60, retries=5):
    print("\n========== БЛОК 6: ПРЕДСТАВИТЕЛЬ ==========")
    w = WebDriverWait(driver, timeout)

    wait_ajax_idle(driver, timeout=30)
    time.sleep(3.0)  # ← было 1.5, стало 3.0
    
    # ← ДОБАВЬ: скролл вверх, кнопка может быть вне зоны видимости
    driver.execute_script("window.scrollTo(0, 0);")
    time.sleep(0.5)

    last_err = None
    clicked = False
    for attempt in range(1, retries + 1):
        for loc in [
            (By.XPATH, "//button[contains(.,'Добавить участника процесса')]"),
            (By.CSS_SELECTOR, "button[id*='addPerson']"),
            # ← ДОБАВЬ запасной вариант через JS
        ]:
            try:
                _click_simple(driver, loc, "«Добавить участника процесса» (представитель)", t=timeout)
                clicked = True
                break
            except TimeoutException:
                continue
        
        if not clicked:
            # ← ДОБАВЬ: попробуй кликнуть через JS если обычный клик не работает
            try:
                result = driver.execute_script("""
                    var btns = document.querySelectorAll('button');
                    for(var b of btns){
                        if(b.textContent.includes('Добавить участника')){
                            b.click();
                            return true;
                        }
                    }
                    return false;
                """)
                if result:
                    print("✔ Кликнул через JS (представитель)")
                    clicked = True
            except Exception as js_err:
                print(f"⚠ JS-клик не сработал: {js_err}")
        
        if clicked:
            break
        last_err = f"Попытка {attempt}/{retries}"
        print(f"⚠ {last_err}: кнопка не найдена, жду...")
        wait_ajax_idle(driver, timeout=20)
        time.sleep(3.0)
        if clicked:
            break
        last_err = f"Попытка {attempt}/{retries}"
        print(f"⚠ {last_err}: кнопка не найдена, жду...")
        wait_ajax_idle(driver, timeout=20)
        time.sleep(3.0)

    if not clicked:
        raise TimeoutException(
            f"Не удалось нажать «Добавить участника» для представителя. {last_err}"
        )

    select_person_type_modal(driver, side_text="представитель", person_type="Физ лицо",
                             resident="Да", timeout=timeout)

    IIN_LOC = (By.XPATH,
               "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
               "or contains(@placeholder,'ЖСН')]")
    w.until(EC.presence_of_element_located(IIN_LOC))
    print("✅ Открыта форма представителя")

    side_val = ""
    try:
        side_text_input = driver.find_element(
            By.XPATH,
            "//label[contains(normalize-space(),'Сторона процесса')]"
            "/following::input[@type='text'][1]",
        )
        side_val = (side_text_input.get_attribute("value") or "").strip()
    except Exception as e:
        print(f"⚠ Не удалось прочитать 'Сторона процесса': {e}")

    print(f"🔎 'Сторона процесса' = {repr(side_val)}")
    if side_val and side_val.lower() != "представитель":
        raise RuntimeError(f"Сторона процесса = {side_val!r}, должна быть 'представитель'")
    print("=== Блок 6 завершён ===")


# ========= БЛОК 7: ИИН представителя =========

@timed_step("Блок 7: представитель — фиксированный ИИН")
def block7_fill_representative_from_default(driver, timeout=40):
    print("\n========== БЛОК 7: ПРЕДСТАВИТЕЛЬ ==========")
    w = WebDriverWait(driver, timeout)

    iin = REP_IIN_DEFAULT.strip()
    if not iin.isdigit():
        raise RuntimeError(f"[REP] Некорректный ИИН представителя: {iin!r}")

    IIN_LOC = (By.XPATH,
               "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
               "or contains(@placeholder,'ЖСН')]")
    iin_el = w.until(EC.presence_of_element_located(IIN_LOC))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", iin_el)
    iin_el.clear()
    iin_el.send_keys(iin)
    print(f"✔ Ввел ИИН представителя: {iin}")

    lupa_xpath = (
        "//input[contains(@id,':person-iin') or contains(@placeholder,'ИИН') "
        "or contains(@placeholder,'ЖСН')]"
        "/following-sibling::span[contains(@class,'gbdSearch')]"
    )
    lupa = w.until(EC.element_to_be_clickable((By.XPATH, lupa_xpath)))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", lupa)
    try:
        lupa.click()
    except Exception:
        driver.execute_script("arguments[0].click();", lupa)
    print("✔ Нажал лупу")

    time.sleep(0.3)
    try:
        fio_el = driver.find_element(
            By.XPATH,
            "//input[contains(@id,':person-fio') or contains(@id,':person-lastName') "
            "or contains(@placeholder,'Фамилия')]",
        )
        print(f"ℹ Фамилия представителя: {(fio_el.get_attribute('value') or '').strip()!r}")
    except Exception:
        print("⚠ Не удалось прочитать ФИО представителя")

    click_save_person(driver, timeout=timeout)
    print("=== Блок 7 завершён ===")


# ========= БЛОК 8: Получатель + «Далее» =========

@timed_step("Блок 8: получатель из Excel + 'Далее'")
def block8_fill_recipient_from_excel_and_next(driver, row, timeout=40):
    print("\n========== БЛОК 8: ПОЛУЧАТЕЛЬ ==========")
    w = WebDriverWait(driver, timeout)

    try:
        wb = load_workbook(EXCEL_RESP_FILE, data_only=True)
        ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active
        raw_region = ws[f"{EXCEL_REGION_COL}{row}"].value
        raw_court = ws[f"{EXCEL_COURT_COL}{row}"].value
        region = str(raw_region or "").strip()
        court = str(raw_court or "").strip()
        print(f"[EXCEL] Область: {raw_region!r}, Суд: {raw_court!r}")
    except Exception as e:
        raise RuntimeError(f"[EXCEL] Не удалось прочитать область/суд: {e}")

    if not region:
        raise RuntimeError("[EXCEL] Пустое значение области")
    if not court:
        raise RuntimeError("[EXCEL] Пустое значение суда")

    district_loc = (By.CSS_SELECTOR, "select[id$=':edit-district']")
    sel_district = w.until(EC.presence_of_element_located(district_loc))

    region_upper = region.upper().strip()
    target_region_text = REGION_MAP.get(region_upper, region.strip())
    print(f"[REGION] Ищу: {target_region_text!r}")

    region_value = None
    for opt in sel_district.find_elements(By.TAG_NAME, "option"):
        text = (opt.text or "").strip()
        if text == target_region_text or text.upper().strip() == region_upper:
            region_value = opt.get_attribute("value") or ""
            print(f"[REGION] Найдено: {text!r} → {region_value!r}")
            break

    if not region_value:
        raise RuntimeError(f"[REGION] Не удалось сопоставить область '{region}'")

    _select_by_value_robust(driver, district_loc, region_value, "Область получателя", timeout=timeout)
    time.sleep(2.0)

    court_loc = (By.CSS_SELECTOR, "select[id$=':edit-court']")
    sel_court = w.until(EC.presence_of_element_located(court_loc))
    court_target_lower = court.strip().lower()
    court_options = sel_court.find_elements(By.TAG_NAME, "option")
    print(f"[COURT] Опций: {len(court_options)}, ищу: {court!r}")

    court_value = None
    for opt in court_options:
        if (opt.text or "").strip().lower() == court_target_lower:
            court_value = opt.get_attribute("value") or ""
            print(f"[COURT] Точное: {opt.text!r} → {court_value!r}")
            break

    if not court_value:
        for opt in court_options:
            text = (opt.text or "").strip().lower()
            if court_target_lower in text or text in court_target_lower:
                court_value = opt.get_attribute("value") or ""
                print(f"[COURT] Частичное: {opt.text!r} → {court_value!r}")
                break

    if not court_value:
        raise RuntimeError(f"[COURT] Не удалось сопоставить суд '{court}'")

    _select_by_value_robust(driver, court_loc, court_value, "Судебный орган", timeout=timeout)
    click_next_and_wait_payment(driver, timeout_click=30, timeout_payment=90)
    print("✔ Перешёл на вкладку 'Оплата'")
    print("=== Блок 8 завершён ===")


# ========= ОПЛАТА =========

def load_payment_data(row):
    wb = load_workbook(EXCEL_RESP_FILE, data_only=True)
    ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active

    fio = str(ws[f"{EXCEL_FIO_COL}{row}"].value or "").strip()
    iin = str(ws[f"{EXCEL_IIN_COL}{row}"].value or "").strip()
    sum_ = ws[f"{EXCEL_SUM_COL}{row}"].value
    duty = ws[f"{EXCEL_DUTY_COL}{row}"].value

    def num_to_str(v):
        if v is None:
            return "0"
        s = str(v).replace(" ", "").replace("\u00A0", "").replace(",", ".")
        return s

    sum_str = num_to_str(sum_)
    duty_str = num_to_str(duty)
    print(f"[EXCEL] ФИО: {fio!r}, ИИН: {iin!r}, Сумма: {sum_str}, Пошлина: {duty_str}")
    return fio, iin, sum_str, duty_str


def get_latest_cases_root():
    base = CASES_BASE_DIR
    if not os.path.isdir(base):
        raise FileNotFoundError(f"Базовая папка не найдена: {base}")

    candidates = []
    for name in os.listdir(base):
        path = os.path.join(base, name)
        if not os.path.isdir(path):
            continue
        m = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", name)
        if not m:
            continue
        try:
            folder_date = datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            continue
        candidates.append((folder_date, path))

    if not candidates:
        raise FileNotFoundError(f"В {base} нет папок с датами")

    candidates.sort(reverse=True)
    latest = candidates[0][1]
    print(f"[ФАЙЛЫ] Корневая папка дел: {latest}")
    return latest


def find_case_folder(fio, iin):
    root = get_latest_cases_root()
    if not os.path.isdir(root):
        raise FileNotFoundError(f"Папка с делами не найдена: {root}")

    candidates_iin = []
    candidates_fio = []
    fio_upper = fio.upper()

    for name in os.listdir(root):
        path = os.path.join(root, name)
        if not os.path.isdir(path):
            continue
        if iin and iin in name:
            candidates_iin.append(path)
        elif fio and fio_upper.split()[0] in name.upper():
            candidates_fio.append(path)

    if candidates_iin:
        folder = sorted(candidates_iin)[0]
        print(f"[ФАЙЛЫ] Папка по ИИН: {folder}")
        return folder
    if candidates_fio:
        folder = sorted(candidates_fio)[0]
        print(f"[ФАЙЛЫ] Папка по ФИО: {folder}")
        return folder

    raise FileNotFoundError(f"Не найдена папка по ИИН={iin!r} / ФИО={fio!r}")


def find_duty_pdf(folder):
    for name in os.listdir(folder):
        if not name.lower().endswith(".pdf"):
            continue
        low = name.lower()
        if any(sub in low for sub in ("оспошлин", "госпошлин")):
            full = os.path.join(folder, name)
            print(f"[ФАЙЛЫ] Квитанция: {full}")
            return full
    raise FileNotFoundError(f"В {folder} не найден PDF с Госпошлиной")


@timed_step("Блок 9: вкладка 'Оплата'")
def block_payment_fill_and_next(driver, row, timeout=60):
    w = WebDriverWait(driver, timeout)

    try:
        kbk_select = w.until(EC.presence_of_element_located(
            (By.CSS_SELECTOR, "select[id$='selectKbk']")
        ))
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", kbk_select)
        print("✅ Вкладка 'Оплата' найдена")
    except TimeoutException:
        raise TimeoutException("Не вижу select КБК")

    fio, iin, sum_str, duty_str = load_payment_data(row)

    _select_by_value_robust(
        driver, (By.CSS_SELECTOR, "select[id$='selectKbk']"), "2", "КБК", timeout=timeout
    )

    sum_input = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "input[id$=':edit-totalSum']")
    ))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", sum_input)
    sum_input.clear()
    sum_input.send_keys(sum_str)
    print(f"✔ Сумма иска: {sum_str}")

    duty_input = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "input[id$=':edit-duty']")
    ))
    duty_input.clear()
    duty_input.send_keys(duty_str)
    print(f"✔ Госпошлина: {duty_str}")

    folder = find_case_folder(fio, iin)
    pdf_path = find_duty_pdf(folder)

    file_input = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "span[id$='selectPaymentScanUploader1'] input[type='file']")
    ))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", file_input)
    file_input.send_keys(pdf_path)
    print(f"✔ Госпошлина загружена: {pdf_path}")

    # Ждём загрузки PDF
    time.sleep(10)

    next_btn = w.until(EC.element_to_be_clickable((
        By.XPATH,
        "//a[contains(@class,'button-orange') and normalize-space()='Далее' "
        "and contains(@onclick,'goToDocuments')]",
    )))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", next_btn)
    try:
        next_btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", next_btn)
    print("✔ Нажал 'Далее' (к документам)")
    print("=== Блок 'Оплата' завершён ===")


# ========= ДОКУМЕНТЫ =========

def find_claim_doc_and_text(folder):
    claim_path = None
    for name in os.listdir(folder):
        low = name.lower()
        if "исковое заявление" in low and low.endswith(".docx"):
            claim_path = os.path.join(folder, name)
            break
    if not claim_path:
        raise FileNotFoundError(f"В {folder} не найден .docx с 'Исковое заявление'")

    doc = Document(claim_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    print(f"[ИСК] Файл: {claim_path}")
    return claim_path, text


def get_current_attached_filenames(driver):
    FILE_LIST_XPATH = "//div[contains(@id,'selectFileUploader')]//span"
    labels = []
    for s in driver.find_elements(By.XPATH, FILE_LIST_XPATH):
        txt = (s.text or "").strip()
        if txt:
            labels.append(txt)
    return labels


def normalize_name_for_compare(name):
    if not name:
        return ""
    s = name.lower().replace("ё", "е").replace("\u00a0", " ").replace(" ", "")
    return s


# ========= ПРОВЕРКА ФАКТИЧЕСКИ ПРИКРЕПЛЁННЫХ ФАЙЛОВ =========

def get_lawsuit_attached_filename(driver):
    """Читает имя фактически прикреплённого искового заявления из скрытого поля."""
    try:
        hid = driver.find_element(By.CSS_SELECTOR, "input[id$=':lawsuitScanHid']")
        return (hid.get_attribute("value") or "").strip()
    except Exception:
        return ""


def get_attached_extra_filenames(driver):
    """
    Возвращает список имён файлов, реально показанных в блоке
    'Документы, прикладываемые к исковому заявлению' (attachPanel).
    В этом списке сайт также показывает квитанцию об оплате и повторно
    исковое заявление — это нормально, лишние записи не мешают проверке.
    """
    names = []
    els = driver.find_elements(
        By.XPATH,
        "//div[contains(@id,':attachPanel')]//div[contains(@class,'flex-jc-between')]/div[1]",
    )
    for el in els:
        txt = (el.text or "").strip()
        if txt:
            names.append(txt)
    return names


def compute_missing_documents(driver, expected_extra_paths, claim_doc_path):
    """
    Сравнивает ожидаемый локальный список файлов с фактически прикреплёнными на сайте.
    Возвращает (missing_extra_paths, claim_ok, lawsuit_val_on_site).
    """
    actual_extra_raw = get_attached_extra_filenames(driver)
    actual_extra_norm = {normalize_name_for_compare(n) for n in actual_extra_raw}

    missing_extra_paths = []
    for p in expected_extra_paths:
        name = os.path.basename(p)
        if normalize_name_for_compare(name) not in actual_extra_norm:
            missing_extra_paths.append(p)

    claim_name = os.path.basename(claim_doc_path)
    lawsuit_val = get_lawsuit_attached_filename(driver)
    claim_ok = normalize_name_for_compare(claim_name) == normalize_name_for_compare(lawsuit_val)
    # запасной вариант: исковое иногда также попадает в attachPanel
    if not claim_ok:
        claim_ok = normalize_name_for_compare(claim_name) in actual_extra_norm

    return missing_extra_paths, claim_ok, lawsuit_val


def reupload_extra_files(driver, file_paths, timeout=20):
    """Повторно отправляет список файлов через input[type=file] блока доп. документов."""
    if not file_paths:
        return
    file_input = WebDriverWait(driver, timeout).until(
        EC.presence_of_element_located(
            (By.CSS_SELECTOR, "div[id$='selectFileUploader'] input[type='file']")
        )
    )
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", file_input)
    file_input.send_keys("\n".join(file_paths))
    print(f"🔄 Повторно отправил {len(file_paths)} недостающих файлов")
    time.sleep(5.0)


def reupload_claim_file(driver, claim_doc_path, timeout=20):
    """Повторно отправляет исковое заявление, если оно не прикрепилось."""
    claim_file_input = WebDriverWait(driver, timeout).until(
        EC.presence_of_element_located(
            (By.CSS_SELECTOR, "div[id$='selectLawsuitScanUploader'] input[type='file']")
        )
    )
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", claim_file_input)
    claim_file_input.send_keys(claim_doc_path)
    print("🔄 Повторно отправил исковое заявление")
    time.sleep(5.0)


def verify_and_fix_attachments(driver, all_files, claim_doc_path, max_retries=1, timeout=20):
    """
    Сверяет реально прикреплённые файлы со списком ожидаемых.
    Если чего-то не хватает — пытается догрузить недостающее (max_retries раз),
    и только после этого бросает RuntimeError с понятным перечнем проблем.
    """
    print("\n[ПРОВЕРКА ВЛОЖЕНИЙ] Сверяю фактически прикреплённые файлы...")

    for attempt in range(max_retries + 1):
        # ждём стабилизации списка (на случай, если файлы ещё дозагружаются)
        try:
            WebDriverWait(driver, timeout).until(
                lambda d: len(get_attached_extra_filenames(d)) > 0 or len(all_files) == 0
            )
        except TimeoutException:
            pass
        time.sleep(1.0)

        missing_extra, claim_ok, lawsuit_val = compute_missing_documents(
            driver, all_files, claim_doc_path
        )

        print(
            f"[ПРОВЕРКА] Ожидалось доп. файлов: {len(all_files)}, "
            f"не хватает: {len(missing_extra)}"
        )
        print(
            f"[ПРОВЕРКА] Исковое заявление на странице: {lawsuit_val!r} "
            f"({'OK' if claim_ok else 'НЕ СОВПАДАЕТ'})"
        )

        if not missing_extra and claim_ok:
            print("✅ Все документы прикреплены и совпадают со списком")
            return

        if attempt >= max_retries:
            problems = []
            if missing_extra:
                names = [os.path.basename(p) for p in missing_extra]
                problems.append(f"не прикрепились доп. файлы: {names}")
            if not claim_ok:
                problems.append(
                    f"исковое заявление не совпадает: ожидалось "
                    f"{os.path.basename(claim_doc_path)!r}, на странице {lawsuit_val!r}"
                )
            raise RuntimeError("Проверка вложений не пройдена: " + "; ".join(problems))

        print(
            f"⚠ Обнаружены расхождения — пробую догрузить "
            f"(попытка {attempt + 2}/{max_retries + 1})..."
        )
        if missing_extra:
            reupload_extra_files(driver, missing_extra, timeout=timeout)
        if not claim_ok:
            reupload_claim_file(driver, claim_doc_path, timeout=timeout)
        dismiss_filetype_reject_alert(driver, timeout=3)


def dismiss_filetype_reject_alert(driver, timeout=2):
    try:
        btn = WebDriverWait(driver, timeout).until(
            EC.element_to_be_clickable((
                By.CSS_SELECTOR,
                "input.button.button-primary[type='button'][value='OK']"
                "[onclick*='hideFileTypeRejectAlert']"
            ))
        )
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", btn)
        try:
            btn.click()
        except Exception:
            driver.execute_script("arguments[0].click();", btn)
        try:
            WebDriverWait(driver, 3).until(EC.staleness_of(btn))
        except Exception:
            pass
        print("✔ Закрыл алерт неверного формата")
        return True
    except TimeoutException:
        return False


@timed_step("Блок 10: электронный бланк + документы")
def block_documents_fill_and_next(driver, row, timeout=120):
    print("\n========== БЛОК 10: ЭЛЕКТРОННЫЙ БЛАНК + ДОКУМЕНТЫ ==========")
    w = WebDriverWait(driver, timeout)

    desc_area = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "textarea[id$='edit-plaint-description']")
    ))
    add_area = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "textarea[id$='edit-plaint-additional']")
    ))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", desc_area)
    print("✅ Вкладка документов найдена")

    fio, iin, _, _ = load_payment_data(row)
    folder = find_case_folder(fio, iin)
    claim_doc_path, claim_text = find_claim_doc_and_text(folder)
    print(f"[DOCX] Символов в иске: {len(claim_text)}")

    # Заполняем текстовые поля
    for area, label in [
        (desc_area, "Исковые требования"),
        (add_area, "Обстоятельства и доказательства"),
    ]:
        driver.execute_script(
            "arguments[0].value = arguments[1];"
            " arguments[0].dispatchEvent(new Event('input', {bubbles: true}));",
            area, claim_text,
        )
        print(f"✔ Текст иска → {label}")

    # Загружаем исковое заявление
    claim_file_input = w.until(EC.presence_of_element_located(
        (By.CSS_SELECTOR, "div[id$='selectLawsuitScanUploader'] input[type='file']")
    ))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", claim_file_input)
    claim_file_input.send_keys(claim_doc_path)
    print(f"✔ Иск загружен: {claim_doc_path}")
    time.sleep(5.0)

    # Собираем все остальные файлы
    files_with_mtime = []
    for name in os.listdir(folder):
        full = os.path.join(folder, name)
        if not os.path.isfile(full):
            continue
        if full == claim_doc_path:
            continue
        if "госпошлин" in name.lower():
            continue
        try:
            mtime = os.path.getmtime(full)
        except Exception:
            mtime = 0
        files_with_mtime.append((mtime, full))

    files_with_mtime.sort(key=lambda x: x[0], reverse=True)
    all_files = [p for _, p in files_with_mtime]
    total = len(all_files)
    print(f"[ДОКУМЕНТЫ] К загрузке: {total} файлов")

    if total > 0:
        file_input = w.until(EC.presence_of_element_located(
            (By.CSS_SELECTOR, "div[id$='selectFileUploader'] input[type='file']")
        ))
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", file_input)

        # ====== ЗАГРУЖАЕМ ВСЕ ФАЙЛЫ ОДНОЙ КОМАНДОЙ ======
        all_paths = "\n".join(all_files)
        file_input.send_keys(all_paths)
        print(f"✔ Отправил {total} файлов одной командой")

        # Считаем общий размер для расчёта таймаута
        total_mb = sum(os.path.getsize(p) / (1024 * 1024) for p in all_files)
        # ~5 секунд на МБ + 30 секунд базовых, но не менее 60 и не более 300
        wait_total = max(60, min(300, int(total_mb * 5) + 30))
        print(f"[ДОКУМЕНТЫ] Общий размер: {total_mb:.1f} МБ, жду до {wait_total}с")

        # Ждём пока хоть что-то появится в списке
        def _any_file_appeared(d):
            try:
                # Пробуем разные локаторы для списка файлов
                for xpath in [
                    "//div[contains(@id,'selectFileUploader')]//li",
                    "//div[contains(@id,'selectFileUploader')]//span[string-length(normalize-space(.)) > 3]",
                    "//ul[contains(@id,'selectFileUploader')]//li",
                    "//div[contains(@class,'ui-fileupload-files')]//span",
                ]:
                    els = d.find_elements(By.XPATH, xpath)
                    if els:
                        return True
            except Exception:
                pass
            return False

        try:
            WebDriverWait(driver, wait_total).until(_any_file_appeared)
            print("✔ Файлы появились в списке")
        except TimeoutException:
            print(f"⚠ Файлы не появились в списке за {wait_total}с — продолжаю всё равно")

        # Закрываем алерт если появился
        dismiss_filetype_reject_alert(driver, timeout=3)

        # Финальная пауза чтобы все файлы успели загрузиться на сервер
        time.sleep(5.0)

        # Показываем что получилось
        print("\n[ДОКУМЕНТЫ] Итого в списке:")
        for xpath in [
            "//div[contains(@id,'selectFileUploader')]//li",
            "//div[contains(@id,'selectFileUploader')]//span[string-length(normalize-space(.)) > 3]",
        ]:
            els = driver.find_elements(By.XPATH, xpath)
            if els:
                for el in els:
                    txt = (el.text or "").strip()
                    if txt:
                        print(f" • {txt}")
                break
    else:
        print("[ДОКУМЕНТЫ] Нет дополнительных файлов.")

    # ====== ПРОВЕРКА ЧТО ВСЁ РЕАЛЬНО ПРИКРЕПИЛОСЬ ======
    verify_and_fix_attachments(
        driver,
        all_files=all_files,
        claim_doc_path=claim_doc_path,
        max_retries=1,
    )

    # Нажимаем Далее
    next_btn = w.until(EC.element_to_be_clickable((
        By.XPATH,
        "//a[contains(@class,'button-orange') and normalize-space()='Далее' "
        "and contains(@onclick,'goToSign')]",
    )))
    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", next_btn)
    try:
        next_btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", next_btn)
    print("✔ Нажал 'Далее' (goToSign)")
    print("=== Блок 10 завершён ===")

# ========= ПРОВЕРКА ПАКЕТА ДОКУМЕНТОВ =========

REQUIRED_DOC_PATTERNS = {
    "Госпошлина": ["госпошлин"],
    "Доверенность представителя": ["доверен", "представител"],
    "Договор цессии": ["договор", "цессии"],
    "Реестр договора цессии": ["реестр", "цессии"],
    "Досудебная претензия": ["досудебн", "претенз"],
    "Исковое заявление (.docx)": ["исковое заявление", ".docx"],
    "Исполнительная надпись": ["исполнительная", "надпис"],
    "Оферта": ["оферта"],
    "Постановление об отмене": ["постановление", "отмене"],
    "Приказ на директора": ["приказ", "директор"],
    "Расчёт задолженности": ["расч", "задолж"],
    "Уведомление об уступке": ["уведомлен", "уступк"],
    "Устав ТОО": ["устав", "тоо"],
    "Скрин отправки Досудебной претензии": ["скрин", "отправк", "досудебн", "претенз"],
}
EXCLUDE_FOR_BEREKE = {"Уведомление об уступке", "Реестр договора цессии"}
def get_product_for_row(row):
    """Читает значение продукта должника из колонки B (PRODUCT_COL)."""
    wb = load_workbook(EXCEL_RESP_FILE, data_only=True)
    ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active
    raw = ws[f"{PRODUCT_COL}{row}"].value
    return str(raw or "").strip()


@timed_step("Проверка пакета документов")
def check_required_docs_for_case(row):
    print("\n========== ПРОВЕРКА ПАКЕТА ДОКУМЕНТОВ ==========")
    fio, iin, _, _ = load_payment_data(row)
    folder = find_case_folder(fio, iin)
    print(f"[CHECK] Папка: {folder}")

    product = get_product_for_row(row)
    is_bereke = "bereke" in product.lower()
    print(f"[CHECK] Продукт: {product!r} → Bereke Bank: {is_bereke}")

    required_patterns = REQUIRED_DOC_PATTERNS
    if is_bereke:
        required_patterns = {
            k: v for k, v in REQUIRED_DOC_PATTERNS.items()
            if k not in EXCLUDE_FOR_BEREKE
        }
        print(f"[CHECK] Bereke Bank — пропускаю проверку: {sorted(EXCLUDE_FOR_BEREKE)}")

    files = [
        f.lower().replace("ё", "е")
        for f in os.listdir(folder)
        if os.path.isfile(os.path.join(folder, f))
    ]
    print("[CHECK] Файлы:")
    for f in files:
        print(" -", f)

    missing = []
    for human_name, patterns in required_patterns.items():
        ok = any(all(p in fname for p in patterns) for fname in files)
        if not ok:
            missing.append(human_name)

    if missing:
        print("\n❌ Не хватает:")
        for m in missing:
            print(" -", m)
        raise RuntimeError("Пакет документов неполный.")

    print("\n✅ Пакет документов полный.")

# ========= ORCHESTRATOR =========
# -*- coding: utf-8 -*-

# ========= ИСПРАВЛЕННЫЕ ФУНКЦИИ =========
# Вставьте эти функции вместо оригинальных в ваш скрипт

# Ключевые изменения:
# 1. check_required_docs_for_case вынесена ДО цикла перезапусков браузера
# 2. Добавлена запись статуса в Excel (колонка S)
# 3. input() заменён на таймаут 60 секунд
# 4. Добавлена пауза между строками чтобы сайт не банил
# 5. Улучшена обработка "мягких" ошибок (неполный пакет не тратит попытки браузера)

from openpyxl import load_workbook
import time
import logging

STATUS_COL = "S"   # Колонка для записи статуса обработки

def verify_participant_added(driver, role_keyword, timeout=15):
    """
    Проверяет что участник с нужной ролью появился в таблице участников.
    role_keyword: 'ответчик', 'истец', 'представитель'
    """
    locators = [
        f"//td[contains(translate(normalize-space(.),'АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ','абвгдеёжзийклмнопрстуфхцчшщъыьэюя'),'{role_keyword}')]",
        f"//span[contains(translate(normalize-space(.),'АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ','абвгдеёжзийклмнопрстуфхцчшщъыьэюя'),'{role_keyword}')]",
        f"//div[contains(translate(normalize-space(.),'АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ','абвгдеёжзийклмнопрстуфхцчшщъыьэюя'),'{role_keyword}')]",
    ]
    end = time.time() + timeout
    while time.time() < end:
        for xpath in locators:
            els = driver.find_elements(By.XPATH, xpath)
            if els:
                print(f"✅ Участник '{role_keyword}' найден на странице")
                return True
        time.sleep(0.5)
    raise RuntimeError(f"❌ Участник '{role_keyword}' НЕ появился на странице за {timeout}с — форма не сохранилась")


def write_status_to_excel(row, status_text):
    """Записывает статус обработки строки в Excel."""
    try:
        wb = load_workbook(EXCEL_RESP_FILE)
        ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active
        ws[f"{STATUS_COL}{row}"] = status_text
        wb.save(EXCEL_RESP_FILE)
        print(f"[EXCEL] Статус строки {row}: {status_text}")
    except Exception as e:
        log_warn(f"Не удалось записать статус в Excel для строки {row}: {e}")


def run_full_flow_for_row(driver, row):
    print(f"\n=== ОБРАБОТКА СТРОКИ {row} ===")

    login(driver)
    go_to_send_docs(driver)
    send_claim(driver)

    block3_fill_case_and_org_participant(driver)
    verify_participant_added(driver, "истец")        # ← проверка
    print("=== Блок 3 завершён ===")

    block4_open_respondent_form(driver)
    block5_fill_person_from_excel_and_save(driver, row=row)
    verify_participant_added(driver, "ответчик")     # ← проверка
    
    block6_open_representative_form(driver)
    block7_fill_representative_from_default(driver)
    verify_participant_added(driver, "представитель") # ← проверка

    ensure_character_and_simple_process(driver)
    print("=== Параметры дела восстановлены ===")

    block8_fill_recipient_from_excel_and_next(driver, row=row)
    block_payment_fill_and_next(driver, row=row)
    block_documents_fill_and_next(driver, row=row)

    print(f"\n=== СТРОКА {row} ЗАВЕРШЕНА ===\n")


def main():
    wb = load_workbook(EXCEL_RESP_FILE, data_only=True)
    ws = wb[EXCEL_RESP_SHEET] if EXCEL_RESP_SHEET else wb.active

    rows_to_process = []
    for row in range(EXCEL_FIRST_ROW, ws.max_row + 1):
        iin_val = ws[f"{EXCEL_RESP_COL}{row}"].value
        if iin_val and str(iin_val).strip():
            rows_to_process.append(row)

    if not rows_to_process:
        raise RuntimeError("В Excel нет строк с данными.")

    print(f"[EXCEL] Строк для обработки: {len(rows_to_process)}: {rows_to_process}")

    driver = init_driver()
    failed_rows = []
    skipped_rows = []   # строки с неполным пакетом документов

    try:
        for row in rows_to_process:
            print("\n" + "=" * 40)
            print(f"=== СТРОКА {row} ===")
            print("=" * 40 + "\n")

            # ШАГ 1: Проверка документов ДО попыток с браузером.
            # Если пакет неполный — пропускаем строку сразу, без перезапуска браузера.
            try:
                check_required_docs_for_case(row)
            except (FileNotFoundError, RuntimeError) as doc_err:
                msg = f"Пропускаю строку {row}: неполный пакет — {doc_err}"
                log_warn(msg)
                print(f"\n⚠ {msg}")
                write_status_to_excel(row, f"ПРОПУЩЕНО: {doc_err}")
                skipped_rows.append(row)
                continue   # сразу к следующей строке, браузер не трогаем

            # ШАГ 2: Попытки с браузером
            success = False

            for restart_attempt in range(MAX_BROWSER_RESTARTS + 1):
                try:
                    run_full_flow_for_row(driver, row)
                    success = True
                    write_status_to_excel(row, "УСПЕШНО")
                    break

                except Exception as e:
                    err_msg = str(e)
                    log_err(
                        f"Строка {row}, попытка {restart_attempt + 1}"
                        f"/{MAX_BROWSER_RESTARTS + 1}: {err_msg}"
                    )

                    if restart_attempt < MAX_BROWSER_RESTARTS:
                        print(
                            f"\n🔄 Перезапускаю браузер и повторяю строку {row} "
                            f"(попытка {restart_attempt + 2}/{MAX_BROWSER_RESTARTS + 1})..."
                        )
                        driver = restart_driver(driver)
                        time.sleep(3)
                    else:
                        print(
                            f"\n❌ Строка {row}: исчерпаны все "
                            f"{MAX_BROWSER_RESTARTS + 1} попытки. Перехожу к следующей."
                        )
                        write_status_to_excel(row, f"ОШИБКА: {err_msg[:200]}")

            if not success:
                failed_rows.append(row)

            # Пауза между строками чтобы не перегружать сайт
            time.sleep(3)

        # ========= ИТОГИ =========
        print("\n" + "=" * 40)
        print("=== Обработка завершена ===")

        if skipped_rows:
            print(f"\n⚠ Пропущены (неполный пакет): {skipped_rows}")
        if failed_rows:
            print(f"\n❌ Не удалось обработать (ошибки браузера): {failed_rows}")
        if not failed_rows and not skipped_rows:
            print("✅ Все строки обработаны успешно.")
        elif not failed_rows:
            print("✅ Все доступные строки обработаны (часть пропущена из-за документов).")

        log_step(f"Итого: пропущено={skipped_rows}, ошибки={failed_rows}")

    finally:
        # Вместо input() — автозакрытие через 60 секунд
        print("\nБраузер закроется автоматически через 60 секунд.")
        print("Нажмите Ctrl+C чтобы закрыть сейчас.")
        try:
            time.sleep(60)
        except KeyboardInterrupt:
            print("Закрываю по запросу...")
        safe_quit_driver(driver)



if __name__ == "__main__":
    main()