# -*- coding: utf-8 -*-
"""
Блок 9: Формирование расчётов задолженности из шаблона Word
"""
import os
from pathlib import Path
from datetime import datetime, date
import threading
import pythoncom
import win32com.client
from openpyxl import load_workbook
from docx import Document


from config import (
    MAIN_EXCEL,
    ROOT,
    TARGET_BASE
)

from utils import (
    ensure_client_folder,
    safe_log,
    safe_update_summary
)

# ═══════════════════════════════════════════════════════════════
# НАСТРОЙКИ
# ═══════════════════════════════════════════════════════════════

SHEET_NAME = "Данные для шаблонов"

TEMPLATE_DOC = rf"{ROOT}\Документы для подачи Исков\Шаблоны документов\Шаблоны по Расчетам задолженности.docx"  # шаблон для расчёта задолженности
_word_lock = threading.Lock()
# ═══════════════════════════════════════════════════════════════
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ═══════════════════════════════════════════════════════════════
# -*- coding: utf-8 -*-

log_file_path = 'log_not_copied.txt'
# ---------- ЛОГИ ----------
try:
    log_file_path
except NameError:
    os.makedirs(TARGET_BASE, exist_ok=True)
    log_file_path = os.path.join(TARGET_BASE, "лог_сбор_документов.txt")


# счётчики для расчёта задолженности
calc_found = 0
calc_total = 0
calc_not_found = []


# -------------------------------------------------
# Вспомогательные функции
# -------------------------------------------------
def nice_case(s: str) -> str:
    if not s:
        return ""
    return " ".join(word.capitalize() for word in str(s).split())


def replace_placeholders_in_paragraph(paragraph, mapping: dict):
    original_text = paragraph.text
    text = original_text
    changed = False

    for key, value in mapping.items():
        if key in text:
            text = text.replace(key, value)
            changed = True

    if not changed:
        return

    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
        target_run = paragraph.runs[0]
    else:
        target_run = paragraph.add_run(text)

def convert_to_pdf(docx_path: str, pdf_path: str):
    """Потокобезопасная конвертация docx → pdf через Word COM."""
    with _word_lock:
        pythoncom.CoInitialize()
        try:
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            try:
                doc = word.Documents.Open(docx_path)
                doc.SaveAs(pdf_path, FileFormat=17)  # 17 = wdFormatPDF
                doc.Close()
            finally:
                word.Quit()
        finally:
            pythoncom.CoUninitialize()
            
def fill_document(row_dict: dict) -> Document:
    doc = Document(TEMPLATE_DOC)
    mapping = {}

    for col_name, val in row_dict.items():
        if not col_name:
            continue

        placeholder = "{" + col_name + "}"
        lower = col_name.lower()

        if val is None or val != val:
            value_str = ""
        else:
            if isinstance(val, (datetime, date)):
                value_str = val.strftime("%d.%m.%Y")
            else:
                if "фио" in lower:
                    value_str = nice_case(val)
                else:
                    value_str = str(val)

        mapping[placeholder] = value_str

    mapping["{CurrDate}"] = datetime.now().strftime("%d.%m.%Y")

    for p in doc.paragraphs:
        replace_placeholders_in_paragraph(p, mapping)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_placeholders_in_paragraph(p, mapping)

    return doc


def find_debtor_folder(root_folder: Path, iin: str) -> Path:
    for sub in root_folder.iterdir():
        if sub.is_dir() and iin in sub.name:
            return sub

    new_folder = root_folder / iin
    new_folder.mkdir(parents=True, exist_ok=True)
    return new_folder


# -------------------------------------------------
# Основная логика
# -------------------------------------------------
def run(df_main):
    global calc_found, calc_total, calc_not_found
    count_success = 0
    count_failed  = 0
    
    import pythoncom
    pythoncom.CoInitialize()
    
    try:
        root_path = Path(TARGET_BASE)
        root_path.mkdir(parents=True, exist_ok=True)

        wb = load_workbook(MAIN_EXCEL, read_only=True, data_only=True)
        ws = wb[SHEET_NAME]
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue

            row_dict = dict(zip(headers, row))
            fio = nice_case(row[2]) if len(row) > 2 else ""
            iin = str(row[3]).strip() if len(row) > 3 and row[3] else ""

            if not fio or not iin:
                continue

            calc_total += 1
            print(f"\n[STEP CALC] {calc_total} — {fio}, {iin}")

            try:
                for key in row_dict.keys():
                    if "фио" in key.lower():
                        row_dict[key] = fio
                    if "иин" in key.lower():
                        row_dict[key] = iin

                doc = fill_document(row_dict)
                debtor_folder = find_debtor_folder(root_path, iin)

                base = f"Расчёт задолженности, {fio}, {iin}"
                docx_path = debtor_folder / (base + ".docx")
                pdf_path  = debtor_folder / (base + ".pdf")

                doc.save(docx_path)
                
                # Конвертация с явной проверкой результата
                try:
                    convert_to_pdf(str(docx_path), str(pdf_path))
                    if pdf_path.exists():
                        os.remove(docx_path)
                        print(f"✅ Сформирован PDF: {pdf_path}")
                        count_success += 1
                        calc_found += 1
                    else:
                        print(f"⚠️  PDF не создан, оставлен DOCX: {docx_path}")
                        count_failed += 1
                        calc_not_found.append(f"{fio}, {iin}")
                except Exception as conv_err:
                    print(f"[ERROR CONVERT] Ошибка конвертации для {fio}, {iin}: {conv_err}")
                    safe_log(f"[РАСЧЁТ ЗАДОЛЖЕННОСТИ] Ошибка конвертации для {fio}, {iin}: {conv_err}")
                    count_failed += 1
                    calc_not_found.append(f"{fio}, {iin}")

            except Exception as e:
                label = f"{fio}, {iin}"
                calc_not_found.append(label)
                count_failed += 1
                print(f"[ERROR CALC] Ошибка для {label}: {e}")
                safe_log(f"[РАСЧЁТ ЗАДОЛЖЕННОСТИ] Ошибка для {label}: {e}")

        wb.close()

    finally:
        pythoncom.CoUninitialize()

    safe_update_summary("РАСЧЁТ ЗАДОЛЖЕННОСТИ", {
        "found": calc_found,
        "total": calc_total,
        "not_found": calc_not_found,
    })

    print(f"\n=== ГОТОВО: расчёты задолженности сформированы ===")
    print(f"[ИТОГ] {calc_found} из {calc_total}")

    return count_success, count_failed