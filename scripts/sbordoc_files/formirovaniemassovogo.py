# blocks/formirovaniemassovogo.py
# -*- coding: utf-8 -*-
"""
Блок 15: Формирование массового заявления в Медеуский суд
Создание массовых исков (кроме Vivus) по группам до 10 ответчиков
"""

import os
import re
import pandas as pd
from pathlib import Path
from datetime import datetime, date

from openpyxl import load_workbook
from docx import Document

from config import MAIN_EXCEL, TARGET_BASE
from utils import safe_log, safe_update_summary

# ═══════════════════════════════════════════════════════════════
# НАСТРОЙКИ
# ═══════════════════════════════════════════════════════════════

SHEET_NAME = "Данные для шаблонов"
TEMPLATE_DOC_MASS = r"\\PC002\work folder\Документы для подачи Исков\Шаблоны документов\Шаблон массового иска.docx"
GROUP_SIZE = 10  # Максимум ответчиков в одном иске

# ═══════════════════════════════════════════════════════════════
# КАРТА ПЕРЕИМЕНОВАНИЯ ЗАГОЛОВКОВ
# ═══════════════════════════════════════════════════════════════

HEADERS_REMAP = {
    # При необходимости добавьте соответствия заголовков
    # "Старое название": "Новое название",
}

# Колонки с денежными суммами (форматируем как '16 487,00')
MONEY_COLUMNS = {
    "Остаток задолженности",
    "Сумма выдачи займа",
    "Основной долг",
    "Вознаграждение",
    "Штраф",
    "Исполнительная надпись (расходы)",
    "Госпошлина (расходы)",
    "ОД+проценты+штрафы",
}

# ═══════════════════════════════════════════════════════════════
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ═══════════════════════════════════════════════════════════════

def nice_case(s: str) -> str:
    """ФИО -> 'Фамилия Имя Отчество'"""
    if not s:
        return ""
    return " ".join(word.capitalize() for word in str(s).split())

def format_money(value) -> str:
    """Формат денежной суммы: 16487 -> '16 487,00'"""
    try:
        f = float(str(value).replace(" ", "").replace(",", "."))
    except Exception:
        return str(value)
    f = round(f + 1e-9, 2)
    s = f"{f:,.2f}"
    s = s.replace(",", " ")
    s = s.replace(".", ",")
    return s

def num_to_words_ru(n: int) -> str:
    """Число в русские слова (для тенге, без слова 'тенге')"""
    n = int(n)
    if n == 0:
        return "ноль"
    
    units = {
        0: ("ноль", "ноль"),
        1: ("один", "одна"),
        2: ("два", "две"),
        3: ("три", "три"),
        4: ("четыре", "четыре"),
        5: ("пять", "пять"),
        6: ("шесть", "шесть"),
        7: ("семь", "семь"),
        8: ("восемь", "восемь"),
        9: ("девять", "девять"),
    }
    teens = {
        10: "десять", 11: "одиннадцать", 12: "двенадцать",
        13: "тринадцать", 14: "четырнадцать", 15: "пятнадцать",
        16: "шестнадцать", 17: "семнадцать", 18: "восемнадцать",
        19: "девятнадцать",
    }
    tens = {
        2: "двадцать", 3: "тридцать", 4: "сорок",
        5: "пятьдесят", 6: "шестьдесят", 7: "семьдесят",
        8: "восемьдесят", 9: "девяносто",
    }
    hundreds = {
        1: "сто", 2: "двести", 3: "триста", 4: "четыреста",
        5: "пятьсот", 6: "шестьсот", 7: "семьсот",
        8: "восемьсот", 9: "девятьсот",
    }
    orders = [
        ("", "", ""),
        ("тысяча", "тысячи", "тысяч"),
        ("миллион", "миллиона", "миллионов"),
        ("миллиард", "миллиарда", "миллиардов"),
    ]
    
    words = []
    order = 0
    while n > 0:
        n, rem = divmod(n, 1000)
        if rem == 0:
            order += 1
            continue
        
        part = []
        h = rem // 100
        t_u = rem % 100
        
        if h:
            part.append(hundreds[h])
        
        if 10 <= t_u <= 19:
            part.append(teens[t_u])
            u = 0
        else:
            t = t_u // 10
            u = t_u % 10
            if t:
                part.append(tens[t])
            if u:
                fem = (order == 1)
                part.append(units[u][1 if fem else 0])
        
        if order > 0:
            if u == 1 and t_u != 11:
                form = orders[order][0]
            elif u in (2, 3, 4) and not 12 <= t_u <= 14:
                form = orders[order][1]
            else:
                form = orders[order][2]
            if form:
                part.append(form)
        
        words.append(" ".join(part))
        order += 1
    
    return " ".join(reversed(words))

def replace_substring_in_runs(runs, index_map, start, end, new_text):
    """Замена подстроки в runs без потери форматирования"""
    if start >= end:
        return
    
    start_pos = start
    end_pos = end - 1
    
    start_run_idx, start_off = index_map[start_pos]
    end_run_idx, end_off = index_map[end_pos]
    
    start_run = runs[start_run_idx]
    end_run = runs[end_run_idx]
    
    start_text = start_run.text
    end_text = end_run.text
    
    prefix = start_text[:start_off]
    suffix = end_text[end_off + 1:]
    
    if start_run_idx == end_run_idx:
        start_run.text = prefix + new_text + suffix
    else:
        start_run.text = prefix + new_text
        end_run.text = suffix
        for ri in range(start_run_idx + 1, end_run_idx):
            runs[ri].text = ""

def replace_placeholders_in_paragraph(paragraph, mapping: dict):
    """Замена плейсхолдеров {Колонка_N} в абзаце"""
    runs = paragraph.runs
    if not runs:
        return
    
    full_text = "".join(r.text for r in runs)
    if not full_text:
        return
    
    if not any(k in full_text for k in mapping.keys()):
        return
    
    index_map = []
    for ri, r in enumerate(runs):
        for ci in range(len(r.text)):
            index_map.append((ri, ci))
    
    replacements = []
    for key, value in mapping.items():
        if not key:
            continue
        search_from = 0
        while True:
            pos = full_text.find(key, search_from)
            if pos == -1:
                break
            replacements.append((pos, pos + len(key), str(value)))
            search_from = pos + len(key)
    
    if not replacements:
        return
    
    replacements.sort(key=lambda x: x[0], reverse=True)
    
    for start, end, new_text in replacements:
        replace_substring_in_runs(runs, index_map, start, end, new_text)

def fill_sum_in_words(doc: Document):
    """Заполнение {сумма прописью} без потери форматирования"""
    pattern = re.compile(r'(\d[\d\s]*,\d{2})\s*\(\{сумма прописью\}\)')
    
    for p in doc.paragraphs:
        runs = p.runs
        if not runs:
            continue
        
        full_text = "".join(r.text for r in runs)
        if "{сумма прописью}" not in full_text:
            continue
        
        matches = list(pattern.finditer(full_text))
        if not matches:
            continue
        
        index_map = []
        for ri, r in enumerate(runs):
            for ci in range(len(r.text)):
                index_map.append((ri, ci))
        
        for m in reversed(matches):
            start, end = m.span()
            num_str = m.group(1)
            digits = num_str.replace(" ", "").replace("\u00A0", "")
            if "," in digits:
                int_part = digits.split(",")[0]
            else:
                int_part = digits
            try:
                n = int(int_part)
            except Exception:
                n = 0
            words = num_to_words_ru(n)
            new_sub = f"{num_str} ({words})"
            
            replace_substring_in_runs(runs, index_map, start, end, new_sub)

def cleanup_unused_respondents(doc: Document, count: int, max_count: int = 10):
    """Удаление блоков по ответчикам, которых нет в текущем пуле"""
    
    # 1. Шапка "Ответчики"
    for n in range(count + 1, max_count + 1):
        i = 0
        while i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            if f"Ответчик №{n}:" in p.text:
                for _ in range(5):
                    if i < len(doc.paragraphs):
                        p_del = doc.paragraphs[i]
                        p_del._element.getparent().remove(p_del._element)
                    else:
                        break
                continue
            i += 1
    
    # 2. Фактическая часть до ПРОШУ
    i = 0
    while i < len(doc.paragraphs):
        paras = doc.paragraphs
        if i >= len(paras):
            break
        
        p = paras[i]
        text = p.text or ""
        
        m = re.search(r"Ответчик\s*№\s*(\d+)", text)
        if not m:
            i += 1
            continue
        
        num = int(m.group(1))
        
        if num <= count or num > max_count:
            i += 1
            continue
        
        start = i
        j = i + 1
        while j < len(paras):
            t2 = paras[j].text or ""
            t2_strip = t2.strip()
            
            if re.search(r"^\s*\d+\.\s*Ответчик\s*№\s*\d+", t2_strip):
                break
            
            if (t2_strip.startswith("Согласно ст. 339") or
                t2_strip.startswith("На основании вышеизложенного") or
                t2_strip.startswith("ПРОШУ:")):
                break
            
            j += 1
        
        for _ in range(j - start):
            p_del = doc.paragraphs[start]
            p_del._element.getparent().remove(p_del._element)
    
    # 3. Хвост в разделе "ПРОШУ"
    for n in range(count + 1, max_count + 1):
        mark1 = f"Ответчик №{n}"
        mark2 = f"Ответчика №{n}"
        to_delete = []
        for p in doc.paragraphs:
            txt = p.text or ""
            if mark1 in txt or mark2 in txt:
                to_delete.append(p)
        for p in to_delete:
            try:
                p._element.getparent().remove(p._element)
            except Exception:
                pass

def fill_document_mass(group_rows: list, norm_headers: list) -> Document:
    """Заполнение массового документа"""
    doc = Document(TEMPLATE_DOC_MASS)
    mapping = {}
    
    for idx, row_dict in enumerate(group_rows, start=1):
        for col_name in norm_headers:
            if not col_name:
                continue
            
            value = row_dict.get(col_name, "")
            
            if value is None or value != value:
                value_str = ""
            elif isinstance(value, (datetime, date)):
                value_str = value.strftime("%d.%m.%Y")
            else:
                col_lower = col_name.lower()
                if "фио" in col_lower:
                    value_str = nice_case(value)
                elif col_name in MONEY_COLUMNS:
                    value_str = format_money(value)
                else:
                    value_str = str(value)
            
            placeholder = "{" + f"{col_name}_{idx}" + "}"
            mapping[placeholder] = value_str
    
    for idx in range(len(group_rows) + 1, GROUP_SIZE + 1):
        for col_name in norm_headers:
            if not col_name:
                continue
            placeholder = "{" + f"{col_name}_{idx}" + "}"
            mapping.setdefault(placeholder, "")
    
    mapping["{CurrDate}"] = datetime.now().strftime("%d.%m.%Y")
    mapping["{RespondentCount}"] = str(len(group_rows))
    
    for p in doc.paragraphs:
        replace_placeholders_in_paragraph(p, mapping)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_placeholders_in_paragraph(p, mapping)
    
    fill_sum_in_words(doc)
    cleanup_unused_respondents(doc, len(group_rows), max_count=GROUP_SIZE)
    
    return doc

# ═══════════════════════════════════════════════════════════════
# ГЛАВНАЯ ФУНКЦИЯ БЛОКА
# ═══════════════════════════════════════════════════════════════

def run(df_main):
    """
    Блок 15: Формирование массового заявления в Медеуский суд
    
    Args:
        df_main: DataFrame с данными клиентов (не используется, читаем из Excel напрямую)
    
    Returns:
        tuple: (count_success, count_failed)
    """
    
    print("\n=== БЛОК 15: Формирование массового заявления ===\n")
    
    base_path = Path(TARGET_BASE)
    base_path.mkdir(parents=True, exist_ok=True)
    
    mass_docs_created = 0
    mass_docs_failed = 0
    
    try:
        # Читаем из Excel напрямую с листа "Данные для шаблонов"
        wb = load_workbook(MAIN_EXCEL, read_only=True, data_only=True)
        ws = wb[SHEET_NAME]
        
        raw_headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        
        print(f"[DEBUG] Найдено заголовков: {len(raw_headers)}")
        
        norm_headers = [HEADERS_REMAP.get(h, h) for h in raw_headers]
        
        fio_header = norm_headers[2] if len(norm_headers) > 2 else "ФИО"
        iin_header = norm_headers[3] if len(norm_headers) > 3 else "ИИН"
        
        rows_for_mass = []
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            
            # Пропускаем Vivus
            prod_value = row[1] if len(row) > 1 else None
            if prod_value and "vivus" in str(prod_value).lower():
                continue
            
            raw_row_dict = dict(zip(raw_headers, row))
            row_dict = {}
            for k, v in raw_row_dict.items():
                if not k:
                    continue
                new_k = HEADERS_REMAP.get(k, k)
                row_dict[new_k] = v
            
            fio = nice_case(row[2]) if len(row) > 2 and row[2] else ""
            iin = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            
            if not fio or not iin:
                continue
            
            for key in list(row_dict.keys()):
                if not key:
                    continue
                lk = key.lower()
                if "фио" in lk:
                    row_dict[key] = fio
                if "иин" in lk:
                    row_dict[key] = iin
            
            row_dict[fio_header] = fio
            row_dict[iin_header] = iin
            
            rows_for_mass.append(row_dict)
        
        wb.close()
        
        if not rows_for_mass:
            print("[INFO] Не найдено строк для массовых исков (кроме Vivus)")
            safe_update_summary("МАССОВЫЕ ИСКИ (КРОМЕ VIVUS)", {
                "found": 0,
                "total": 0,
                "not_found": [],
            })
            return 0, 0
        
        print(f"[INFO] Всего должников (кроме Vivus): {len(rows_for_mass)}")
        
        # Разбиваем на группы по GROUP_SIZE
        for start in range(0, len(rows_for_mass), GROUP_SIZE):
            group_rows = rows_for_mass[start:start + GROUP_SIZE]
            group_index = start // GROUP_SIZE + 1
            
            print(f"[STEP] Группа {group_index}: ответчиков {len(group_rows)}")
            
            try:
                doc = fill_document_mass(group_rows, norm_headers)
                
                base_name = f"Массовый иск в Медеуски суд, пул {group_index}"
                for ch in r'\/:*?"<>|':
                    base_name = base_name.replace(ch, "_")
                
                # Сохраняем в папку соответствующего Пула
                pool_dir = base_path / f"Пул {group_index}"
                pool_dir.mkdir(parents=True, exist_ok=True)
                
                docx_path = pool_dir / (base_name + ".docx")
                doc.save(docx_path)
                
                print(f"[OK] Сформирован: {docx_path}")
                mass_docs_created += 1
                
            except Exception as e:
                msg = f"[ERROR] Ошибка при формировании группы {group_index}: {e}"
                print(msg)
                mass_docs_failed += 1
                safe_log(f"[МАССОВЫЕ ИСКИ] {msg}")
                continue
        
        print("\n=== ГОТОВО: массовые иски сформированы ===")
        print(f"[ИТОГ] Создано документов: {mass_docs_created}")
        
    except Exception as e:
        print(f"[ERROR] Критическая ошибка: {e}")
        safe_log(f"[МАССОВЫЕ ИСКИ] Критическая ошибка: {e}")
        mass_docs_failed = 1
    
    print(f"\n--- ИТОГ МАССОВЫХ ИСКОВ ---")
    print(f"✔️ Создано: {mass_docs_created}")
    print(f"❌ Ошибок: {mass_docs_failed}")
    
    safe_update_summary("МАССОВЫЕ ИСКИ (КРОМЕ VIVUS)", {
        "found": mass_docs_created,
        "total": len(rows_for_mass) if 'rows_for_mass' in locals() else 0,
        "not_found": [],
    })
    
    # ✅ ОБЯЗАТЕЛЬНО ВОЗВРАЩАЕМ ДВА ЗНАЧЕНИЯ!
    return mass_docs_created, mass_docs_failed

