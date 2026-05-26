# -*- coding: utf-8 -*-
"""
Возврат госпошлины — оркестратор.

Принимает ДВА файла:
  1) --pdf_file    — «Реквизиты …».pdf: пачка платёжных поручений на оплату
                     госпошлины (по одному на страницу);
  2) --reestr_file — «Реестр_на_возврат_госпошлины_*.xlsx», ранее сформированный
                     scripts/reestr_gosposhliny.py и отправленный в WhatsApp
                     (A=Уникальный номер, B=ИИН, C=ФИО, D=Госпошлина 3%,
                      E=УГД, F=БИН УГД, G=адрес, H=Суд).

Этапы (--stage 1|2|3, по умолчанию 3):
  Этап 1 — распарсить PDF, сопоставить плательщиков с реестром по ФИО (точное
           совпадение), сформировать LoansImport.xlsx в папке автоимпорта
           Дельты (CREDENTIALS['path_crm']) + копия в out/; отчёт сопоставления.
  Этап 2 — по EID сопоставленных строк вытянуть из БД crm данные в формате
           «Отчёта по отменам» (2 SQL — листы «Отмены» и «Данные для шаблонов»),
           адрес/регион/суд подставить из реестра, сохранить
           out/Отчёт_реестр_ГП_*.xlsx; затем запустить scripts/sbor.py на этот
           файл (--excel_path).
  Этап 3 — запустить scripts/podacha_iska_v2.py на тот же файл (--excel_path).

Старые scripts/sbor.py и scripts/podacha_iska_v2.py НЕ меняются — откат = просто
не пользоваться этим скриптом.
"""

import os
import re
import sys
import time
import shutil
import argparse
import subprocess
import unicodedata
from pathlib import Path
from datetime import datetime


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workdir', type=str, default=None)
    parser.add_argument('--company_id', type=str, default=None)
    parser.add_argument('--pdf_file', type=str, default=None)
    parser.add_argument('--reestr_file', type=str, default=None)
    parser.add_argument('--stage', type=str, default='3',
                        help="до какого этапа доходить: "
                             "1 — только LoansImport; "
                             "2 — + отчёт по БД и сбор документов; "
                             "3 — + подача иска (по умолчанию 3)")
    return parser.parse_known_args()[0]


args = parse_args()

if not args.company_id or not args.company_id.strip():
    print("❌ ОШИБКА: не передан --company_id — компания не определена, запуск остановлен.")
    sys.exit(1)
os.environ['COMPANY_ID'] = args.company_id.strip()

import pdfplumber
import pandas as pd
from openpyxl import Workbook, load_workbook

from config import CREDENTIALS, DB_COMPANY_FILTER, ROOT

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# ============================================================
# НАСТРОЙКИ / ЛОГ
# ============================================================

WORKDIR = Path(args.workdir) if args.workdir else Path(".")
OUT_DIR = WORKDIR / "out"
OUT_DIR.mkdir(parents=True, exist_ok=True)

PATH_CRM = CREDENTIALS.get("path_crm") or ""
LOANS_IMPORT_XLSX = (
    Path(PATH_CRM) / "LoansImport.xlsx" if PATH_CRM
    else OUT_DIR / "LoansImport.xlsx"
)

# Отчёт в формате «Отчёта по отменам» — его читают sbor.py и podacha_iska_v2.py.
OTMENY_REPORT_XLSX = OUT_DIR / f"Отчёт_реестр_ГП_{datetime.now():%Y%m%d_%H%M}.xlsx"

# БД crm — те же реквизиты, что в scripts/reestr_gosposhliny.py.
DB_CONFIG = {
    "server":   "DBSRV",
    "database": "crm",
    "username": "user",
    "password": "Log1cF",
}


def log(msg):
    print(time.strftime("[%H:%M:%S]"), msg, flush=True)


# ============================================================
# ЗАГОЛОВКИ LoansImport (по образцу C:\...\Downloads\LoansImport.xlsx)
# ============================================================

LOANS_IMPORT_HEADERS = [
    "Уникальный номер сделки",                       # 1  A
    "Продукт",                                       # 2
    "Номер договора займа",                          # 3
    "Дата выдачи займа",                             # 4
    "Фамилия",                                       # 5
    "Имя",                                           # 6
    "Отчество",                                      # 7
    "ИИН",                                           # 8  H
    "Сумма займа",                                   # 9
    "ОД _при покупке",                               # 10
    "Вознаграждение _при покупке",                   # 11
    "Пени/штрафы _при покупке",                      # 12
    "Остаток задолженности _при покупке",            # 13
    "Услуги нотариуса _расходы после цессии",        # 14
    "Услуги медиатора _расходы после цессии",        # 15
    "Госпошлина по суду _расходы после цессии",      # 16 P
    "Кредитор",                                      # 17
    "Статус СУСН",                                   # 18
    "Номер паспорта",                                # 19
    "Дата выдачи паспорта",                          # 20
    "Срок займа",                                    # 21
    "DPD _при покупке",                              # 22
    "Сумма оплаты до цессии",                        # 23
    "Дата последнего платежа до цессии",             # 24
    "E-mail",                                        # 25
    "Номер договора цессии",                         # 26
    "Дата договора цессии",                          # 27
    "Планируемый месяц урегулирования",              # 28
    "Дата смерти",                                   # 29
    "Дата начала воинской службы",                   # 30
    "Сумма покупки долга",                           # 31
    "Годовая эффективная ставка",                    # 32
    "Номинальная ставка",                            # 33
    "Форма Медиации",                                # 34
    "Услуги нотариуса _при покупке",                 # 35
    "Услуги медиатора _при покупке",                 # 36
    "Госпошлина по суду _при покупке",               # 37
    "Иные расходы _при покупке",                     # 38
    "Сумма возврата переплаты",                      # 39
    "Дата возврата средств по переплате",            # 40
    "Орган выдачи паспорта",                         # 41
    "ФИО Ответственного",                            # 42
    "Сумма последнего платежа до цессии",            # 43
    "Дата окончания воинской службы",                # 44
    "ОД _расходы после цессии",                      # 45
    "Вознаграждение _расходы после цессии",          # 46
    "Пени/штрафы _расходы после цессии",             # 47
    "Иные расходы _расходы после цессии",            # 48
    "Остаток задолженности _расходы после цессии",   # 49
    "Комментарий по графикам",                       # 50
    "Дата возврата госпошлины",                      # 51
]
LI_COL_UNIQ = 1     # A
LI_COL_IIN = 8      # H
LI_COL_GP = 16      # P


# ============================================================
# ЭТАП 1: PDF -> сопоставление с реестром -> LoansImport.xlsx
# ============================================================

# Каз. буквы -> рус. (как в scripts/reestr_gosposhliny.py::normalize_kz_text)
_KZ_FOLD = str.maketrans({
    "қ": "к", "ғ": "г", "ң": "н", "ү": "у", "ұ": "у",
    "ө": "о", "ә": "а", "һ": "х", "і": "и", "ё": "е",
})


def norm_fio_key(value) -> str:
    """Ключ для ТОЧНОГО сопоставления ФИО: регистр не важен, каз.->рус.,
    порядок слов не важен, лишние пробелы схлопнуты."""
    s = unicodedata.normalize("NFC", str(value or "")).lower().translate(_KZ_FOLD)
    s = re.sub(r"[^а-яёa-z\s-]", " ", s)
    tokens = sorted(t for t in re.split(r"[\s-]+", s) if t)
    return " ".join(tokens)


def norm_amount(whole: str, kop: str) -> float:
    """'8 894' + '16' -> 8894.16"""
    return round(int(re.sub(r"\D", "", whole)) + int(kop) / 100.0, 2)


def norm_uniq(value) -> str:
    s = str(value if value is not None else "").strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def norm_iin(value) -> str:
    s = re.sub(r"\D", "", str(value if value is not None else ""))
    return s.zfill(12) if s else ""


def parse_gp_pdf(pdf_path: Path) -> list[dict]:
    """Каждая страница = одно платёжное поручение. Возвращает список:
    {page, doc_no, doc_date, fio, amount, kbk, beneficiary_bin}."""
    # Строка назначения платежа: "<ФИО> Сумма 8 894-16 теңге 108126"
    re_naz = re.compile(
        r"(?m)^(.+?)\s+Сумма\s+([\d ]+)[-–](\d{2})\s+теңге\s+(\d+)\s*$"
    )
    re_no = re.compile(r"№\s*(\d+)\s*от\s*(\d{2}\.\d{2}\.\d{4})")
    re_bek = re.compile(r"Бе[Кк]/[КК]Бе:\s*\d+\s*(\d{12})")

    rows: list[dict] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            m_naz = re_naz.search(text)
            if not m_naz:
                rows.append({"page": i, "fio": None, "amount": None,
                             "doc_no": None, "doc_date": None,
                             "kbk": None, "beneficiary_bin": None,
                             "raw_ok": False})
                continue
            m_no = re_no.search(text)
            m_bek = re_bek.search(text)
            rows.append({
                "page": i,
                "fio": m_naz.group(1).strip(),
                "amount": norm_amount(m_naz.group(2), m_naz.group(3)),
                "kbk": m_naz.group(4),
                "doc_no": m_no.group(1) if m_no else None,
                "doc_date": m_no.group(2) if m_no else None,
                "beneficiary_bin": m_bek.group(1) if m_bek else None,
                "raw_ok": True,
            })
    return rows


def load_reestr(reestr_path: Path) -> tuple[dict, list[dict]]:
    """Читает реестр -> (индекс norm_fio_key -> [строки], список всех строк).
    Строка: {row, uniq, iin, fio, gp, ugd, ugd_bin, address, court}."""
    wb = load_workbook(str(reestr_path), data_only=True)
    ws = wb["Реестр"] if "Реестр" in wb.sheetnames else wb.active

    headers = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h:
            headers[str(h).strip().lower()] = c

    def col(*substrings, default=None):
        for key, idx in headers.items():
            if all(s in key for s in substrings):
                return idx
        return default

    c_uniq = col("уникальн") or 1
    c_iin = col("иин") or 2
    c_fio = col("фио") or 3
    c_gp = col("госпошл") or 4
    c_ugd = col("угд") or 5
    c_bin = col("бин") or 6
    c_addr = col("адрес") or 7
    c_court = col("судебн", "орган") or 8

    index: dict[str, list[dict]] = {}
    all_rows: list[dict] = []
    for r in range(2, ws.max_row + 1):
        fio = ws.cell(row=r, column=c_fio).value
        uniq = ws.cell(row=r, column=c_uniq).value
        if not fio and not uniq:
            continue
        rec = {
            "row": r,
            "uniq": norm_uniq(uniq),
            "iin": norm_iin(ws.cell(row=r, column=c_iin).value),
            "fio": str(fio or "").strip(),
            "gp": ws.cell(row=r, column=c_gp).value,
            "ugd": ws.cell(row=r, column=c_ugd).value,
            "ugd_bin": ws.cell(row=r, column=c_bin).value,
            "address": ws.cell(row=r, column=c_addr).value,
            "court": ws.cell(row=r, column=c_court).value,
        }
        all_rows.append(rec)
        index.setdefault(norm_fio_key(fio), []).append(rec)
    return index, all_rows


def build_loans_import(pdf_rows: list[dict], reestr_index: dict) -> dict:
    """Сопоставляет платежи с реестром по точному ключу ФИО и пишет
    LoansImport.xlsx. Несопоставленные — пропускаются, дубли по ФИО —
    все строки реестра с той же суммой. Возвращает сводку."""
    matched: list[dict] = []      # {pdf, reestr}
    unmatched: list[dict] = []    # pdf-строки без совпадения / без парсинга
    ambiguous: list[dict] = []    # pdf-строка -> несколько строк реестра

    for pr in pdf_rows:
        if not pr.get("raw_ok") or not pr.get("fio"):
            unmatched.append({**pr, "reason": "не распарсилась страница PDF"})
            continue
        hits = reestr_index.get(norm_fio_key(pr["fio"]), [])
        if not hits:
            unmatched.append({**pr, "reason": "ФИО не найдено в реестре"})
            continue
        for rr in hits:
            matched.append({"pdf": pr, "reestr": rr})
        if len(hits) > 1:
            ambiguous.append({**pr, "count": len(hits),
                              "uniqs": [h["uniq"] for h in hits]})

    wb = Workbook()
    wsx = wb.active
    wsx.title = "Лист1"
    for c, h in enumerate(LOANS_IMPORT_HEADERS, start=1):
        wsx.cell(row=1, column=c, value=h)

    out_row = 2
    for m in matched:
        wsx.cell(row=out_row, column=LI_COL_UNIQ, value=m["reestr"]["uniq"])
        wsx.cell(row=out_row, column=LI_COL_IIN, value=m["reestr"]["iin"])
        wsx.cell(row=out_row, column=LI_COL_GP, value=m["pdf"]["amount"])
        out_row += 1

    os.makedirs(LOANS_IMPORT_XLSX.parent, exist_ok=True)
    wb.save(LOANS_IMPORT_XLSX)
    log(f"✅ LoansImport сформирован: {LOANS_IMPORT_XLSX} (строк: {out_row - 2})")

    if LOANS_IMPORT_XLSX.parent != OUT_DIR:
        try:
            shutil.copy2(LOANS_IMPORT_XLSX, OUT_DIR / "LoansImport.xlsx")
        except Exception as e:
            log(f"⚠ Не удалось сохранить копию LoansImport.xlsx в out/: {e}")

    return {"matched": matched, "unmatched": unmatched, "ambiguous": ambiguous,
            "written_rows": out_row - 2}


def write_report(pdf_rows: list[dict], summary: dict):
    """Отчёт сопоставления в out/ — что попало в LoansImport, что нет."""
    path = OUT_DIR / f"GP_import_отчёт_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Сопоставление"
    ws.append(["Стр. PDF", "№ поручения", "Дата", "ФИО (PDF)", "Сумма",
               "БИН бенефициара", "Статус", "Уникальный номер", "ИИН",
               "Комментарий"])

    matched_by_page: dict[int, list[dict]] = {}
    for m in summary["matched"]:
        matched_by_page.setdefault(m["pdf"]["page"], []).append(m["reestr"])
    unmatched_pages = {u["page"]: u.get("reason", "") for u in summary["unmatched"]}
    ambiguous_pages = {a["page"] for a in summary["ambiguous"]}

    for pr in pdf_rows:
        base = [pr["page"], pr.get("doc_no"), pr.get("doc_date"),
                pr.get("fio"), pr.get("amount"), pr.get("beneficiary_bin")]
        if pr["page"] in matched_by_page:
            recs = matched_by_page[pr["page"]]
            status = "СОПОСТАВЛЕНО" + (" (дубль)" if pr["page"] in ambiguous_pages else "")
            comment = "" if len(recs) == 1 else f"строк реестра: {len(recs)}"
            for rr in recs:
                ws.append(base + [status, rr["uniq"], rr["iin"], comment])
        else:
            ws.append(base + ["ПРОПУЩЕНО", "", "",
                              unmatched_pages.get(pr["page"], "нет совпадения")])

    wb.save(path)
    log(f"📄 Отчёт сопоставления: {path}")
    return path


def stage1_pdf_to_loans_import() -> dict:
    log("=== ЭТАП 1: PDF → LoansImport ===")
    pdf_path = Path(args.pdf_file) if args.pdf_file else None
    reestr_path = Path(args.reestr_file) if args.reestr_file else None

    if not pdf_path or not pdf_path.exists():
        log(f"❌ Не найден PDF госпошлин: {pdf_path}")
        sys.exit(1)
    if not reestr_path or not reestr_path.exists():
        log(f"❌ Не найден файл реестра: {reestr_path}")
        sys.exit(1)

    pdf_rows = parse_gp_pdf(pdf_path)
    log(f"📥 Страниц в PDF: {len(pdf_rows)}; распознано платежей: "
        f"{sum(1 for r in pdf_rows if r.get('raw_ok'))}")

    reestr_index, reestr_rows = load_reestr(reestr_path)
    log(f"📋 Строк в реестре: {len(reestr_rows)}")

    summary = build_loans_import(pdf_rows, reestr_index)
    write_report(pdf_rows, summary)

    log(f"— сопоставлено платежей: "
        f"{len({m['pdf']['page'] for m in summary['matched']})}")
    log(f"— пропущено (нет в реестре / не распознано): {len(summary['unmatched'])}")
    if summary["unmatched"]:
        for u in summary["unmatched"]:
            log(f"    · стр.{u['page']}: {u.get('fio') or '—'} — {u.get('reason')}")
    if summary["ambiguous"]:
        log(f"— ⚠ дубли по ФИО (записаны все строки): {len(summary['ambiguous'])}")
        for a in summary["ambiguous"]:
            log(f"    · стр.{a['page']}: {a['fio']} → {a['count']} строк "
                f"({', '.join(a['uniqs'])})")

    return summary


# ============================================================
# ЭТАП 2: отчёт по БД (формат «Отчёта по отменам») + сбор документов
# ============================================================

# Регион из адреса СК — как в scripts/reestr_gosposhliny.py.
_REGION_NAMES = [
    "Акмолинская область", "Актюбинская область", "Алматинская область",
    "город Алматы", "город Астана", "Атырауская область",
    "Восточно-Казахстанская область", "Жамбылская область",
    "Западно-Казахстанская область", "Карагандинская область",
    "Костанайская область", "Кызылординская область",
    "Мангистауская область", "Область Абай", "Область Жетісу",
    "Область Ұлытау", "Павлодарская область",
    "Северо-Казахстанская область", "Туркестанская область", "город Шымкент",
]


def region_from_address(address) -> str:
    if not address:
        return ""
    up = str(address).upper()
    if re.search(r"\bАСТАНА\b", up) or "НУР-СУЛТАН" in up:
        return "город Астана"
    if re.search(r"\bШЫМКЕНТ\b", up):
        return "город Шымкент"
    if re.search(r"\bАЛМАТЫ\b", up) and "АЛМАТИНСКАЯ ОБЛАСТЬ" not in up:
        return "город Алматы"
    low = up.lower()
    for reg in _REGION_NAMES:
        if reg.lower() in low:
            return reg
    return ""


# Заголовки листов — точь-в-точь как в «Отчёте по отменам».
OTMENY_HEADERS = [
    "Уникальный номер", "Продукт", "ФИО", "ИИН", "Остаток задолженности",
    "Статус исполнительного документа", "Статус АИС ОИП", "Статус кредита",
    "Кол-во займов", "Остаток без расходов за надпись и без госпошлины",
    "Госпошлина 3%", "УГД", "БИН УГД", "Оплаченная ранее госпошлина",
    "Актуальный адрес с Судебного кабинета", "Регион с Судебного кабинета",
    "Судебный орган с Судебного кабинета",
]
TEMPLATES_HEADERS = [
    "Уникальный номер", "Продукт", "ФИО", "ИИН", "Остаток задолженности",
    "Email", "Адрес регистрации", "Дата минус 7 дней", "Мобильный", "Кредитор",
    "Номер договора займа", "Дата выдачи займа", "Сумма выдачи займа",
    "Номер договора цессии", "Дата договора цессии", "Основной долг",
    "Вознаграждение", "Штраф", "Исполнительная надпись (расходы)",
    "Госпошлина (расходы)", "Срок займа", "Дата окончания займа",
    "Судебный орган", "Дата договора цессии + 5 дней", "ОД+проценты+штрафы",
    "Остаток без расходов за надпись и без госпошлины",
]

# --- SQL: «Отчёт по отменам» (лист «Отмены»), фильтр только по EID ------------
SQL_OTMENY = """
WITH LoanCounts AS (
    SELECT c.F293 AS ИИН, COUNT(*) AS LoanCount
    FROM loans l JOIN clients c ON l.CID = c.ID
    WHERE l.F209 = N'{company}'
    GROUP BY c.F293
)
SELECT
    l.EID AS [Уникальный номер],
    dF246.F249 AS [Продукт],
    (SELECT STRING_AGG(CASE WHEN LEN(value) > 1
            THEN UPPER(LEFT(value,1)) + LOWER(SUBSTRING(value,2,LEN(value)-1))
            ELSE UPPER(value) END, ' ')
     FROM STRING_SPLIT(c.FIO, ' ')) AS [ФИО],
    c.F293 AS [ИИН],
    CAST(l.F34 AS DECIMAL(18,2)) AS [Остаток задолженности],
    constF235.Caption AS [Статус исполнительного документа],
    constF236.Caption AS [Статус АИС ОИП],
    s.Caption AS [Статус кредита],
    ISNULL(lc.LoanCount, 0) AS [Кол-во займов],
    CAST((ISNULL(l.F34,0) - ISNULL(l.F315,0) - ISNULL(l.F307,0)) AS DECIMAL(18,2))
        AS [Остаток без расходов за надпись и без госпошлины],
    CAST(CEILING(CASE
            WHEN (ISNULL(l.F34,0) - ISNULL(l.F315,0) - ISNULL(l.F307,0)) > 0
                THEN (ISNULL(l.F34,0) - ISNULL(l.F315,0) - ISNULL(l.F307,0)) * 0.03
            ELSE 0 END) AS DECIMAL(18,0)) AS [Госпошлина 3%],
    NULL AS [УГД],
    NULL AS [БИН УГД],
    CAST(ISNULL(l.F307, 0) AS DECIMAL(18,2)) AS [Оплаченная ранее госпошлина]
FROM loans l (NOLOCK)
JOIN clients c (NOLOCK) ON l.CID = c.ID
JOIN states s ON s.ID = l.State
JOIN Dictionary dF246 ON dF246.ID = l.F246
LEFT JOIN Constants constF235 ON constF235.ID = l.F235
JOIN Constants constF236 ON constF236.ID = l.F236
LEFT JOIN LoanCounts lc ON c.F293 = lc.ИИН
WHERE l.EID IN ({eids}) AND l.F209 = N'{company}'
GROUP BY l.EID, dF246.F249, c.FIO, c.F293, l.F34, constF235.Caption,
    constF236.Caption, s.Caption, lc.LoanCount, l.F315, l.F307;
"""

# --- SQL: «Данные для шаблонов», фильтр только по EID ------------------------
SQL_TEMPLATES = """
DECLARE @base datetime2 = DATEADD(day, -7, CAST(GETDATE() AS date));
WITH LoanCounts AS (
    SELECT c.F293 AS ИИН, COUNT(*) AS LoanCount
    FROM loans l JOIN clients c ON l.CID = c.ID
    WHERE l.F209 = N'{company}'
    GROUP BY c.F293
)
SELECT
    l.EID AS [Уникальный номер],
    dF246.F249 AS [Продукт],
    (SELECT STRING_AGG(CASE WHEN LEN(value) > 1
            THEN UPPER(LEFT(value,1)) + LOWER(SUBSTRING(value,2,LEN(value)-1))
            ELSE UPPER(value) END, ' ')
     FROM STRING_SPLIT(c.FIO, ' ')) AS [ФИО],
    c.F293 AS [ИИН],
    CAST(l.F34 AS DECIMAL(18,2)) AS [Остаток задолженности],
    MAX(e.Email) AS [Email],
    NULL AS [Адрес регистрации],
    CONVERT(varchar(19), DATEADD(SECOND,
        32400 + CAST(RAND(CHECKSUM(NEWID())) * 36000 AS int), @base), 120)
        AS [Дата минус 7 дней],
    MAX(ph.PhoneNumber) AS [Мобильный],
    constF234.Caption AS [Кредитор],
    l.F287 AS [Номер договора займа],
    l.F15  AS [Дата выдачи займа],
    l.F200 AS [Сумма выдачи займа],
    l.F69  AS [Номер договора цессии],
    l.F68  AS [Дата договора цессии],
    CAST((ISNULL(l.F34,0) - ISNULL(l.F315,0) - ISNULL(l.F307,0)
        - ISNULL(l.F38,0) - ISNULL(l.F8,0)) AS DECIMAL(18,2)) AS [Основной долг],
    l.F8  AS [Вознаграждение],
    l.F38 AS [Штраф],
    CAST(ISNULL(l.F315, 0) AS DECIMAL(18,2)) AS [Исполнительная надпись (расходы)],
    CAST(ISNULL(l.F307, 0) AS DECIMAL(18,2)) AS [Госпошлина (расходы)],
    l.F43 AS [Срок займа],
    DATEADD(DAY, l.F43, l.F15) AS [Дата окончания займа],
    NULL AS [Судебный орган],
    DATEADD(DAY, 5, l.F68) AS [Дата договора цессии + 5 дней],
    (l.F32 + l.F8 + l.F38) AS [ОД+проценты+штрафы],
    CAST((ISNULL(l.F34,0) - ISNULL(l.F315,0) - ISNULL(l.F307,0)) AS DECIMAL(18,2))
        AS [Остаток без расходов за надпись и без госпошлины]
FROM loans l (NOLOCK)
JOIN clients c (NOLOCK) ON l.CID = c.ID
JOIN states s ON s.ID = l.State
JOIN Dictionary dF246 ON dF246.ID = l.F246
JOIN Constants constF234 ON constF234.ID = l.F234
LEFT JOIN Constants constF235 ON constF235.ID = l.F235
JOIN Constants constF236 ON constF236.ID = l.F236
LEFT JOIN LoanCounts lc ON c.F293 = lc.ИИН
LEFT JOIN Emails e (NOLOCK) ON e.CID = l.CID
OUTER APPLY (
    SELECT TOP 1 p.PhoneNumber FROM phones p
    WHERE p.CID = c.ID AND p.State = 165 ORDER BY p.ID DESC
) ph
WHERE l.EID IN ({eids}) AND l.F209 = N'{company}'
GROUP BY l.EID, dF246.F249, c.FIO, c.F293, l.F34, constF234.Caption,
    constF235.Caption, constF236.Caption, s.Caption, lc.LoanCount,
    l.F315, l.F307, l.F287, l.F15, l.F200, l.F69, l.F68, l.F32, l.F8, l.F38,
    l.F43;
"""


def _db_connect():
    conn_str = (
        "DRIVER={ODBC Driver 18 for SQL Server};"
        f"SERVER={DB_CONFIG['server']};"
        f"DATABASE={DB_CONFIG['database']};"
        f"UID={DB_CONFIG['username']};"
        f"PWD={DB_CONFIG['password']};"
        "TrustServerCertificate=yes;Encrypt=no;"
    )
    import pyodbc
    return pyodbc.connect(conn_str, timeout=30)


def _eid(v) -> str:
    """EID к строке из одних цифр. l.EID из БД crm приходит как Decimal
    ('107020.00'), поэтому берём только целую часть до точки/запятой —
    иначе '107020.00' превращалось в '10702000' и сопоставление с реестром
    ломалось (адрес/суд не подставлялись, отчёт был неполный)."""
    s = str(v if v is not None else "").strip()
    s = re.split(r"[.,]", s, maxsplit=1)[0]
    return re.sub(r"\D", "", s)


def _eid_sql_list(eids) -> str:
    clean = sorted({_eid(e) for e in eids if _eid(e)})
    return ", ".join(clean) if clean else "NULL"


def stage2_build_report(matched: list[dict]) -> Path:
    """Тянет данные из БД по EID сопоставленных строк, добавляет адрес/регион/суд
    из реестра, пишет отчёт в формате «Отчёта по отменам» (2 листа)."""
    log("=== ЭТАП 2: отчёт по БД (формат «Отчёта по отменам») ===")

    # EID -> (реестр-строка, сумма ГП из PDF); при дублях берём первую строку
    by_eid: dict[str, dict] = {}
    for m in matched:
        eid = _eid(m["reestr"]["uniq"])
        if not eid:
            continue
        by_eid.setdefault(eid, {"reestr": m["reestr"], "gp_paid": m["pdf"]["amount"]})

    if not by_eid:
        log("❌ Нет сопоставленных EID — отчёт не формируется, этапы 2–3 пропущены.")
        return None

    eids = list(by_eid)
    company = DB_COMPANY_FILTER.replace("'", "''")
    eid_in = _eid_sql_list(eids)

    log(f"🔌 Подключение к БД crm ({len(eids)} EID)...")
    conn = _db_connect()
    try:
        df_otm = pd.read_sql(SQL_OTMENY.format(company=company, eids=eid_in), conn)
        df_tpl = pd.read_sql(SQL_TEMPLATES.format(company=company, eids=eid_in), conn)
    finally:
        conn.close()
    log(f"✅ БД: лист «Отмены» — {len(df_otm)} строк, «Данные для шаблонов» — {len(df_tpl)} строк")

    in_db = {_eid(v) for v in df_otm["Уникальный номер"].tolist()}
    missing = sorted(set(eids) - in_db)
    if missing:
        log(f"⚠ Нет в БД по фильтру компании: {', '.join(missing)}")

    wb = Workbook()
    ws_o = wb.active
    ws_o.title = "Отмены"
    ws_o.append(OTMENY_HEADERS)
    for _, row in df_otm.iterrows():
        eid = _eid(row["Уникальный номер"])
        rec = by_eid.get(eid, {})
        rr = rec.get("reestr", {})
        addr = rr.get("address")
        vals = [row.get(h) for h in OTMENY_HEADERS[:14]]
        # K (кол. 11) «Госпошлина 3%» -> точная сумма из PDF
        if rec.get("gp_paid") is not None:
            vals[10] = rec["gp_paid"]
        vals += [addr, region_from_address(addr), rr.get("court")]
        ws_o.append(vals)

    ws_t = wb.create_sheet("Данные для шаблонов")
    ws_t.append(TEMPLATES_HEADERS)
    for _, row in df_tpl.iterrows():
        eid = _eid(row["Уникальный номер"])
        rr = by_eid.get(eid, {}).get("reestr", {})
        vals = [row.get(h) for h in TEMPLATES_HEADERS]
        vals[6] = rr.get("address")   # «Адрес регистрации» <- адрес из реестра
        vals[22] = rr.get("court")    # «Судебный орган»    <- суд из реестра
        ws_t.append(vals)

    wb.save(OTMENY_REPORT_XLSX)
    log(f"💾 Отчёт сохранён: {OTMENY_REPORT_XLSX}")
    return OTMENY_REPORT_XLSX


def _run_subscript(script_rel: str, extra: list[str]) -> int:
    cmd = [sys.executable, "-u", str(PROJECT_ROOT / script_rel),
           "--company_id", str(args.company_id).strip()]
    if args.workdir:
        cmd += ["--workdir", str(args.workdir)]
    cmd += extra
    log(f"▶ {' '.join(cmd)}")
    proc = subprocess.run(cmd, cwd=str(PROJECT_ROOT))
    log(f"↩ {Path(script_rel).name} завершился с кодом {proc.returncode}")
    return proc.returncode


def stage2a_place_gp_pdf():
    """Блок «Госпошлины» в sbor.py (sbordoc_files/gosposhliny.py) сам сканирует
    {ROOT}\\Госпошлины\\<год>\\<самая свежая датированная папка> и не знает про
    --pdf_file. Кладём загруженный PDF в папку «ГП от <сегодня>» — она станет
    самой свежей, и блок разложит платёжки по клиентским папкам ЭТОЙ партии."""
    if not args.pdf_file:
        return
    src = Path(args.pdf_file)
    if not src.exists():
        return
    now = datetime.now()
    gp_dir = Path(ROOT) / "Госпошлины" / str(now.year) / f"ГП от {now:%d.%m.%Y}"
    try:
        gp_dir.mkdir(parents=True, exist_ok=True)
        dst = gp_dir / f"{src.stem}_reestr_gp{src.suffix}"
        shutil.copy2(src, dst)
        log(f"📎 PDF госпошлин размещён для сбора: {dst}")
    except Exception as e:
        log(f"⚠ Не удалось разместить PDF госпошлин в {gp_dir}: {e}")


def stage2b_run_sbor(report_path: Path):
    log("=== ЭТАП 2б: сбор документов (scripts/sbor.py) ===")
    stage2a_place_gp_pdf()
    _run_subscript("scripts/sbor.py", ["--excel_path", str(report_path)])


def stage3_run_podacha(report_path: Path):
    log("=== ЭТАП 3: подача иска (scripts/podacha_iska_v2.py) ===")
    _run_subscript("scripts/podacha_iska_v2.py", ["--excel_path", str(report_path)])


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    stage = str(args.stage).strip()
    log(f"Компания: {DB_COMPANY_FILTER} | этап: до {stage}")

    summary = stage1_pdf_to_loans_import()

    if stage not in {"2", "3"}:
        log("=== ГОТОВО (этап 1) ===")
        sys.exit(0)

    if not summary["matched"]:
        log("❌ Нет сопоставленных платежей — этапы 2–3 пропущены.")
        sys.exit(0)

    report_path = stage2_build_report(summary["matched"])
    if report_path is None:
        sys.exit(1)

    stage2b_run_sbor(report_path)

    if stage == "3":
        stage3_run_podacha(report_path)

    log("=== ГОТОВО ===")
