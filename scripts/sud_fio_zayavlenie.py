# -*- coding: utf-8 -*-
"""
Судебный кабинет (portal-sot.kz) — ФИО и адрес по ИИН + AddressesImport.

Заменяет прежнее содержимое этого файла (старая логика подачи иска в
office.sud.kz + автоподстановка ФИО/адреса участника — office.sud.kz для
этого больше не используется, см. scripts/podacha_iska_v2.py для подачи
исков сейчас). Перенесено из ноутбука fio_address_parsing_SK.ipynb, два
блока которого объединены в один проход (раньше между ними был ручной
шаг — скопировать результат блока 1 в "... — копия.xlsx" для блока 2):

1) Читает Excel с ИИН (колонка B, со 2-й строки — тот же формат, что и в
   ноутбуке), логинится на portal-sot.kz через ЭЦП (NCALayer) и по каждому
   ИИН тянет ФИО/адрес из ГБД ФЛ (API, тот же, что и в scripts/poiskvsk.py —
   вход и запрос к ГБД ФЛ переиспользованы оттуда 1:1, включая уже
   исправленный баг с обрезкой ответа до 3000 символов).
2) Из БД CRM по ИИН находит все EID сделок этой компании и строит
   AddressesImport.xlsx — сохраняет в out/ задачи и копирует в папку
   автоимпорта CRM компании (CREDENTIALS['path_crm']).

ВАЖНО про вход по ЭЦП: он привязан к физическому сертификату на этой
машине и к отдельному Chrome-профилю с уже выданным разрешением
portal-sot.kz → NCALayer — см. PORTAL_SOT_BY_COMPANY в scripts/config.py.
Пока настроено ТОЛЬКО для company_id=1 (Омега); для остальных компаний
скрипт останавливается с понятной ошибкой, а не подписывает чужим
сертификатом.
"""

import os
import sys
import argparse

# =========================
# АРГУМЕНТЫ ЗАПУСКА — ДО ИМПОРТА config!
# =========================

def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workdir', default=None)
    parser.add_argument('--company_id', type=str, default=None)
    parser.add_argument('--excel_file', default=None,
                         help='Excel с ИИН в колонке B (со 2-й строки)')
    return parser.parse_known_args()[0]


args = parse_args()
if not args.company_id or not args.company_id.strip():
    print("❌ ОШИБКА: не передан --company_id — компания не определена, запуск остановлен.")
    sys.exit(1)
os.environ['COMPANY_ID'] = args.company_id.strip()  # ← ОБЯЗАТЕЛЬНО до import config

if not args.excel_file or not args.excel_file.strip():
    print("❌ ОШИБКА: не передан --excel_file (Excel с ИИН в колонке B).")
    sys.exit(1)

# =========================
# ДОБАВЛЯЕМ КОРЕНЬ ПРОЕКТА В sys.path
# =========================

project_root = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(project_root)
if project_root not in sys.path:
    sys.path.insert(0, project_root)

# =========================
# УЧЁТНЫЕ ДАННЫЕ ПО КОМПАНИИ — ИЗ config.py
# =========================

from config import CREDENTIALS, _company_id, DB_COMPANY_FILTER, PORTAL_SOT_BY_COMPANY

print(f"🏢 Судебный кабинет (ФИО, адрес) для компании ID={_company_id}")

_portal_cfg = PORTAL_SOT_BY_COMPANY.get(_company_id)
if not _portal_cfg:
    print(
        f"❌ ОШИБКА: для компании ID={_company_id} не настроен вход на portal-sot.kz "
        f"(нет записи в PORTAL_SOT_BY_COMPANY в scripts/config.py). Нужны: свой "
        f"ЭЦП-сертификат на этой машине и отдельный Chrome-профиль с разрешением "
        f"portal-sot.kz → NCALayer. Пока поддержана только компания '1' (Омега)."
    )
    sys.exit(1)

PORTAL_EDS_PASSWORD = _portal_cfg['eds_password']
PORTAL_LOGIN_PASSWORD = _portal_cfg['portal_password']
PORTAL_CHROME_PROFILE = _portal_cfg['chrome_profile']

PORTAL_SOT_BASE = "https://portal-sot.kz"
GBDFL_BY_IIN_URL = PORTAL_SOT_BASE + "/api/secure/gbdfl/v2/byIin/"
NCALAYER_WS_PORT = 13579
NCALAYER_PATH = _portal_cfg.get(
    "ncalayer_path", r"C:\Users\User\AppData\Local\Programs\NCALayer\NCALayer.exe"
)

SAVE_EVERY = 20  # автосохранение отчёта каждые N обработанных строк


# =========================
# ИМПОРТЫ
# =========================

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

from openpyxl import load_workbook

from pathlib import Path
from datetime import datetime

import base64
import json
import re
import shutil
import time
import traceback

import pandas as pd
import pyodbc
import requests


# =========================
# ПАПКА ЗАПУСКА
# =========================

if args.workdir:
    JOB_DIR = Path(args.workdir)
else:
    dt_string = datetime.now().strftime("%Y-%m-%d___%H-%M")
    JOB_DIR = Path(r"C:\Users\User\Desktop") / f"sud_fio_{_company_id}_{dt_string}"

OUT_DIR = JOB_DIR / "out"
OUT_DIR.mkdir(parents=True, exist_ok=True)

REPORT_XLSX = OUT_DIR / "Адрес с СК.xlsx"


def log4(msg):
    print(time.strftime("[%H:%M:%S]"), msg, flush=True)


# =========================
# БД — EID ПО ИИН ДЛЯ ЭТОЙ КОМПАНИИ
# =========================

DB_CONFIG = {
    "server": "DBSRV",
    "database": "crm",
    "username": "user",
    "password": "Log1cF",
}


def get_db_connection():
    conn_str = (
        f"DRIVER={{ODBC Driver 18 for SQL Server}};"
        f"SERVER={DB_CONFIG['server']};"
        f"DATABASE={DB_CONFIG['database']};"
        f"UID={DB_CONFIG['username']};"
        f"PWD={DB_CONFIG['password']};"
        f"TrustServerCertificate=yes;"
    )
    return pyodbc.connect(conn_str, timeout=30)


def load_eids_by_iin() -> dict:
    """Все EID сделок ЭТОЙ компании (l.F209 = DB_COMPANY_FILTER) по ИИН."""
    print("🔌 Подключение к БД...")
    sql = f"""
    SELECT
        c.F293 AS IIN,
        l.EID AS EID
    FROM dbo.loans l WITH (NOLOCK)
    JOIN dbo.clients c WITH (NOLOCK) ON c.ID = l.CID
    WHERE
        l.F209 = N'{DB_COMPANY_FILTER}'
        AND NULLIF(LTRIM(RTRIM(c.F293)), N'') IS NOT NULL
        AND l.EID IS NOT NULL
    ORDER BY c.F293, l.EID
    """
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute(sql)
        mapping = {}
        for iin_raw, eid_raw in cursor.fetchall():
            iin = _norm_iin(iin_raw)
            if len(iin) != 12:
                continue
            eid = str(eid_raw).strip()
            mapping.setdefault(iin, [])
            if eid not in mapping[iin]:
                mapping[iin].append(eid)
        cursor.close()
    finally:
        conn.close()
    print(f"✅ Загружено ИИН с EID: {len(mapping)}")
    return mapping


# =========================
# ВСПОМОГАТЕЛЬНОЕ
# =========================

def _norm_iin(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    try:
        if "e" in s.lower():
            s = str(int(float(s)))
    except Exception:
        pass
    s = "".join(ch for ch in s if ch.isdigit())
    return s[:12]


REGION_NAMES = [
    "Акмолинская область", "Актюбинская область", "Алматинская область",
    "город Алматы", "город Астана", "Атырауская область",
    "Восточно-Казахстанская область", "Жамбылская область",
    "Западно-Казахстанская область", "Карагандинская область",
    "Костанайская область", "Кызылординская область",
    "Мангистауская область", "Область Абай", "Область Жетісу",
    "Область Ұлытау", "Павлодарская область",
    "Северо-Казахстанская область", "Туркестанская область",
    "город Шымкент",
]


def extract_region_from_address(address: str) -> str:
    """Область/город республиканского значения по тексту адреса — по
    белому списку официальных названий. Сделано так специально: поля
    ГБД ФЛ (regDistrictNameRu/regRegionNameRu/regCity) называются не по
    смыслу того, что реально в них приходит (проверено на живых ответах
    API), поэтому доверять области конкретному полю по имени нельзя —
    надёжнее найти в готовом адресе точное совпадение с официальным
    названием региона, независимо от того, в каком поле оно пришло."""
    if not address:
        return ""
    up = str(address).upper()
    if re.search(r'\bАСТАНА\b', up) or "НУР-СУЛТАН" in up:
        return "город Астана"
    if re.search(r'\bШЫМКЕНТ\b', up):
        return "город Шымкент"
    if re.search(r'\bАЛМАТЫ\b', up) and "АЛМАТИНСКАЯ ОБЛАСТЬ" not in up:
        return "город Алматы"
    addr_low = up.lower()
    for reg in REGION_NAMES:
        if reg.lower() in addr_low:
            return reg
    return ""


# =========================
# ВХОД НА portal-sot.kz ЧЕРЕЗ ЭЦП (перенесено 1:1 из scripts/poiskvsk.py)
# =========================

def _portal_init_driver():
    opts = webdriver.ChromeOptions()
    opts.add_argument("--start-maximized")
    opts.add_argument("--disable-notifications")
    opts.add_argument("--disable-popup-blocking")
    opts.add_argument("--disable-blink-features=AutomationControlled")
    # Постоянный профиль этой компании, где уже нажато «Разрешить» для
    # portal-sot.kz → NCALayer.
    opts.add_argument(r"--user-data-dir=" + PORTAL_CHROME_PROFILE)
    driver = webdriver.Chrome(options=opts)
    driver.set_page_load_timeout(120)
    return driver


def _ncalayer_ws_open() -> bool:
    import socket
    try:
        with socket.create_connection(("127.0.0.1", NCALAYER_WS_PORT), timeout=1):
            return True
    except OSError:
        return False


def _portal_ensure_ncalayer(timeout=70):
    """NCALayer должен быть запущен ДО клика «Войти», иначе окно подписи
    не появится вовсе. Если ws-порт 13579 закрыт — запускаем NCALayer.exe."""
    if _ncalayer_ws_open():
        log4("NCALayer уже запущен (порт 13579)")
        return

    import subprocess
    if os.path.isfile(NCALAYER_PATH):
        log4(f"Запускаю NCALayer: {NCALAYER_PATH}")
        try:
            subprocess.Popen([NCALAYER_PATH], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception as e:
            log4(f"⚠ Не смог запустить NCALayer ({e}) — запустите его вручную")
    else:
        log4(f"⚠ NCALayer.exe не найден: {NCALAYER_PATH}")

    deadline = time.time() + timeout
    while time.time() < deadline:
        if _ncalayer_ws_open():
            log4("✅ NCALayer готов (порт 13579)")
            time.sleep(3)
            return
        time.sleep(1)
    raise RuntimeError(
        f"NCALayer не поднялся за {timeout}s (порт {NCALAYER_WS_PORT} закрыт). "
        f"Запустите NCALayer вручную и перезапустите скрипт."
    )


def _portal_wait_ncalayer(timeout=60):
    from pywinauto import Desktop
    desktop = Desktop(backend="uia")
    t0 = time.time()
    while time.time() - t0 < timeout:
        try:
            nca = desktop.window(title="NCALayer")
            if nca.exists() and nca.is_visible():
                return nca
        except Exception:
            pass
        time.sleep(0.4)
    raise RuntimeError("NCALayer не появился")


def _portal_wait_nca_button(title, timeout=60):
    from pywinauto import Desktop
    t0 = time.time()
    while time.time() - t0 < timeout:
        try:
            nca = Desktop(backend="uia").window(title="NCALayer")
            if nca.exists() and nca.is_visible():
                btn = nca.child_window(title=title, control_type="Button")
                if btn.exists():
                    return nca, btn
        except Exception:
            pass
        time.sleep(0.4)
    raise RuntimeError(f"Кнопка NCALayer «{title}» не появилась")


def _portal_click_login(driver):
    wait = WebDriverWait(driver, 30)
    btn = wait.until(EC.element_to_be_clickable((
        By.XPATH, "//*[self::button or self::a][contains(normalize-space(.),'Войти')]",
    )))
    try:
        btn.click()
    except Exception:
        driver.execute_script("arguments[0].click();", btn)


def _portal_jwt_claims(token: str) -> dict:
    try:
        payload = token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        return json.loads(base64.urlsafe_b64decode(payload))
    except Exception:
        return {}


def _portal_login(driver, force_fresh=False):
    _portal_ensure_ncalayer()

    driver.get(PORTAL_SOT_BASE + "/cabinet")
    time.sleep(3)

    if not force_fresh:
        access = driver.execute_script("return localStorage.getItem('access_token');")
        has_pwd_field = driver.execute_script(
            "return !!document.querySelector(\"input[type='password']\");"
        )
        if access and "/cabinet" in driver.current_url.lower() and not has_pwd_field:
            claims = _portal_jwt_claims(access)
            exp = claims.get("exp", 0) or 0
            if exp and exp > time.time() + 30:
                refresh = driver.execute_script("return localStorage.getItem('refresh_token');")
                log4(f"✅ Сессия portal-sot.kz жива (токен действует до "
                     f"{datetime.fromtimestamp(exp):%H:%M:%S})")
                return access, refresh
            log4("⚠ access_token в localStorage просрочен/битый — полный вход через ЭЦП")

    driver.get(PORTAL_SOT_BASE)
    time.sleep(2)

    nca = None
    for attempt in range(1, 3):
        _portal_click_login(driver)
        log4(f"Нажал «Войти» (попытка {attempt}), жду NCALayer…")
        try:
            nca = _portal_wait_ncalayer(45)
            break
        except RuntimeError:
            log4("⚠ Окно NCALayer не появилось — проверяю службу и пробую ещё раз")
            _portal_ensure_ncalayer()
            time.sleep(2)
    if nca is None:
        raise RuntimeError(
            "Окно подписи NCALayer так и не появилось после клика «Войти». "
            "Убедитесь, что NCALayer запущен и в нём для portal-sot.kz нажато «Разрешить»."
        )
    nca.set_focus()

    pwd = nca.child_window(title="Пароль", control_type="Edit")
    if not pwd.exists():
        pwd = nca.child_window(auto_id="JavaFX39", control_type="Edit")
    open_btn = nca.child_window(title="Открыть", control_type="Button")
    if not open_btn.exists():
        open_btn = nca.child_window(auto_id="JavaFX47", control_type="Button")

    pwd.click_input()
    time.sleep(0.2)
    pwd.type_keys("^a{BACKSPACE}")
    pwd.type_keys(PORTAL_EDS_PASSWORD, with_spaces=True, pause=0.03)
    time.sleep(0.3)
    open_btn.click_input()

    nca, sign_btn = _portal_wait_nca_button("Подписать", 60)
    nca.set_focus()
    time.sleep(0.3)
    sign_btn.click_input()
    log4("✅ ЭЦП подписана")

    portal_pwd = WebDriverWait(driver, 60).until(
        EC.visibility_of_element_located((By.XPATH, "//input[@type='password']"))
    )
    portal_pwd.clear()
    portal_pwd.send_keys(PORTAL_LOGIN_PASSWORD)

    final_btn = WebDriverWait(driver, 30).until(EC.element_to_be_clickable((
        By.XPATH, "//button[contains(normalize-space(.),'Войти')]",
    )))
    final_btn.click()

    WebDriverWait(driver, 60).until(lambda d: "/cabinet" in d.current_url.lower())

    access = driver.execute_script("return localStorage.getItem('access_token');")
    refresh = driver.execute_script("return localStorage.getItem('refresh_token');")
    if not access:
        raise RuntimeError("access_token не найден после входа")

    log4("✅ Вход на portal-sot.kz выполнен, access_token получен")
    return access, refresh


def _portal_make_session(driver, access_token):
    s = requests.Session()
    try:
        ua = driver.execute_script("return navigator.userAgent")
    except Exception:
        ua = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/152.0.0.0 Safari/537.36")
    s.headers.update({
        "Authorization": f"Bearer {access_token}",
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "ru",
        "Accept-Encoding": "gzip, deflate",
        "Origin": PORTAL_SOT_BASE,
        "Referer": PORTAL_SOT_BASE + "/",
        "User-Agent": ua,
    })
    return s


def _portal_gbdfl_raw(session, iin):
    """GET /api/secure/gbdfl/v2/byIin/<iin>. ВАЖНО: тело НЕ обрезаем — см.
    исправленный баг в scripts/poiskvsk.py (обрезка до 3000 символов ломала
    json.loads() на любом ответе длиннее лимита)."""
    url = PORTAL_SOT_BASE + "/api/secure/gbdfl/v2/byIin/" + str(iin)
    r = session.get(url, timeout=(10, 60), allow_redirects=True)
    body = r.text or ""
    return {
        "status": r.status_code,
        "statusText": r.reason or "",
        "respURL": r.url,
        "headers": "\n".join(f"{k}: {v}" for k, v in r.headers.items()),
        "bodyLen": len(body),
        "body": body,
    }


def _portal_build_address(result: dict) -> str:
    if not result:
        return ""
    parts = []
    for x in (result.get("regDistrictNameRu"), result.get("regRegionNameRu"),
              result.get("regCity"), result.get("regStreet")):
        if x is None:
            continue
        s = str(x).strip()
        if s and s.lower() != "null":
            parts.append(s)
    building = result.get("regBuilding")
    flat = result.get("regFlat")
    if building:
        parts.append(f"дом {str(building).strip()}")
    if flat:
        parts.append(f"кв. {str(flat).strip()}")
    return ", ".join(parts)


def _portal_fetch_by_iin(session, iin, retries=3):
    """При 401 бросает RuntimeError('PORTAL_TOKEN_EXPIRED')."""
    iin = _norm_iin(iin)
    if len(iin) != 12:
        raise ValueError(f"Некорректный ИИН: {iin!r}")

    last_error = None
    for attempt in range(1, retries + 1):
        try:
            env = _portal_gbdfl_raw(session, iin)
        except Exception as e:
            last_error = e
            log4(f"   ⚠ запрос ГБД ФЛ упал (попытка {attempt}/{retries}): {type(e).__name__}: {e}")
            time.sleep(min(1.5 * attempt, 5))
            continue

        status = int(env.get("status", 0) or 0)
        body = env.get("body") or ""
        body_stripped = body.strip()

        if status == 401:
            raise RuntimeError("PORTAL_TOKEN_EXPIRED")

        if status in (0, -1, -2, -3, -9, 429, 500, 502, 503, 504) or not body_stripped:
            last_error = RuntimeError(f"HTTP {status} ({env.get('statusText', '')}), bodyLen={env.get('bodyLen')}")
            log4(f"   ⚠ ГБД ФЛ ответ невалиден (попытка {attempt}/{retries}): status={status}")
            time.sleep(min(2 * attempt, 8))
            continue

        try:
            js = json.loads(body)
        except Exception:
            last_error = RuntimeError(f"не JSON, status={status}, body[:300]={body_stripped[:300]!r}")
            log4(f"   ⚠ ГБД ФЛ вернул не-JSON (попытка {attempt}/{retries}): {body_stripped[:300]!r}")
            time.sleep(min(2 * attempt, 8))
            continue

        result = js.get("result") or {}
        code = str((js.get("status") or {}).get("code") or "").strip()
        if not result:
            log4(f"   ✗ ИИН не найден в ГБД ФЛ (status={status}, code={code})")
            return {"iin": iin, "sur": "", "name": "", "patr": "", "live": "",
                     "settlement": "", "raw": {}, "status_code": code}

        return {
            "iin": str(result.get("iin") or iin),
            "sur": str(result.get("lastName") or "").strip(),
            "name": str(result.get("firstName") or "").strip(),
            "patr": str(result.get("patronymic") or "").strip(),
            "live": _portal_build_address(result),
            "settlement": str(result.get("regCity") or result.get("regRegionNameRu") or "").strip(),
            "raw": result,
            "status_code": code,
        }

    raise RuntimeError(f"GBDFL_REQUEST_FAILED: {last_error}")


# =========================
# ЭТАП 1: ФИО/адрес по ИИН -> отчёт
# =========================

def fill_people_from_portal(input_xlsx: str) -> pd.DataFrame:
    log4("=" * 90)
    log4("СУДЕБНЫЙ КАБИНЕТ — ФИО И АДРЕС ПО ИИН")
    log4("=" * 90)

    wb = load_workbook(input_xlsx)
    ws = wb.active

    headers = {
        1: "порядковый", 2: "ИИН", 3: "Фамилия", 4: "Имя", 5: "Отчество",
        6: "ФИО", 7: "Место жительства", 8: "Область", 9: "Населенный пункт",
        10: "Статус",
    }
    for col, title in headers.items():
        ws.cell(row=1, column=col).value = title

    log4("🔐 Вход на portal-sot.kz через ЭЦП...")
    driver = _portal_init_driver()
    try:
        access_token, refresh_token = _portal_login(driver)
        session = _portal_make_session(driver, access_token)

        processed = ok_count = error_count = 0

        for row in range(2, ws.max_row + 1):
            iin = _norm_iin(ws.cell(row=row, column=2).value)
            if not iin:
                continue

            processed += 1
            ws.cell(row=row, column=1).value = processed
            log4(f"#{processed} | Excel row {row} | ИИН {iin}")

            if len(iin) != 12:
                ws.cell(row=row, column=10).value = "НЕВЕРНЫЙ ИИН"
                error_count += 1
                continue

            try:
                data = _portal_fetch_by_iin(session, iin)
            except RuntimeError as e:
                if str(e) == "PORTAL_TOKEN_EXPIRED":
                    log4("   🔄 access_token истёк, повторный вход через ЭЦП")
                    access_token, refresh_token = _portal_login(driver, force_fresh=True)
                    session = _portal_make_session(driver, access_token)
                    try:
                        data = _portal_fetch_by_iin(session, iin)
                    except Exception as e2:
                        ws.cell(row=row, column=10).value = f"ОШИБКА: {e2}"
                        error_count += 1
                        log4(f"   ❌ {iin}: {e2}")
                        continue
                else:
                    ws.cell(row=row, column=10).value = f"ОШИБКА: {e}"
                    error_count += 1
                    log4(f"   ❌ {iin}: {e}")
                    continue
            except Exception as e:
                ws.cell(row=row, column=10).value = f"ОШИБКА: {e}"
                error_count += 1
                log4(f"   ❌ {iin}: {e}")
                continue

            if not data.get("raw"):
                ws.cell(row=row, column=10).value = data.get("status_code") or "НЕ НАЙДЕН"
                error_count += 1
                log4(f"   ❌ {iin}: {data.get('status_code') or 'НЕ НАЙДЕН'}")
                continue

            address = data.get("live", "")
            ws.cell(row=row, column=3).value = data.get("sur", "")
            ws.cell(row=row, column=4).value = data.get("name", "")
            ws.cell(row=row, column=5).value = data.get("patr", "")
            ws.cell(row=row, column=6).value = f"{data.get('sur', '')} {data.get('name', '')} {data.get('patr', '')}".strip()
            ws.cell(row=row, column=7).value = address
            ws.cell(row=row, column=8).value = extract_region_from_address(address)
            ws.cell(row=row, column=9).value = data.get("settlement", "")
            ws.cell(row=row, column=10).value = "OK"

            ok_count += 1
            log4(f"   ✅ {iin}: {ws.cell(row=row, column=6).value} | {address}")

            if processed % SAVE_EVERY == 0:
                wb.save(REPORT_XLSX)
                log4(f"💾 Autosave после {processed} строк")

            time.sleep(0.25)

        wb.save(REPORT_XLSX)

        log4("")
        log4("=" * 90)
        log4("✅ ЭТАП 1 ГОТОВ")
        log4(f"Обработано: {processed} | Успешно: {ok_count} | Ошибок: {error_count}")
        log4(f"Файл: {REPORT_XLSX}")
        log4("=" * 90)

    finally:
        try:
            driver.quit()
        except Exception:
            pass

    return pd.read_excel(REPORT_XLSX, dtype=str)


# =========================
# ЭТАП 2: AddressesImport.xlsx
# =========================

def clean_region_for_import(region: str) -> str:
    """
    СЕВЕРО-КАЗАХСТАНСКАЯ ОБЛАСТЬ -> Северо Казахстанская
    МАНГИСТАУСКАЯ ОБЛАСТЬ -> Мангистауская
    ОБЛАСТЬ ЖЕТІСУ -> Жетісу
    НУР-СУЛТАН -> Астана
    """
    if not region:
        return ""
    s = str(region).strip()
    if "НУР-СУЛТАН" in s.upper():
        return "Астана"
    s = s.replace("-", " ")
    s = re.sub(r"\bОБЛАСТЬ\b", " ", s, flags=re.IGNORECASE)
    s = re.sub(r"\s+", " ", s).strip()
    return s.title()


def address_after_region(full_address: str, region: str) -> str:
    """Всё в адресе ПОСЛЕ найденного региона (или разумный fallback,
    если регион в тексте не встретился дословно)."""
    if not full_address:
        return ""
    parts = [x.strip() for x in str(full_address).split(",") if x.strip()]
    if not parts:
        return ""

    region_norm = re.sub(r"\s+", " ", str(region or "").strip()).upper()
    region_index = None
    for i, part in enumerate(parts):
        if region_norm and re.sub(r"\s+", " ", part).upper() == region_norm:
            region_index = i
            break

    if region_index is not None:
        return ", ".join(parts[region_index + 1:]).strip()
    if len(parts) >= 3 and parts[0].upper() == "КАЗАХСТАН":
        return ", ".join(parts[2:]).strip()
    if len(parts) > 1:
        return ", ".join(parts[1:]).strip()
    return ""


def build_addresses_import(df_people: pd.DataFrame) -> pd.DataFrame:
    log4("")
    log4("=== ЭТАП 2: ФОРМИРОВАНИЕ AddressesImport.xlsx ===")

    eid_map = load_eids_by_iin()

    output_rows = []
    no_eid = []
    skipped = []

    for _, row in df_people.iterrows():
        iin = _norm_iin(row.get("ИИН"))
        if not iin:
            continue

        status = str(row.get("Статус") or "").strip()
        if status.upper() != "OK":
            skipped.append((iin, status))
            continue

        full_address = str(row.get("Место жительства") or "").strip()
        region_raw = str(row.get("Область") or "").strip()
        region_import = clean_region_for_import(region_raw)
        address_tail = address_after_region(full_address, region_raw)

        eids = eid_map.get(iin, [])
        if not eids:
            no_eid.append(iin)
            continue

        for eid in eids:
            output_rows.append({
                "Уникальный номер сделки": str(eid),
                "Страна": "КАЗАХСТАН",
                "Регион": region_import,
                "Город/Район": address_tail,
                "Населенный пункт": "",
                "Улица": "",
                "Дом": "",
                "Квартира": "",
                "Тип адреса": "Из судебного кабинета",
                "Тип владельца": "Клиент",
                "Статус адреса": "Актуальный",
            })

    df_out = pd.DataFrame(output_rows, columns=[
        "Уникальный номер сделки", "Страна", "Регион", "Город/Район",
        "Населенный пункт", "Улица", "Дом", "Квартира", "Тип адреса",
        "Тип владельца", "Статус адреса",
    ])

    log4(f"Исходных ИИН со статусом OK: {len(df_people[df_people['Статус'].astype(str).str.upper() == 'OK']) if 'Статус' in df_people.columns else 0}")
    log4(f"Создано строк: {len(df_out)}")
    log4(f"ИИН без EID в этой компании: {len(no_eid)}")
    log4(f"Пропущено (статус != OK): {len(skipped)}")
    if no_eid:
        log4(f"ИИН без EID (первые 50): {no_eid[:50]}")

    return df_out


def save_addresses_import(df_out: pd.DataFrame):
    local_path = OUT_DIR / "AddressesImport.xlsx"
    df_out.to_excel(local_path, index=False, sheet_name="AddressesImport")
    log4(f"✅ AddressesImport сохранён: {local_path}")

    try:
        network_dir = Path(CREDENTIALS["path_crm"])
        network_dir.mkdir(parents=True, exist_ok=True)
        network_path = network_dir / "AddressesImport.xlsx"
        shutil.copy2(local_path, network_path)
        log4(f"✅ AddressesImport скопирован в папку автоимпорта CRM: {network_path}")
    except Exception as e:
        log4(f"⚠ Не удалось скопировать AddressesImport в сетевую папку CRM: {e}")


# =========================
# ЗАПУСК
# =========================

def main():
    df_people = fill_people_from_portal(args.excel_file)
    df_out = build_addresses_import(df_people)
    save_addresses_import(df_out)

    log4("")
    log4("=" * 90)
    log4("✅ ГОТОВО")
    log4("=" * 90)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print("КРИТИЧЕСКАЯ ОШИБКА")
        traceback.print_exc()
        raise
