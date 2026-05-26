# -*- coding: utf-8 -*-
"""
Судебный кабинет — определения о медиации через HTTP/API.

Портировано из ноутбука SUD_DOWNLOAD_MEDIATION_V11_RU_KZ_DOC_DOCX.ipynb.
Поддерживает русские и казахские судебные определения, а также DOC и DOCX.
Казахское определение распознаётся по ҰЙҒАРЫМ + признакам медиации/
татуласу/бітімгершілік + утверждению (бекітілсін, бекіту туралы и др.).
Определения только о принятии и возбуждении дела отсекаются.

Отличия от ноутбука — только интеграция с интерфейсом:
- логин/пароль СК и Chrome remote-debug порт/профиль берутся по
  --company_id (как в scripts/podacha_iska_v2.py), чтобы у каждой компании
  была своя изолированная браузерная сессия и не было гонки за порт 9222;
- входной Excel передаётся файлом через веб-интерфейс (--excel_file),
  а не берётся с жёстко заданного пути на рабочем столе;
- результаты (скачанные DOCX/DOC и обновлённый Excel со статусами в
  колонках T/U) складываются в --workdir/out, чтобы их можно было
  скачать через интерфейс как обычный результат задачи.
"""

import os
import sys
import argparse
import shutil


# =========================
# АРГУМЕНТЫ ЗАПУСКА — ДО ИМПОРТА config!
# =========================

def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workdir', type=str, default=None)
    parser.add_argument('--company_id', type=str, default=None)
    parser.add_argument('--excel_file', type=str, default=None)
    return parser.parse_known_args()[0]


args = parse_args()

if not args.company_id or not args.company_id.strip():
    print("❌ ОШИБКА: не передан --company_id — компания не определена, запуск остановлен.")
    sys.exit(1)
os.environ['COMPANY_ID'] = args.company_id.strip()  # ← ОБЯЗАТЕЛЬНО до import config

if not args.excel_file or not args.excel_file.strip():
    print("❌ ОШИБКА: не передан --excel_file — нужен отчёт «Отчёт по Возврату Госпошлин» для этой компании.")
    sys.exit(1)

# =========================
# ДОБАВЛЯЕМ КОРЕНЬ ПРОЕКТА В sys.path
# =========================

project_root = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.dirname(project_root)
if project_root not in sys.path:
    sys.path.insert(0, project_root)

# =========================
# УЧЁТНЫЕ ДАННЫЕ ПО КОМПАНИИ — ИЗ config.py
# =========================

from config import CREDENTIALS, _company_id

print(f"🏢 Запуск скачивания определений о медиации для компании ID={_company_id}")

# =========================
# ОСТАЛЬНЫЕ ИМПОРТЫ
# =========================

import copy
import json
import logging
import mimetypes
import re
import time
import socket
import subprocess
import zipfile
import xml.etree.ElementTree as ET

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


# ============================================================
# НАСТРОЙКИ
# ============================================================

BASE_URL = "https://office.sud.kz"

AUTO_LOGIN = True

# СК логин/пароль — по company_id, как и в остальных скриптах, работающих
# с office.sud.kz (podacha_iska_v2.py, chsi_podacha_runner.py, sud.py).
SUD_LOGIN = CREDENTIALS['sk_login']
SUD_PASSWORD = CREDENTIALS['sk_password']

CHROME_PATH = Path(
    r"C:\Program Files\Google\Chrome\Application\chrome.exe"
)

# Порт и профиль Chrome — свои на каждую компанию (см. подробное объяснение
# в scripts/podacha_iska_v2.py): иначе параллельный запуск двух компаний
# подключался бы ко ОДНОМУ и тому же браузеру и путал бы их сессии в СК.
# Совпадает по формуле с podacha_iska_v2.py — если для этой компании уже
# открыт залогиненный Chrome-профиль (после запуска "Подача иска в СК"),
# этот скрипт переиспользует ту же сессию.
_company_id_for_chrome = args.company_id.strip() or "1"
if _company_id_for_chrome.isdigit():
    _chrome_port_offset = int(_company_id_for_chrome)
else:
    import zlib as _zlib
    _chrome_port_offset = _zlib.crc32(_company_id_for_chrome.encode("utf-8")) % 1000
CHROME_DEBUG_PORT = 9222 + _chrome_port_offset
CHROME_DEBUG_ADDRESS = f"127.0.0.1:{CHROME_DEBUG_PORT}"

CHROME_USER_DATA_DIR = Path(
    rf"C:\Users\User\АИС ОИП1\chrome_sud_profile_company{_company_id_for_chrome}"
)

LOGIN_TIMEOUT = 90

# ============================================================
# ЗАДАЧА: СКАЧИВАНИЕ ОПРЕДЕЛЕНИЙ О МЕДИАЦИИ
# ============================================================

WORKDIR = Path(args.workdir) if args.workdir else Path(".")
OUT_DIR = WORKDIR / "out"
OUT_DIR.mkdir(parents=True, exist_ok=True)

_source_excel = Path(args.excel_file)
if not _source_excel.exists():
    print(f"❌ ОШИБКА: входной Excel не найден: {_source_excel}")
    sys.exit(1)

# Работаем с копией загруженного файла внутри out/ — так результат
# (обновлённые колонки T/U) сразу попадает в скачиваемый архив задачи,
# а не остаётся только в uploads/.
EXCEL_FILE = OUT_DIR / _source_excel.name
shutil.copy2(_source_excel, EXCEL_FILE)

EXCEL_SHEET = None          # None = активный лист
START_ROW = 2
END_ROW = None              # None = до последней строки

COL_FIO = "C"
COL_IIN = "D"
COL_CASE_NUMBER = "K"
COL_STATUS = "T"          # статус обработки
COL_RESULT_FILE = "U"     # имя скачанного файла

# Папка со скачанными документами — внутри out/ конкретной задачи.
OUTPUT_DIR = OUT_DIR / "Определения о медиации"

MY_CASES_URL = BASE_URL + "/form/cases/mycases.xhtml"
TARGET_TEXTS = [
    # -------------------------
    # РУССКИЙ
    # -------------------------
    "определение об утверждении соглашения об урегулировании спора",
    "определение об утверждении мирового соглашения сторон",
    "утверждении соглашения об урегулировании спора (конфликта) в порядке медиации",
    "вынесено определение о прекращении (упрощенное производство)",

    # -------------------------
    # ҚАЗАҚША
    # -------------------------
    "дауды медиация тәртібінде реттеу туралы келісімді бекіту",
    "дауды медиация тәртібімен реттеу туралы келісімді бекіту",
    "дауды (жанжалды) медиация тәртібімен реттеу туралы келісімді бекіту",
    "дауды (дау-шарды) реттеу туралы келісімді бекіту",
    "медиативтік келісімді бекіту",
    "медиация тәртібіндегі келісімді бекіту",
    "татуласу келісімін бекіту туралы",
    "бітімгершілік келісімін бекіту туралы",
]

HTTP_TIMEOUT = 120
VERIFY_TLS = True
PAUSE_BETWEEN_CASES = 0.4
CONTINUE_ON_ROW_ERROR = True

LOG_FILE = (WORKDIR / "sud_download_mediation.log") if args.workdir else Path("sud_download_mediation.log")
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[
        logging.FileHandler(LOG_FILE, encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger("sud-http")


# ============================================================
# АВТОМАТИЧЕСКИЙ ВХОД И ПЕРЕДАЧА COOKIES В HTTP-СЕССИЮ
# ============================================================

@dataclass
class BrowserAuth:
    cookies: list[dict[str, Any]]
    user_agent: str
    current_url: str


def _debug_port_is_open(
    host: str = "127.0.0.1",
    port: int = CHROME_DEBUG_PORT,
) -> bool:
    try:
        with socket.create_connection((host, port), timeout=1):
            return True
    except OSError:
        return False


def _start_debug_chrome() -> None:
    """
    Запускает отдельный Chrome-профиль для Судебного кабинета — на своих
    порту и профиле для текущей компании. Если Chrome уже запущен на этом
    порту (например, тем же profile из podacha_iska_v2.py), ничего не делает.
    """
    if _debug_port_is_open():
        log.info("Chrome remote debugging уже запущен: %s", CHROME_DEBUG_ADDRESS)
        return

    if not CHROME_PATH.exists():
        raise FileNotFoundError(
            f"Chrome не найден: {CHROME_PATH}. "
            "Исправьте CHROME_PATH в настройках."
        )

    CHROME_USER_DATA_DIR.mkdir(parents=True, exist_ok=True)

    command = [
        str(CHROME_PATH),
        f"--remote-debugging-port={CHROME_DEBUG_PORT}",
        f"--user-data-dir={CHROME_USER_DATA_DIR}",
        "--start-maximized",
        "--disable-notifications",
        BASE_URL + "/form/proceedings/services.xhtml",
    ]

    subprocess.Popen(
        command,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )

    deadline = time.time() + 30
    while time.time() < deadline:
        if _debug_port_is_open():
            log.info("Chrome запущен в режиме remote debugging.")
            return
        time.sleep(0.5)

    raise TimeoutError(
        f"Chrome запущен, но порт remote debugging {CHROME_DEBUG_PORT} не открылся."
    )


def _first_visible_element(driver, selectors: list[tuple[str, str]]):
    from selenium.webdriver.common.by import By

    by_map = {
        "css": By.CSS_SELECTOR,
        "id": By.ID,
        "name": By.NAME,
        "xpath": By.XPATH,
    }

    for selector_type, selector in selectors:
        try:
            elements = driver.find_elements(by_map[selector_type], selector)
        except Exception:
            continue

        for element in elements:
            try:
                if element.is_displayed() and element.is_enabled():
                    return element
            except Exception:
                continue

    return None


def _looks_authenticated(driver) -> bool:
    """
    Проверяет реальную авторизацию.

    Важно: страница входа тоже имеет URL /form/proceedings/services.xhtml,
    поэтому по одному URL авторизацию определять нельзя.
    """
    html = str(driver.page_source or "").casefold()

    try:
        password_fields = driver.find_elements(
            "css selector",
            "input[type='password']",
        )
        for field_el in password_fields:
            if field_el.is_displayed():
                return False
    except Exception:
        pass

    login_markers = (
        "жсн/бсн",
        "құпия сөз",
        "кіру",
        "войти",
        "авторизация",
        "type=\"password\"",
    )
    if any(marker in html for marker in login_markers):
        return False

    auth_markers = (
        "шығу",
        "выйти",
        "менің істерім",
        "мои дела",
        "мои заявления",
        "құжаттар беру",
        "подача документов",
    )

    return any(marker in html for marker in auth_markers)


def _perform_login(driver) -> None:
    """
    Вход строго по фактической русской форме портала:

    1. Открывает браузер.
    2. Переключает язык на РУС.
    3. Печатает ИИН/БИН посимвольно.
    4. Печатает пароль посимвольно.
    5. Нажимает точную кнопку Войти.
    """
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.common.by import By
    from selenium.webdriver.common.action_chains import ActionChains
    from selenium.common.exceptions import StaleElementReferenceException

    login_url = BASE_URL + "/form/proceedings/services.xhtml"
    driver.get(login_url)

    WebDriverWait(driver, 30).until(
        lambda d: d.execute_script("return document.readyState") == "complete"
    )

    if _looks_authenticated(driver):
        log.info("Пользователь уже авторизован в Chrome-профиле.")
        return

    if not SUD_LOGIN or not SUD_PASSWORD:
        raise RuntimeError("Не заполнены SUD_LOGIN и SUD_PASSWORD.")

    ru_link = WebDriverWait(driver, 20).until(
        lambda d: _first_visible_element(
            d,
            [
                ("css", "a[onclick*='selectLanguageRu']"),
                ("xpath", "//a[normalize-space()='РУС']"),
            ],
        )
    )

    log.info("Переключаю портал на русский язык.")

    try:
        ActionChains(driver).move_to_element(ru_link).pause(0.3).click().perform()
    except Exception:
        driver.execute_script("arguments[0].click();", ru_link)

    WebDriverWait(driver, 20).until(
        lambda d: bool(d.find_elements(By.CSS_SELECTOR, "input[placeholder='ИИН/БИН']"))
        and bool(d.find_elements(By.CSS_SELECTOR, "input[placeholder='Пароль']"))
    )

    time.sleep(1)

    login_element = WebDriverWait(driver, 20).until(
        lambda d: _first_visible_element(
            d,
            [
                ("css", "input#j_idt71\\:auth\\:xin"),
                ("css", "input[name='j_idt71:auth:xin']"),
                ("css", "input[placeholder='ИИН/БИН']"),
            ],
        )
    )

    password_element = WebDriverWait(driver, 20).until(
        lambda d: _first_visible_element(
            d,
            [
                ("css", "input#j_idt71\\:auth\\:password"),
                ("css", "input[name='j_idt71:auth:password']"),
                ("css", "input[placeholder='Пароль']"),
            ],
        )
    )

    log.info("Печатаю логин посимвольно.")

    login_element.click()

    driver.execute_script(
        """
        const el = arguments[0];
        el.value = '';
        el.dispatchEvent(new Event('input', {bubbles: true}));
        el.dispatchEvent(new Event('change', {bubbles: true}));
        """,
        login_element,
    )

    time.sleep(0.3)

    for char in SUD_LOGIN:
        login_element.send_keys(char)
        time.sleep(0.10)

    time.sleep(0.5)

    password_element = WebDriverWait(driver, 20).until(
        lambda d: _first_visible_element(
            d,
            [
                ("css", "input#j_idt71\\:auth\\:password"),
                ("css", "input[name='j_idt71:auth:password']"),
                ("css", "input[placeholder='Пароль']"),
            ],
        )
    )

    log.info("Печатаю пароль посимвольно.")

    password_element.click()

    driver.execute_script(
        """
        const el = arguments[0];
        el.value = '';
        el.dispatchEvent(new Event('input', {bubbles: true}));
        el.dispatchEvent(new Event('change', {bubbles: true}));
        """,
        password_element,
    )

    time.sleep(0.3)

    for char in SUD_PASSWORD:
        password_element.send_keys(char)
        time.sleep(0.10)

    time.sleep(0.8)

    def find_login_button():
        return _first_visible_element(
            driver,
            [
                ("css", "input#j_idt71\\:auth\\:j_idt78[value='Войти']"),
                ("css", "input[name='j_idt71:auth:j_idt78'][value='Войти']"),
                ("css", "input[type='submit'][value='Войти']"),
            ],
        )

    login_element = WebDriverWait(driver, 20).until(
        lambda d: _first_visible_element(
            d,
            [
                ("css", "input#j_idt71\\:auth\\:xin"),
                ("css", "input[name='j_idt71:auth:xin']"),
                ("css", "input[placeholder='ИИН/БИН']"),
            ],
        )
    )

    actual_login = driver.execute_script("return arguments[0].value;", login_element)

    if str(actual_login or "") != SUD_LOGIN:
        raise RuntimeError(
            f"Логин введён некорректно: {actual_login!r}. Ожидалось: {SUD_LOGIN!r}"
        )

    button = WebDriverWait(driver, 20).until(lambda d: find_login_button())

    log.info("Логин проверен: %s. Нажимаю точную кнопку Войти.", actual_login)

    clicked = False
    for attempt in range(1, 4):
        try:
            button = find_login_button()
            if button is None:
                raise RuntimeError("Кнопка Войти не найдена.")

            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", button)
            time.sleep(0.4)

            ActionChains(driver).move_to_element(button).pause(0.3).click().perform()

            clicked = True
            log.info("Кнопка Войти нажата, попытка %s/3.", attempt)
            break

        except StaleElementReferenceException:
            log.warning("Кнопка Войти перерисовалась. Повтор %s/3.", attempt)
            time.sleep(0.8)

    if not clicked:
        raise RuntimeError("Не удалось нажать кнопку Войти после трёх попыток.")

    try:
        WebDriverWait(driver, LOGIN_TIMEOUT).until(lambda d: _looks_authenticated(d))
    except Exception as exc:
        debug_dir = Path(rf"C:\Users\User\АИС ОИП1\login_debug")
        debug_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        screenshot_path = debug_dir / f"login_error_{timestamp}.png"
        html_path = debug_dir / f"login_error_{timestamp}.html"

        try:
            driver.save_screenshot(str(screenshot_path))
            html_path.write_text(driver.page_source, encoding="utf-8")
            log.error("Диагностика входа сохранена: %s | %s", screenshot_path, html_path)
        except Exception:
            pass

        page_text = str(driver.page_source or "").casefold()

        if "пользователь не найден" in page_text:
            reason = "портал сообщил: пользователь не найден"
        elif "проверьте правильность" in page_text:
            reason = "портал сообщил: проверьте логин или пароль"
        else:
            reason = (
                "кнопка Войти нажата, но личный кабинет не открылся; "
                f"URL={driver.current_url}"
            )

        raise RuntimeError(f"Автоматический вход не выполнен: {reason}.") from exc

    log.info("Автоматический вход выполнен. URL: %s", driver.current_url)


def get_auth_from_chrome_auto(
    debug_address: str = CHROME_DEBUG_ADDRESS,
) -> BrowserAuth:
    """
    Сам запускает Chrome, при необходимости выполняет вход,
    затем передаёт cookies в requests.
    """
    from selenium import webdriver

    _start_debug_chrome()

    options = webdriver.ChromeOptions()
    options.add_experimental_option("debuggerAddress", debug_address)

    driver = webdriver.Chrome(options=options)
    try:
        if AUTO_LOGIN:
            _perform_login(driver)

        auth = BrowserAuth(
            cookies=driver.get_cookies(),
            user_agent=driver.execute_script("return navigator.userAgent"),
            current_url=driver.current_url,
        )

        if not auth.cookies:
            raise RuntimeError("После входа Chrome не вернул cookies авторизации.")

        log.info("Получено cookies из Chrome: %s шт.", len(auth.cookies))
        log.info("Текущая страница Chrome: %s", auth.current_url)
        return auth
    finally:
        # Не закрываем Chrome: профиль остаётся доступным для проверки
        # и для следующего скрипта этой же компании.
        try:
            driver.service.stop()
        except Exception:
            pass
        del driver


def make_http_session(auth: BrowserAuth) -> requests.Session:
    session = requests.Session()
    session.headers.update({
        "User-Agent": auth.user_agent,
        "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
        "Origin": BASE_URL,
    })

    for cookie in auth.cookies:
        kwargs = {
            "name": cookie["name"],
            "value": cookie["value"],
            "path": cookie.get("path", "/"),
        }
        domain = cookie.get("domain")
        if domain:
            kwargs["domain"] = domain
        session.cookies.set(**kwargs)

    return session


# ============================================================
# HTTP/API-ФУНКЦИИ ДЛЯ ПОИСКА ДЕЛА И СКАЧИВАНИЯ DOCX
# ============================================================

VIEWSTATE_NAMES = ("javax.faces.ViewState", "jakarta.faces.ViewState")
DEBUG_DIR = OUTPUT_DIR / "_debug"


def extract_viewstate(html: str) -> str | None:
    soup = BeautifulSoup(html, "lxml")
    for name in VIEWSTATE_NAMES:
        node = soup.find("input", attrs={"name": name})
        if node and node.get("value"):
            return str(node["value"])
    return None


def looks_like_login(response: requests.Response) -> bool:
    text = response.text[:10000].casefold()
    return (
        "войти в судебный кабинет" in text
        or 'type="password"' in text
        or "placeholder=\"пароль\"" in text
    )


def checked_get(url: str, *, referer: str | None = None) -> requests.Response:
    headers = {"Referer": referer} if referer else {}
    response = SESSION.get(
        url,
        headers=headers,
        timeout=HTTP_TIMEOUT,
        verify=VERIFY_TLS,
        allow_redirects=True,
    )
    response.raise_for_status()
    if looks_like_login(response):
        raise RuntimeError("Авторизация потеряна. Повторно выполните блок входа.")
    return response


def checked_post(
    url: str,
    data: list[tuple[str, str]],
    *,
    referer: str,
    ajax: bool = True,
) -> requests.Response:
    headers = {
        "Referer": referer,
        "Origin": BASE_URL,
        "Accept": "*/*",
        "Content-Type": "application/x-www-form-urlencoded;charset=UTF-8",
    }
    if ajax:
        headers["Faces-Request"] = "partial/ajax"

    response = SESSION.post(
        url,
        data=data,
        headers=headers,
        timeout=HTTP_TIMEOUT,
        verify=VERIFY_TLS,
        allow_redirects=True,
    )
    response.raise_for_status()
    if looks_like_login(response):
        raise RuntimeError("Авторизация потеряна. Повторно выполните блок входа.")
    return response


def parse_partial_response(xml_text: str) -> tuple[str, str | None, str | None]:
    """
    Возвращает:
      1) объединённый HTML всех <update>;
      2) URL redirect;
      3) новый ViewState.
    """
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError:
        return xml_text, None, extract_viewstate(xml_text)

    fragments: list[str] = []
    redirect_url = None
    new_viewstate = None

    for node in root.iter():
        tag = node.tag.rsplit("}", 1)[-1]

        if tag == "update":
            node_id = str(node.attrib.get("id") or "")
            content = node.text or ""

            if node_id in VIEWSTATE_NAMES or node_id.endswith("ViewState"):
                new_viewstate = content.strip()
            elif content:
                fragments.append(content)

        elif tag == "redirect":
            redirect_url = node.attrib.get("url")

        elif tag == "eval" and node.text and not redirect_url:
            match = re.search(
                r"(?:window\.)?location(?:\.href)?\s*=\s*['\"]([^'\"]+)['\"]",
                node.text,
                flags=re.I,
            )
            if match:
                redirect_url = match.group(1)

    return "\n".join(fragments), redirect_url, new_viewstate


def save_debug(name: str, content: str) -> Path:
    DEBUG_DIR.mkdir(parents=True, exist_ok=True)
    path = DEBUG_DIR / name
    path.write_text(content, encoding="utf-8", errors="replace")
    return path


def normalize_case_number(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().lstrip("№").strip()


def normalize_iin(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    return digits.zfill(12) if digits else ""


def sanitize_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return value[:180] or "Без имени"


def read_excel_rows() -> list[dict[str, str]]:
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(f"Excel-файл не найден: {EXCEL_FILE}")

    wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    try:
        ws = wb[EXCEL_SHEET] if EXCEL_SHEET else wb.active
        last_row = END_ROW or ws.max_row
        rows = []

        for row in range(START_ROW, last_row + 1):
            case_number = normalize_case_number(ws[f"{COL_CASE_NUMBER}{row}"].value)
            fio = str(ws[f"{COL_FIO}{row}"].value or "").strip()
            iin = normalize_iin(ws[f"{COL_IIN}{row}"].value)

            if not case_number:
                log.info("Строка %s пропущена: колонка %s пустая.", row, COL_CASE_NUMBER)
                continue

            rows.append({
                "row": row,
                "case_number": case_number,
                "fio": fio,
                "iin": iin,
            })
        return rows
    finally:
        wb.close()


def discover_filter_ids(page_html: str) -> dict[str, str]:
    """
    Находит реальные JSF/RichFaces ID прямо в HTML страницы.
    HAR показал шаблон:
      ...:filter-form:filter-requestNumber
      кнопка Найти -> ...:j_idt50
    """
    soup = BeautifulSoup(page_html, "lxml")

    number_input = soup.find(
        "input",
        attrs={"name": re.compile(r"filter-requestNumber$", re.I)}
    )
    if not number_input:
        raise RuntimeError("Не найдено поле filter-requestNumber на странице «Мои дела».")

    field_name = str(number_input.get("name"))
    form = number_input.find_parent("form")
    if not form:
        raise RuntimeError("Поле номера дела не находится внутри JSF-формы.")

    form_name = str(form.get("name") or form.get("id") or field_name.rsplit(":", 1)[0])

    search_button = None
    for node in form.find_all(["input", "button", "a"]):
        caption = " ".join([
            str(node.get("value") or ""),
            node.get_text(" ", strip=True),
            str(node.get("title") or ""),
        ]).casefold()
        if "найти" in caption:
            search_button = node
            break

    if not search_button:
        raise RuntimeError("Не найдена кнопка «Найти».")

    source_id = str(search_button.get("id") or search_button.get("name") or "").strip()
    if not source_id:
        onclick = str(search_button.get("onclick") or "")
        match = re.search(
            r"(?:jsf\.ajax\.request|RichFaces\.ajax)\s*\(\s*['\"]([^'\"]+)['\"]",
            onclick,
            flags=re.I,
        )
        if match:
            source_id = match.group(1)

    if not source_id:
        raise RuntimeError("Не удалось определить JSF ID кнопки «Найти».")

    page_scripts = "\n".join(
        script.get_text(" ", strip=True)
        for script in soup.find_all("script")
    )

    open_match = re.search(
        r'viewSelectedRequest\s*=\s*function\s*\([^)]*\)\s*\{\s*'
        r'RichFaces\.ajax\("([^"]+)"',
        page_scripts,
        flags=re.I | re.S,
    )

    if not open_match:
        open_match = re.search(
            r'viewSelectedRequest.*?RichFaces\.ajax\("([^"]+)"',
            page_scripts,
            flags=re.I | re.S,
        )

    open_source = (
        open_match.group(1).replace("\\u002D", "-")
        if open_match
        else None
    )

    return {
        "form_name": form_name,
        "number_field": field_name,
        "search_source": source_id,
        "open_source": open_source,
    }


def exact_richfaces_payload(
    *,
    form_name: str,
    number_field: str,
    case_number: str,
    viewstate: str,
    source_id: str,
    extra: list[tuple[str, str]] | None = None,
    include_click_event: bool = False,
) -> list[tuple[str, str]]:
    """
    Точный формат POST из перехваченного HAR.
    """
    data = [
        (form_name, form_name),
        (number_field, case_number),
        ("javax.faces.ViewState", viewstate),
        ("javax.faces.source", source_id),
    ]

    if include_click_event:
        data.append(("javax.faces.partial.event", "click"))

    data.extend([
        ("javax.faces.partial.execute", f"{source_id} @component"),
        ("javax.faces.partial.render", "@component"),
    ])

    if extra:
        data.extend(extra)

    data.extend([
        ("org.richfaces.ajax.component", source_id),
        (source_id, source_id),
        ("rfExt", "null"),
        ("AJAX:EVENTS_COUNT", "1"),
        ("javax.faces.partial.ajax", "true"),
    ])
    return data


def submit_search(page: requests.Response, case_number: str) -> dict:
    ids = discover_filter_ids(page.text)
    viewstate = extract_viewstate(page.text)
    if not viewstate:
        raise RuntimeError("На странице не найден javax.faces.ViewState.")

    data = exact_richfaces_payload(
        form_name=ids["form_name"],
        number_field=ids["number_field"],
        case_number=case_number,
        viewstate=viewstate,
        source_id=ids["search_source"],
        include_click_event=True,
    )

    response = checked_post(
        page.url,
        data,
        referer=page.url,
        ajax=True,
    )

    html, redirect, new_viewstate = parse_partial_response(response.text)

    if redirect:
        redirected = checked_get(urljoin(response.url, redirect), referer=response.url)
        return {
            "page_url": redirected.url,
            "html": redirected.text,
            "viewstate": extract_viewstate(redirected.text),
            "ids": ids,
            "raw_response": response.text,
        }

    if not html.strip():
        debug = save_debug("01_search_empty_response.xml", response.text)
        raise RuntimeError(
            f"Поиск вернул пустой partial-response. Диагностика сохранена: {debug}"
        )

    save_debug("01_search_response.xml", response.text)
    save_debug("02_search_updates.html", html)

    return {
        "page_url": response.url,
        "html": html,
        "viewstate": new_viewstate or viewstate,
        "ids": ids,
        "raw_response": response.text,
    }


def discover_case_click(
    result_html: str,
    form_name: str,
    open_source: str | None,
) -> tuple[str, str]:
    """
    В AJAX-результате находится только:
        onclick="viewSelectedRequest(0)"

    Сам JSF-компонент viewSelectedRequest расположен в полной странице
    «Мои дела», поэтому он передаётся из discover_filter_ids().
    """
    soup = BeautifulSoup(result_html, "lxml")

    cards = soup.select(".case-item-container[onclick]")
    if not cards:
        cards = soup.find_all(
            attrs={"onclick": re.compile(r"viewSelectedRequest\s*\(", re.I)}
        )

    for card in cards:
        onclick = str(card.get("onclick") or "")
        match = re.search(
            r"viewSelectedRequest\s*\(\s*(\d+)\s*\)",
            onclick,
            flags=re.I,
        )
        if not match:
            continue

        param1 = match.group(1)

        if open_source:
            return open_source, param1

        return f"{form_name}:j_idt119", param1

    debug = save_debug("02b_unrecognized_case_card.html", result_html)
    raise RuntimeError(
        "Результат поиска получен, но карточка "
        "onclick=viewSelectedRequest(...) не найдена. "
        f"Диагностика: {debug}"
    )


def open_case(search_state: dict, case_number: str) -> requests.Response:
    source_id, param1 = discover_case_click(
        search_state["html"],
        search_state["ids"]["form_name"],
        search_state["ids"].get("open_source"),
    )

    ids = search_state["ids"]
    viewstate = search_state["viewstate"]
    if not viewstate:
        raise RuntimeError("После поиска отсутствует ViewState.")

    data = exact_richfaces_payload(
        form_name=ids["form_name"],
        number_field=ids["number_field"],
        case_number=case_number,
        viewstate=viewstate,
        source_id=source_id,
        extra=[("param1", param1)],
        include_click_event=False,
    )

    response = checked_post(
        search_state["page_url"],
        data,
        referer=search_state["page_url"],
        ajax=True,
    )

    html, redirect, _ = parse_partial_response(response.text)
    save_debug("03_open_case_response.xml", response.text)

    if redirect:
        return checked_get(urljoin(response.url, redirect), referer=response.url)

    if html.strip():
        save_debug("04_open_case_updates.html", html)

        soup = BeautifulSoup(html, "lxml")
        for link in soup.find_all("a", href=True):
            href = str(link.get("href") or "")
            probe = (href + " " + link.get_text(" ", strip=True)).casefold()
            if "case" in probe or "дел" in probe:
                return checked_get(urljoin(response.url, href), referer=response.url)

        if (
            any(text.casefold() in html.casefold() for text in TARGET_TEXTS)
            or "динамика хода рассмотрения дела" in html.casefold()
        ):
            response._content = html.encode(response.encoding or "utf-8", errors="replace")
            return response

    current = checked_get(search_state["page_url"], referer=search_state["page_url"])
    current_text = current.text.casefold()
    if (
        "динамика хода рассмотрения дела" in current_text
        or any(text.casefold() in current_text for text in TARGET_TEXTS)
        or "скачать талон об отправке" in current_text
        or "отправить дополнительные документы" in current_text
    ):
        return current

    debug = save_debug("05_after_open_current_page.html", current.text)
    raise RuntimeError(
        f"POST открытия дела выполнен, но страница дела не распознана. Диагностика: {debug}"
    )


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().casefold()


def _direct_text(node) -> str:
    """
    Только собственный текст элемента, без текста всех вложенных контейнеров.
    Это не даёт принять весь блок страницы за нужный абзац.
    """
    parts = []
    for child in node.children:
        if isinstance(child, str):
            parts.append(child)
        elif getattr(child, "name", None) == "br":
            parts.append(" ")
    return _normalized_text(" ".join(parts))


def find_target_document_link(case_response: requests.Response) -> str:
    """
    Находит именно абзац с TARGET_TEXT и только ближайшую DOCX-ссылку,
    относящуюся к этому абзацу. Общие DIV-контейнеры страницы игнорируются.
    """
    soup = BeautifulSoup(case_response.text, "lxml")
    targets = [_normalized_text(x) for x in TARGET_TEXTS]

    paragraphs = []

    for text_node in soup.find_all(string=True):
        text = _normalized_text(str(text_node))
        if not any(target in text for target in targets):
            continue

        parent = text_node.parent
        if parent and parent.name not in {"script", "style"}:
            paragraphs.append(parent)

    if not paragraphs:
        for node in soup.find_all(["p", "span", "td", "li", "div"]):
            if any(target in _direct_text(node) for target in targets):
                paragraphs.append(node)

    unique_paragraphs = []
    seen = set()
    for node in paragraphs:
        marker = id(node)
        if marker not in seen:
            seen.add(marker)
            unique_paragraphs.append(node)

    if not unique_paragraphs:
        debug = save_debug("06_case_without_target_paragraph.html", case_response.text)
        raise RuntimeError(
            f"Ни один из целевых абзацев не найден: {TARGET_TEXTS}. Диагностика: {debug}"
        )

    for paragraph in unique_paragraphs:
        for link in paragraph.find_all("a", href=True):
            href = str(link.get("href") or "")
            label = link.get_text(" ", strip=True)
            if re.search(r"\.docx?(?:\b|[?#&/])", f"{href} {label}".casefold()):
                return urljoin(case_response.url, href)

        sibling = paragraph.next_sibling
        checked = 0

        while sibling is not None and checked < 8:
            checked += 1

            if isinstance(sibling, str):
                text = _normalized_text(sibling)
                sibling = sibling.next_sibling
                if not text:
                    continue
                if not re.search(r"\.docx?\b", text):
                    break
                continue

            if sibling.name == "a" and sibling.get("href"):
                href = str(sibling.get("href") or "")
                label = sibling.get_text(" ", strip=True)
                if re.search(r"\.docx?(?:\b|[?#&/])", f"{href} {label}".casefold()):
                    return urljoin(case_response.url, href)

            for link in sibling.find_all("a", href=True, recursive=True):
                href = str(link.get("href") or "")
                label = link.get_text(" ", strip=True)
                if re.search(r"\.docx?(?:\b|[?#&/])", f"{href} {label}".casefold()):
                    return urljoin(case_response.url, href)

            sibling_text = _normalized_text(sibling.get_text(" ", strip=True))

            if sibling_text and ".docx" not in sibling_text:
                break

            sibling = sibling.next_sibling

    debug = save_debug(
        "07_exact_paragraph_without_adjacent_docx.html",
        case_response.text,
    )
    raise RuntimeError(
        "Нужный абзац найден, но рядом с ним нет DOC/DOCX-ссылки. "
        "Другие документы не скачивались. "
        f"Диагностика: {debug}"
    )


def extract_docx_text(file_path: Path) -> str:
    """
    Извлекает текст из word/document.xml без открытия файла в Word.
    """
    with zipfile.ZipFile(file_path, "r") as archive:
        xml_bytes = archive.read("word/document.xml")

    root = ET.fromstring(xml_bytes)
    texts = []
    for node in root.iter():
        if node.tag.rsplit("}", 1)[-1] == "t" and node.text:
            texts.append(node.text)

    return _normalized_text(" ".join(texts))


def detect_document_language(text: str) -> str:
    """
    Возвращает KZ или RU. Для казахского проверяются специфические буквы
    казахского алфавита.
    """
    kz_letters = set("әғқңөұүһі")
    kz_count = sum(1 for ch in text.casefold() if ch in kz_letters)
    return "KZ" if kz_count >= 3 else "RU"


def validate_mediation_text(text: str) -> str:
    """
    Двуязычная проверка судебного акта.

    Принимаем:
      - русское определение о медиации / мировом соглашении;
      - казахское ҰЙҒАРЫМ о медиации;
      - казахское ҰЙҒАРЫМ об утверждении татуласу/бітімгершілік келісімі.

    Отсекаем:
      - определение/ұйғарым только о принятии и возбуждении дела,
        даже если внутри общим текстом упоминается возможность медиации.
    """
    text = _normalized_text(text)

    negative_markers = (
        "қарауға қабылдау және азаматтық іс қозғау",
        "азаматтық іс қозғалсын",
        "оңайлатылған (жазбаша) іс жүргізу тәртібімен қабылданып",
        "принять исковое заявление к производству",
        "принято к производству",
        "возбудить гражданское дело",
        "о принятии искового заявления",
    )

    if any(x in text for x in negative_markers):
        raise RuntimeError(
            "Скачанный Word-файл является определением о принятии/возбуждении дела, "
            "а не определением об утверждении медиации."
        )

    has_ru_order = any(x in text for x in (
        "определен",
        "суд определил",
        "определил:",
    ))

    has_kz_order = any(x in text for x in (
        "ұйғарым",
        "ұйғарды",
        "ұйғарым етті",
    ))

    has_ru_mediation = any(x in text for x in (
        "медиаци",
        "урегулировании спора",
        "урегулирования спора",
        "медиативного соглашения",
        "мирового соглашения",
        "утвердить соглашение",
        "утверждении соглашения",
        "соглашение утвердить",
    ))

    has_kz_mediation = any(x in text for x in (
        "медиация тәртібінде",
        "медиация тәртібімен",
        "медиативтік келісім",
        "дауды (жанжалды) медиация",
        "дауды (дау-шарды) реттеу туралы келісім",
        "медиация туралы",
    ))

    has_kz_settlement = any(x in text for x in (
        "татуласу келісім",
        "бітімгершілік келісім",
    ))

    has_approval = any(x in text for x in (
        "утвердить",
        "утверждении",
        "утвержден",
        "утвердил",
        "бекітілсін",
        "бекітілуге",
        "бекіту туралы",
        "бекітеді",
        "бекітілген",
    ))

    ru_ok = has_ru_order and has_ru_mediation and has_approval
    kz_ok = has_kz_order and (has_kz_mediation or has_kz_settlement) and has_approval

    if ru_ok or kz_ok:
        return detect_document_language(text)

    preview = text[:1000]
    raise RuntimeError(
        "Скачанный Word-файл не похож на русское или казахское определение "
        "об утверждении медиации/соглашения. "
        f"RU_ORDER={has_ru_order}, KZ_ORDER={has_kz_order}, "
        f"RU_MEDIATION={has_ru_mediation}, KZ_MEDIATION={has_kz_mediation}, "
        f"KZ_SETTLEMENT={has_kz_settlement}, APPROVAL={has_approval}. "
        f"Начало документа: {preview!r}"
    )


def validate_downloaded_mediation_docx(file_path: Path) -> str:
    text = extract_docx_text(file_path)
    return validate_mediation_text(text)


def download_word_document(url: str, destination: Path, referer: str | None = None) -> Path:
    headers = {"Referer": referer} if referer else {}
    response = SESSION.get(url, headers=headers, timeout=HTTP_TIMEOUT)
    response.raise_for_status()
    content = response.content
    if not content:
        raise RuntimeError("Сервер вернул пустой файл.")
    disp = response.headers.get("Content-Disposition", "").casefold()
    ctype = response.headers.get("Content-Type", "").casefold()
    probe = (response.url + " " + disp).casefold()
    ole = bytes.fromhex("D0CF11E0A1B11AE1")
    if content.startswith(b"PK") or ".docx" in probe or "openxmlformats-officedocument.wordprocessingml.document" in ctype:
        ext = ".docx"
    elif content.startswith(ole) or re.search(r"\.doc(?:\b|[?#&/])", probe) or "application/msword" in ctype:
        ext = ".doc"
    else:
        raise RuntimeError(f"Файл не распознан как DOC/DOCX. Content-Type: {ctype}")
    destination = destination.with_suffix(ext)
    destination.write_bytes(content)
    return destination


def validate_docx_structure(file_path: Path) -> None:
    try:
        with zipfile.ZipFile(file_path, "r") as archive:
            bad_member = archive.testzip()
            if bad_member:
                raise RuntimeError(f"Повреждён элемент DOCX: {bad_member}")

            names = set(archive.namelist())
            required = {
                "[Content_Types].xml",
                "_rels/.rels",
                "word/document.xml",
            }
            missing = required - names
            if missing:
                raise RuntimeError(
                    "В DOCX отсутствуют обязательные элементы: "
                    + ", ".join(sorted(missing))
                )
    except zipfile.BadZipFile as exc:
        raise RuntimeError("Скачанный файл не является корректным DOCX.") from exc


def rebuild_docx_with_python_docx(file_path: Path) -> None:
    """
    Пересобирает DOCX через python-docx. Обычно после этого Word открывает
    файл без предупреждения о повреждённом содержимом.
    """
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Не установлен python-docx. Выполните: pip install python-docx"
        ) from exc

    temp_fixed = file_path.with_name(file_path.stem + ".__fixed__.docx")

    try:
        document = Document(str(file_path))
        document.save(str(temp_fixed))
        validate_docx_structure(temp_fixed)
        temp_fixed.replace(file_path)
    except Exception as exc:
        if temp_fixed.exists():
            try:
                temp_fixed.unlink()
            except OSError:
                pass
        raise RuntimeError(
            "DOCX скачан, но автоматически пересобрать его не удалось."
        ) from exc


def prepare_result_excel() -> None:
    """Добавляет заголовки T/U в тот же Excel."""
    wb = load_workbook(EXCEL_FILE)
    try:
        ws = wb[EXCEL_SHEET] if EXCEL_SHEET else wb.active
        ws[f"{COL_STATUS}1"] = "Статус скачивания"
        ws[f"{COL_RESULT_FILE}1"] = "Имя файла"
        wb.save(EXCEL_FILE)
    finally:
        wb.close()


def write_excel_result(row: int, status: str, filename: str = "") -> None:
    """
    Записывает результат сразу после обработки строки:
      T = статус
      U = имя файла
    """
    wb = load_workbook(EXCEL_FILE)
    try:
        ws = wb[EXCEL_SHEET] if EXCEL_SHEET else wb.active
        ws[f"{COL_STATUS}{row}"] = status
        ws[f"{COL_RESULT_FILE}{row}"] = filename
        wb.save(EXCEL_FILE)
    finally:
        wb.close()


def friendly_status(exc: Exception) -> str:
    text = re.sub(r"\s+", " ", str(exc)).strip()

    if "Точный абзац" in text or "Абзац с текстом" in text:
        return "Нужный абзац не найден"
    if "рядом с ним нет DOCX" in text or "непосредственно после него нет DOCX" in text:
        return "DOCX после нужного абзаца не найден"
    if "не похож на определение о медиации" in text:
        return "Неверный документ — не определение"
    if "не найдено в результатах" in text:
        return "Дело не найдено"
    if "Авторизация потеряна" in text:
        return "Ошибка авторизации"
    if "не является корректным DOCX" in text or "Повреждён элемент DOCX" in text:
        return "Некорректный DOCX"
    if "пересобрать его не удалось" in text:
        return "Ошибка исправления DOCX"

    return ("Ошибка: " + text)[:250]


def validate_word_structure(file_path: Path) -> None:
    if file_path.suffix.lower() == ".docx":
        return validate_docx_structure(file_path)
    if file_path.suffix.lower() == ".doc":
        if file_path.read_bytes()[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
            raise RuntimeError("Скачанный файл не является корректным DOC.")
        return
    raise RuntimeError("Неизвестный формат Word-файла.")


def validate_downloaded_mediation_document(file_path: Path) -> str:
    """
    Проверяет DOCX и старый DOC на русском и казахском языках.
    Возвращает язык: RU или KZ.
    """
    if file_path.suffix.lower() == ".docx":
        return validate_downloaded_mediation_docx(file_path)

    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "Для проверки .doc установите pywin32: py -m pip install pywin32"
        ) from exc

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None

    try:
        doc = word.Documents.Open(
            str(file_path.resolve()),
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        text = _normalized_text(doc.Content.Text)
    finally:
        if doc is not None:
            doc.Close(SaveChanges=False)
        word.Quit()

    return validate_mediation_text(text)


def rebuild_word_document(file_path: Path) -> None:
    if file_path.suffix.lower() == ".docx":
        rebuild_docx_with_python_docx(file_path)


def unique_destination(fio: str, iin: str) -> Path:
    base = sanitize_filename(f"Определение о медиации по {fio}, {iin}")
    path = OUTPUT_DIR / f"{base}.docx"

    counter = 2
    while path.exists():
        path = OUTPUT_DIR / f"{base} ({counter}).docx"
        counter += 1

    return path


# ============================================================
# ЗАПУСК ОБРАБОТКИ ВСЕГО EXCEL
# ============================================================

def main():
    global SESSION

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    auth = get_auth_from_chrome_auto()
    SESSION = make_http_session(auth)
    print("Автоматический вход выполнен. HTTP session готова.")

    prepare_result_excel()

    rows = read_excel_rows()
    if not rows:
        raise RuntimeError("В Excel не найдено ни одной строки с номером судебного дела.")

    log.info("Найдено строк для обработки: %s", len(rows))
    log.info("Папка результата: %s", OUTPUT_DIR)
    log.info("Статусы: колонка T; имя файла: колонка U")

    completed: list[dict] = []
    failed: list[dict] = []

    for position, item in enumerate(rows, start=1):
        row = item["row"]
        case_number = item["case_number"]
        fio = item["fio"]
        iin = item["iin"]

        log.info(
            "=" * 80 + "\n%s/%s | строка %s | дело %s | %s | %s",
            position, len(rows), row, case_number, fio, iin,
        )

        destination = None

        try:
            write_excel_result(row, "В обработке", "")

            my_cases = checked_get(MY_CASES_URL)
            search_state = submit_search(my_cases, case_number)
            case_page = open_case(search_state, case_number)

            docx_url = find_target_document_link(case_page)
            destination = unique_destination(fio, iin)

            destination = download_word_document(docx_url, destination, referer=case_page.url)
            validate_word_structure(destination)

            document_language = validate_downloaded_mediation_document(destination)

            rebuild_word_document(destination)
            validate_word_structure(destination)
            document_language = validate_downloaded_mediation_document(destination)

            completed.append({
                "row": row,
                "case_number": case_number,
                "file": str(destination),
            })

            write_excel_result(row, f"Скачано ({document_language})", destination.name)
            log.info("СКАЧАНО И ПРОВЕРЕНО [%s]: %s", document_language, destination)

        except Exception as exc:
            if destination and destination.exists():
                try:
                    destination.unlink()
                    log.warning("Некорректный файл удалён: %s", destination)
                except OSError:
                    pass

            status = friendly_status(exc)

            try:
                write_excel_result(row, status, "")
            except Exception as excel_exc:
                log.exception("Не удалось записать статус в Excel: %s", excel_exc)

            failed.append({
                "row": row,
                "case_number": case_number,
                "error": str(exc),
            })
            log.exception("ОШИБКА: строка %s, дело %s: %s", row, case_number, exc)

            if not CONTINUE_ON_ROW_ERROR:
                raise

        time.sleep(PAUSE_BETWEEN_CASES)

    print()
    print("=" * 80)
    print(f"Готово. Скачано файлов: {len(completed)}")
    print(f"Ошибок: {len(failed)}")
    print(f"Папка: {OUTPUT_DIR}")
    print(f"Excel: {EXCEL_FILE}")
    print("T = Статус скачивания")
    print("U = Имя файла")

    if failed:
        print("\nОшибки:")
        for item in failed:
            print(f"  строка {item['row']} | дело {item['case_number']} | {item['error']}")


SESSION: requests.Session = None  # type: ignore[assignment]

if __name__ == "__main__":
    try:
        main()
    except Exception:
        import traceback
        print("КРИТИЧЕСКАЯ ОШИБКА В СКРИПТЕ СКАЧИВАНИЯ ОПРЕДЕЛЕНИЙ О МЕДИАЦИИ")
        traceback.print_exc()
        raise
