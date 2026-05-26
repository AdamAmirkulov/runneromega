# -*- coding: utf-8 -*-
import os
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook
from docx import Document
from docx2pdf import convert   # pip install docx2pdf

# ---------------- НАСТРОЙКИ ----------------
EXCEL_PATH   = r"C:\Users\user\Desktop\Подача по АИС ОИП\Список для подачи заявления ЧСИ через АИС ОИП_Omega.xlsx"
TEMPLATE_DOC = r"C:\Users\user\Desktop\Подача по АИС ОИП\Шаблон заявления для ЧСИ.docx"
OUT_FOLDER   = r"C:\Users\user\Desktop\Подача по АИС ОИП\Заявления"

# -------------------------------------------------
# Вспомогательные функции
# -------------------------------------------------
def nice_case(s: str) -> str:
    """Первая буква каждого слова заглавная, остальные строчные."""
    if not s:
        return ""
    parts = str(s).split()
    return " ".join(p.capitalize() for p in parts)

def replace_placeholders_in_paragraph(paragraph, mapping: dict):
    """
    ВАЖНО: работаем с целым paragraph.text,
    а потом пересобираем абзац, чтобы плейсхолдеры,
    разрезанные на несколько run'ов, тоже заменялись.
    """
    text = paragraph.text
    changed = False

    for key, value in mapping.items():
        if key in text:
            text = text.replace(key, value)
            changed = True

    if changed:
        if paragraph.runs:
            paragraph.runs[0].text = text
            for r in paragraph.runs[1:]:
                r.text = ""
        else:
            paragraph.add_run(text)

def fill_document(row_dict: dict) -> Document:
    """Подставляет данные из строки Excel в шаблон DOCX и возвращает готовый Document."""
    doc = Document(TEMPLATE_DOC)

    mapping = {
        "{Уникальный номер}": str(row_dict.get("Уникальный номер", "")),
        "{ФИО}": nice_case(row_dict.get("ФИО", "")),
        "{ИИН}": str(row_dict.get("ИИН", "")),
        "{Остаток задолженности}": str(row_dict.get("Остаток задолженности", "")),
        "{ЧСИ}": nice_case(row_dict.get("ЧСИ", "")),
        "{ProjectName}": str(row_dict.get("Продукт", "")),
        "{CurrDate}": datetime.now().strftime("%d.%m.%Y"),
    }

    # Абзацы обычные
    for p in doc.paragraphs:
        replace_placeholders_in_paragraph(p, mapping)

    # Абзацы внутри таблиц (на всякий случай)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_placeholders_in_paragraph(p, mapping)

    return doc

# -------------------------------------------------
# Основная логика
# -------------------------------------------------
def run():
    os.makedirs(OUT_FOLDER, exist_ok=True)

    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active

    # Заголовки колонок (первая строка)
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue

        row_dict = dict(zip(headers, row))

        fio = nice_case(row_dict.get("ФИО", ""))
        iin = str(row_dict.get("ИИН", "")).strip()

        if not fio or not iin:
            # если нет ФИО или ИИН – пропускаем
            continue

        # 1) Собираем docx
        doc = fill_document(row_dict)

        base_name = f"Заявление ЧСИ на взыскание, {fio}, {iin}"
        docx_path = Path(OUT_FOLDER) / (base_name + ".docx")
        pdf_path  = Path(OUT_FOLDER) / (base_name + ".pdf")

        doc.save(docx_path)

        # 2) Конвертируем в PDF
        convert(str(docx_path), str(pdf_path))

        # 3) Можно удалить промежуточный DOCX, если не нужен
        try:
            os.remove(docx_path)
        except OSError:
            pass

        print(f"Сформирован файл: {pdf_path}")

    wb.close()
    print("Готово.")

