# -*- coding: utf-8 -*-
"""
portal-sot.kz — заявления на выдачу исполнительного листа (ИЛ) с автоподписью
ЭЦП и отправкой. Два вида (--kind):
  opredelenie — по определению суда (медиативное соглашение не исполнено),
                ноутбук PORTAL_SOT_Podacha_Zayavleniya_Po_Opredeleniyu_na_IL.ipynb;
  reshenie    — по решению суда,
                ноутбук PORTAL_SOT_Podacha_Zayavleniya_Po_Resheniy_na_IL.ipynb.
Ноутбуки отличаются только отбором сделок, текстом/шаблоном заявления и
колонкой даты в ProcessImport — всё это собрано в KINDS ниже.

Цепочка по каждой сделке (как в ноутбуках):
    DOCX из шаблона → PDF (Word) → POST LETTER/create → saveAttachment (PDF)
    → POST LETTER/blank (суд + текст) → подпись xmlForSign в NCALayer
    → POST LETTER/sign → строка в ProcessImport.xlsx для Дельты.

Отличия от ноутбуков:
  * параметры --workdir / --company_id / --max_rows / --no_sign (веб-интерфейс);
  * вход и все запросы — через общий PortalSotHttp + login_chain (сохранённые
    токены, темп, общий лок, стоп при блокировке) вместо Chrome на порту 9222;
  * пароль ЭЦП — config.PORTAL_SOT_BY_COMPANY, а не DPAPI credentials.bin;
    пароль БД — config.CRM_DB, а не из старого ноутбука;
  * истец/директор — config.IL_DECISION_CONFIG, фильтр компании в SQL —
    config.DB_COMPANY_FILTER (l.F209), а не '%Омега%';
  * шаблон .docx один на все компании (в нём всё на переменных), а картинка
    «подпись + печать» подменяется на signature_image компании; если её нет —
    заявление уходит без картинки (чужую печать не ставим);
  * шаблоны (.docx, ProcessImport) — scripts/templates/, справочник судов —
    общий с подачей исков, запасной — scripts/data/;
  * журнал и ProcessImport_registry.xlsx (защита от повторной подачи) лежат
    на диске компании: <ROOT>/Заявления на выдачу ИЛ/portal-sot по …;
    DOCX/PDF и ProcessImport текущего запуска — в <workdir>/out/;
  * ProcessImport.xlsx публикуется в папку автоимпорта (path_crm), только если
    в этом запуске что-то отправлено (ноутбуки публиковали и пустой файл);
  * если исход LETTER/sign неизвестен (обрыв сети/5xx), статус сверяется по
    GET LETTER/<id>; не удалось сверить — SENT_PENDING_VERIFY (повторно не подаём).
"""
import argparse
import os
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass


def _parse_args():
    p = argparse.ArgumentParser(description="portal-sot.kz — заявления на выдачу ИЛ")
    p.add_argument("--kind", default="opredelenie", choices=("opredelenie", "reshenie"),
                   help="вид заявления: по определению / по решению суда")
    p.add_argument("--workdir", default=None, help="рабочая папка задачи (из runner)")
    p.add_argument("--company_id", default=None, help="ID компании (config)")
    p.add_argument("--max_rows", default=None, help="ограничить число сделок (тест)")
    p.add_argument("--no_sign", default=None, help="1 = только черновики, без подписи и отправки")
    return p.parse_args()


_args = _parse_args()

if not _args.company_id or not _args.company_id.strip():
    print("❌ ОШИБКА: не передан --company_id — компания не определена.")
    sys.exit(1)

# COMPANY_ID нужен config.py ещё на этапе импорта
COMPANY_ID = _args.company_id.strip()
os.environ["COMPANY_ID"] = COMPANY_ID

_WORKDIR = Path(_args.workdir).resolve() if (_args.workdir and _args.workdir.strip()) else Path.cwd()
_WORKDIR.mkdir(parents=True, exist_ok=True)

import base64  # noqa: E402
import json  # noqa: E402
import mimetypes  # noqa: E402
import re  # noqa: E402
import shutil  # noqa: E402
import ssl  # noqa: E402
import subprocess  # noqa: E402
import threading  # noqa: E402
import time  # noqa: E402
import traceback  # noqa: E402
import zipfile  # noqa: E402
from datetime import datetime  # noqa: E402

import pandas as pd  # noqa: E402
import pyodbc  # noqa: E402
import requests  # noqa: E402
from lxml import etree  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
import portal_sot_auth  # noqa: E402
from portal_sot_http import PortalSotHttp, PortalBlocked  # noqa: E402
from portal_sot_login import login_chain  # noqa: E402
from config import (  # noqa: E402
    COMPANY_ROOTS, CREDENTIALS, CRM_DB, DB_COMPANY_FILTER, PORTAL_SOT_BY_COMPANY, ROOT,
)

try:
    from config import IL_DECISION_CONFIG  # noqa: E402
except ImportError:
    print("❌ В scripts/config.py нет IL_DECISION_CONFIG — добавьте по образцу config.example.py")
    sys.exit(1)

# ============================================================
# НАСТРОЙКИ
# ============================================================
# Чем различаются два вида заявления (всё остальное общее).
KINDS = {
    "opredelenie": {
        "title": "по определению",
        "template": "OMEGA_IL_DEFINITION_TEMPLATE.docx",
        "template_key": "template_definition",   # свой шаблон компании в IL_DECISION_CONFIG
        "pi_template": "ProcessImport_template_OPREDELENIE.xlsx",
        "pi_date_header": "Дата подачи заявления на выдачу СП",
        "pi_current": "ProcessImport_CURRENT_OPREDELENIE",
        "state_dir": "portal-sot по определению",
        "journal": "SK_IL_OPREDELENIE_PORTAL_SOT_results.xlsx",
        "pause": 3.0,
    },
    "reshenie": {
        "title": "по решению",
        "template": "OMEGA_IL_DECISION_TEMPLATE.docx",
        "template_key": "template",
        "pi_template": "ProcessImport_template_RESHENIE.xlsx",
        "pi_date_header": "Дата подачи заявления на выписку ИЛ",
        "pi_current": "ProcessImport_CURRENT_RESHENIE",
        "state_dir": "portal-sot по решению",
        "journal": "SK_IL_RESHENIE_PORTAL_SOT_results.xlsx",
        "pause": 1.0,
    },
}
KIND = _args.kind
KCFG = KINDS[KIND]

# Пока подаём только по Омеге. Остальные компании настроены (IL_DECISION_CONFIG,
# печати, фильтр кредитора), но выключены — чтобы включить, добавьте ID сюда.
ENABLED_COMPANIES = {"1"}
if COMPANY_ID not in ENABLED_COMPANIES:
    print(f"❌ Заявления на ИЛ через portal-sot.kz пока включены только для Омеги "
          f"(компания {COMPANY_ID} выключена в ENABLED_COMPANIES, scripts/portal_sot_il.py).")
    sys.exit(1)

# Orion и InvestWay в CRM идут под одним F209 («ORION») — различаются кредитором
# (l.F234 → Constants.Caption): Moneyman — это Orion, Vivus — InvestWay.
# Для остальных компаний фильтра по кредитору нет.
CREDITOR_LIKE_BY_COMPANY = {
    "3": "%Moneyman%",
    "4": "%Vivus%",
}
CREDITOR_LIKE = CREDITOR_LIKE_BY_COMPANY.get(COMPANY_ID)

BASE_URL = "https://portal-sot.kz"
ORG_BIN = CREDENTIALS["org_bin"]

_SCRIPT_DIR = Path(__file__).resolve().parent

_ILCFG = IL_DECISION_CONFIG.get(COMPANY_ID)
if not _ILCFG or not _ILCFG.get("project_name") or not _ILCFG.get("director"):
    print(f"❌ Нет записи для компании {COMPANY_ID} в IL_DECISION_CONFIG (scripts/config.py): "
          f"нужны project_name и director.")
    sys.exit(1)

# Шаблон заявления — один на все компании: текст целиком на переменных.
# Не на переменных только картинка «подпись директора + печать» — в общем
# шаблоне она Омеги, для остальных подменяется на signature_image (PNG в
# scripts/templates/). Свой шаблон целиком — ключ template_definition
# (по определению) / template (по решению).
_DEFAULT_TEMPLATE = KCFG["template"]
_TEMPLATE_NAME = _ILCFG.get(KCFG["template_key"]) or _DEFAULT_TEMPLATE
_OWN_TEMPLATE = _TEMPLATE_NAME != _DEFAULT_TEMPLATE

PROJECT_CFG = {"project_name": _ILCFG["project_name"], "director": _ILCFG["director"]}
IL_TEMPLATE = _SCRIPT_DIR / "templates" / _TEMPLATE_NAME
TEMPLATE_STAMP_PART = "word/media/image1.png"   # подпись + печать в общем шаблоне

# Прозрачный PNG 4x4 — «без картинки», когда своей печати у компании нет.
_BLANK_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAQAAAAECAYAAACp8Z5+AAAAFElEQVR4nGP8//8/AzJgQuERJQAAxo4DBV7n/JQAAAAASUVORK5CYII="
)

if _ILCFG.get("signature_image"):
    _sig_path = _SCRIPT_DIR / "templates" / _ILCFG["signature_image"]
    if not _sig_path.exists():
        print(f"❌ Не найдена картинка подписи/печати компании {COMPANY_ID}: {_sig_path}")
        sys.exit(1)
    STAMP_IMAGE = _sig_path.read_bytes()
    STAMP_NOTE = f"подпись/печать: {_sig_path.name}"
elif COMPANY_ID == "1" or _OWN_TEMPLATE:
    STAMP_IMAGE = None   # картинка шаблона — этой компании, оставляем как есть
    STAMP_NOTE = "подпись/печать: из шаблона"
else:
    STAMP_IMAGE = _BLANK_PNG
    STAMP_NOTE = (f"⚠ БЕЗ картинки подписи/печати — в IL_DECISION_CONFIG['{COMPANY_ID}'] "
                  f"не задан signature_image (в общем шаблоне печать Омеги, её не ставим)")
PROCESS_IMPORT_TEMPLATE = _SCRIPT_DIR / "templates" / KCFG["pi_template"]

PORTAL_CFG = PORTAL_SOT_BY_COMPANY.get(COMPANY_ID)
if not PORTAL_CFG or not PORTAL_CFG.get("eds_password"):
    print(f"❌ Для компании {COMPANY_ID} не настроен вход по ЭЦП — PORTAL_SOT_BY_COMPANY в scripts/config.py")
    sys.exit(1)
NCALAYER_PATH = PORTAL_CFG.get("ncalayer_path") or os.path.expandvars(
    r"%LOCALAPPDATA%\Programs\NCALayer\NCALayer.exe"
)

# Текущий запуск — в папку задачи (видно в веб-интерфейсе)
OUT_DIR = _WORKDIR / "out"
DOCX_DIR = OUT_DIR / "DOCX"
PDF_DIR = OUT_DIR / "PDF"

# Постоянное состояние — на диске компании: общее для сервера и локального запуска
_STATE_REL = Path("Заявления на выдачу ИЛ") / KCFG["state_dir"]
STATE_DIR = Path(ROOT) / _STATE_REL
RESULT_XLSX = STATE_DIR / KCFG["journal"]
PROCESS_IMPORT_REGISTRY = STATE_DIR / "ProcessImport_registry.xlsx"   # постоянная история EID
PROCESS_IMPORT_OUTPUT = None  # уникальный файл текущего запуска назначается при старте
PROCESS_IMPORT_NETWORK = Path(CREDENTIALS["path_crm"]) / "ProcessImport.xlsx"

COURTS_FILE_NAME = "Справочник судов portal-sot.xlsx"
COURTS_REL = Path("Документы для подачи Исков") / "Шаблоны документов" / COURTS_FILE_NAME


def _int_arg(v):
    try:
        return int(str(v).strip()) if v and str(v).strip() else None
    except ValueError:
        return None


MAX_ROWS = _int_arg(_args.max_rows)
SIGN_AND_SEND = str(_args.no_sign or "").strip().lower() not in ("1", "да", "yes", "true")
SKIP_ALREADY_CREATED = True
PAUSE_BETWEEN_ROWS = KCFG["pause"]

PROCESS_IMPORT_DATE_HEADER = KCFG["pi_date_header"]
PROCESS_IMPORT_REQUIRED_HEADERS = [
    "Уникальный номер сделки", PROCESS_IMPORT_DATE_HEADER,
    "Номер талона письма ЕПСО", "Тип процесса", "Статус процесса",
]
# С этими статусами сделка считается поданной и повторно не подаётся.
DONE_STATUSES = ["SENT", "SENT_PENDING_VERIFY", "SENT_PROCESSIMPORT_PENDING"]
PROCESS_TYPE_DEFAULT = "1. Упрощённое производство"
PROCESS_STATUS_DEFAULT = "Рассмотрение дела"

RESULT_COLUMNS = ["LoanID", "EID", "IIN", "FIO", "CourtCRM", "CourtCode", "CourtPortal",
                  "DeclarationID", "RequestUID", "PDF", "Status", "Error"]


def log(message=""):
    print(f"[{datetime.now():%d.%m.%Y %H:%M:%S}] {message}", flush=True)


# ============================================================
# SQL — ОТБОР СДЕЛОК ДЛЯ ЗАЯВЛЕНИЯ НА ВЫДАЧУ ИЛ
# (условия отбора из ноутбуков не изменены; фильтр компании — DB_COMPANY_FILTER)
#
# ПО ОПРЕДЕЛЕНИЮ:
# - F220 = 1
# - исключаем закрывающие статусы
# - F236 не "На исполнении"
# - F235 не "Погасил"
# - F329 и F326 последнего процесса = NULL
# - F323 = "Судебное примерение"
# В выборку также тянем: F157 = суд, F35 = сумма без госпошлины
# ============================================================
_SQL_OPREDELENIE = r"""
SELECT
    l.ID                                  AS LoanID,
    l.EID                                 AS EID,
    l.F209                                AS Project,
    c.FIO                                 AS FIO,
    c.F293                                AS IIN,
    LTRIM(RTRIM(l.F157))                  AS Sud,
    CAST(l.F35 AS DECIMAL(18,2))          AS DebtRestWithoutStateDuty,
    s.Caption                             AS CreditStatus,
    constF323.Caption                     AS F323,
    constF236.Caption                     AS F236,
    constF235.Caption                     AS F235,
    pLast.F329                            AS ProcessF329,
    pLast.F326                            AS ProcessF326
FROM dbo.Loans l WITH (NOLOCK)

JOIN dbo.Clients c WITH (NOLOCK)
    ON c.ID = l.CID

JOIN dbo.States s WITH (NOLOCK)
    ON s.ID = l.State

LEFT JOIN dbo.Constants constF236 WITH (NOLOCK)
    ON constF236.ID = l.F236

LEFT JOIN dbo.Constants constF235 WITH (NOLOCK)
    ON constF235.ID = l.F235

LEFT JOIN dbo.Constants constF323 WITH (NOLOCK)
    ON constF323.ID = l.F323

OUTER APPLY (
    SELECT TOP (1)
        p.F329,
        p.F326
    FROM dbo.ProcessCreatedInLoans pl WITH (NOLOCK)
    JOIN dbo.Process p WITH (NOLOCK)
        ON p.ID = pl.ProcessID
    WHERE pl.LoanID = l.ID
    ORDER BY p.ID DESC
) pLast

WHERE
    l.F209 = N'{F209}'
    AND l.F220 = 1

    AND s.Caption NOT IN (
        N'Погашен',
        N'Обратный выкуп',
        N'Умерший',
        N'Банкрот',
        N'Мошенничество',
        N'Армия',
        N'В графике'
    )

    AND (
        constF236.ID IS NULL
        OR NULLIF(LTRIM(RTRIM(constF236.Caption)), N'') <> N'На исполнении'
    )

    AND (
        constF235.ID IS NULL
        OR constF235.Caption NOT IN (N'Погасил')
    )

    AND pLast.F329 IS NULL
    AND pLast.F326 IS NULL

    AND constF323.Caption = N'Судебное примерение'
    {CREDITOR}

    AND NULLIF(LTRIM(RTRIM(l.F157)), N'') IS NOT NULL
    AND c.F293 IS NOT NULL

ORDER BY l.ID ASC;
"""
# ПО РЕШЕНИЮ: решение вынесено (F186 первого процесса) более 33 дней назад,
# заявление на ИЛ ещё не подавалось (F318 пуст), F220 = 0/NULL.
_SQL_RESHENIE = r"""
WITH LoanCounts AS (
    SELECT
        c.F293 AS ИИН,
        COUNT(*) AS LoanCount
    FROM loans l WITH (NOLOCK)
    JOIN clients c WITH (NOLOCK) ON l.CID = c.ID
    WHERE l.F209 = N'{F209}'
    GROUP BY c.F293
)
SELECT
    l.ID                                  AS LoanID,
    l.F209                                AS Project,
    l.EID                                 AS EID,
    dF246.F249                            AS Product,
    c.FIO                                 AS FIO,
    c.F293                                AS IIN,
    s.Caption                             AS CreditStatus,
    NULLIF(constF236.Caption, N'')        AS F236,
    psel.F186                             AS DecisionDate,
    constF238.Caption                     AS Region,
    LTRIM(RTRIM(l.F157))                  AS Sud,
    l.F65                                 AS CaseNumber,
    l.F49                                 AS Judge,
    ISNULL(lc.LoanCount, 0)               AS LoanCount
FROM loans l WITH (NOLOCK)

OUTER APPLY (
    SELECT TOP (1)
        p.F186,
        p.F318
    FROM ProcessCreatedInLoans pcl WITH (NOLOCK)
    JOIN dbo.Process p WITH (NOLOCK) ON p.Id = pcl.ProcessId
    WHERE pcl.LoanId = l.Id
    ORDER BY p.Id ASC
) psel

JOIN clients c WITH (NOLOCK) ON l.CID = c.ID
JOIN Constants constF238 WITH (NOLOCK) ON constF238.ID = c.F238
JOIN states s WITH (NOLOCK) ON s.ID = l.State
JOIN Dictionary dF246 WITH (NOLOCK) ON dF246.ID = l.F246
LEFT JOIN Constants constF236 WITH (NOLOCK) ON constF236.ID = l.F236
LEFT JOIN LoanCounts lc ON c.F293 = lc.ИИН

WHERE
    s.Caption NOT IN (N'Погашен', N'Обратный выкуп', N'В графике', N'Умерший')
    AND l.F209 = N'{F209}'
    AND (constF236.ID IS NULL OR constF236.Caption NOT IN (N'На исполнении'))
    AND psel.F186 IS NOT NULL
    AND (l.F220 = 0 OR l.F220 IS NULL)
    AND psel.F186 < DATEADD(DAY, -33, GETDATE())
    AND psel.F318 IS NULL
    {CREDITOR}
    AND NULLIF(LTRIM(RTRIM(l.F157)), N'') IS NOT NULL
    AND c.F293 IS NOT NULL

ORDER BY psel.F186 ASC;
"""

_SQL_TEMPLATE = _SQL_RESHENIE if KIND == "reshenie" else _SQL_OPREDELENIE
_SQL_CREDITOR = (
    "AND EXISTS (SELECT 1 FROM dbo.Constants constF234 WITH (NOLOCK) "
    f"WHERE constF234.ID = l.F234 AND constF234.Caption LIKE N'{CREDITOR_LIKE}')"
    if CREDITOR_LIKE else ""
)
SQL_QUERY = (_SQL_TEMPLATE
             .replace("{F209}", DB_COMPANY_FILTER.replace("'", "''"))
             .replace("{CREDITOR}", _SQL_CREDITOR))


# ============================================================
# ГЕНЕРАЦИЯ ЗАЯВЛЕНИЯ (DOCX → PDF + текст электронного бланка)
# ============================================================
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def normalize_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def format_money_kz(value):
    if value is None or pd.isna(value):
        return "0,00"
    x = float(value)
    s = f"{x:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", " ")


def _id_str(value):
    """l.ID / l.EID приходят как int, float или Decimal('107020.00') — берём целую часть."""
    m = re.match(r"\s*(\d+)", str(value if value is not None else ""))
    return m.group(1) if m else ""


def defendant_line(row):
    """Строка ответчика — общая для DOCX ({DefendantsData}) и электронного бланка."""
    fio = normalize_text(row["FIO"])
    iin = normalize_text(row["IIN"])
    if KIND == "reshenie":
        return f"1. {fio}, ИИН {iin}"
    # В шаблоне «по определению» номер «1.» — часть списка Word, в бланк он дописывается отдельно.
    return f"{fio}, ИИН {iin}, сумма задолженности {format_money_kz(row['DebtRestWithoutStateDuty'])} тенге"


def build_letter_text_reshenie(row, project_cfg):
    sud = normalize_text(row["Sud"])
    judge = normalize_text(row.get("Judge", ""))
    project_name = project_cfg["project_name"]

    return (
        f"В {sud}\n"
        f"От: Директора {project_name}\n"
        f"{project_cfg['director']}\n"
        f"Судье: {judge}\n\n"
        f"ЗАЯВЛЕНИЕ\n\n"
        f"В производстве судебного органа – {sud} находятся гражданские дела по иску "
        f"{project_name} о взыскании задолженности к следующим ответчикам:\n"
        f"{defendant_line(row)}\n\n"
        f"На данный момент, в отношении указанных должников были вынесены решения суда. "
        f"В связи с этим, просим предоставить исполнительные листы по указанным ответчикам, "
        f"в соответствии с ранее вынесенными решениями суда.\n\n"
        f"ПРОШУ:\n\n"
        f"Выписать исполнительный лист по Решению суда о взыскании задолженности в пользу "
        f"{project_name} и закрепить исполнительный документ за региональной палатой "
        f"частных судебных исполнителей посредством информационной системы «Төрелік», "
        f"по месту регистрации Ответчика.\n\n"
        f"Директор {project_name} {project_cfg['director']}"
    )


def build_letter_text_opredelenie(row, project_cfg):
    amount = format_money_kz(row["DebtRestWithoutStateDuty"])
    sud = normalize_text(row["Sud"])
    fio = normalize_text(row["FIO"])
    iin = normalize_text(row["IIN"])
    project_name = project_cfg["project_name"]

    return (
        f"В {sud}\n"
        f"От: Директора {project_name}\n"
        f"{project_cfg['director']}\n\n"
        f"ЗАЯВЛЕНИЕ\n\n"
        f"В производстве - {sud} находится гражданское дело по иску "
        f"{project_name} о взыскании задолженности к следующему ответчику:\n"
        f"1. {fio}, ИИН {iin}, сумма задолженности {amount} тенге\n\n"
        f"Ранее между {project_name} и вышеуказанным Ответчикам было заключено "
        f"медиативное соглашение, однако до настоящего времени со стороны Ответчика "
        f"не было произведено оплаты в пользу погашения задолженности перед "
        f"{project_name}. Таким образом Ответчик нарушил условия медиативного соглашения.\n\n"
        f"Исполнение данного соглашения возможно только в принудительном порядке, "
        f"поскольку Ответчик по делу исполнять решение добровольно не желает.\n"
        f"На основании вышеизложенного и в соответствии со ст. 178, 180, 241 ГПК РК\n\n"
        f"ПРОШУ:\n\n"
        f"Выписать исполнительный лист по определению суда о взыскании задолженности "
        f"в пользу {project_name} и направить его для исполнения в региональную "
        f"палату частных судебных исполнителей через информационную систему «Төрелік» "
        f"по месту регистрации Ответчика.\n\n"
        f"Директор\n"
        f"{project_name}    {project_cfg['director']}"
    )


def _replace_placeholder_across_text_nodes(nodes, placeholder, replacement):
    """
    Заменяет placeholder даже если Word разорвал его на несколько <w:t>.
    Форматирование первого текстового узла сохраняется.
    """
    if not nodes:
        return False

    full_text = "".join((node.text or "") for node in nodes)
    if placeholder not in full_text:
        return False

    new_text = full_text.replace(placeholder, replacement)

    # Весь новый текст кладём в первый w:t, остальные очищаем.
    nodes[0].text = new_text
    for node in nodes[1:]:
        node.text = ""

    return True


def fill_docx_ooxml(template_path, output_path, replacements, stamp_image=None):
    """
    Меняет переменные непосредственно в OOXML.
    Важно: python-docx здесь НЕ используется, потому что он может потерять
    плавающий графический объект подписи/печати.
    stamp_image — PNG, которым заменяется картинка подписи/печати шаблона
    (рамка и положение остаются от шаблона); None — оставить как в шаблоне.
    """
    with zipfile.ZipFile(template_path, "r") as zin, \
         zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zout:

        for item in zin.infolist():
            data = zin.read(item.filename)

            if item.filename == "word/document.xml":
                root = etree.fromstring(data)

                for p in root.xpath(".//w:p", namespaces=NS):
                    nodes = p.xpath(".//w:t", namespaces=NS)
                    if not nodes:
                        continue

                    for key, value in replacements.items():
                        _replace_placeholder_across_text_nodes(nodes, key, str(value))

                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")

            elif item.filename == TEMPLATE_STAMP_PART and stamp_image is not None:
                data = stamp_image

            zout.writestr(item, data)


def convert_docx_to_pdf(docx_path, pdf_path):
    """
    Сначала Microsoft Word COM (лучше сохраняет верстку).
    Если pywin32/Word недоступны — пробуем LibreOffice.
    """
    docx_path = Path(docx_path).resolve()
    pdf_path = Path(pdf_path).resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    # Вариант 1: MS Word
    try:
        import win32com.client

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        try:
            doc = word.Documents.Open(str(docx_path))
            # 17 = wdFormatPDF
            doc.SaveAs(str(pdf_path), FileFormat=17)
            doc.Close(False)
        finally:
            word.Quit()

        if pdf_path.exists() and pdf_path.stat().st_size > 0:
            return pdf_path

    except Exception as e:
        log(f"Word PDF fallback: {e}")

    # Вариант 2: LibreOffice
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "Не удалось конвертировать DOCX в PDF. "
            "Нужен Microsoft Word + pywin32 или LibreOffice."
        )

    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
        check=True,
        capture_output=True,
    )

    generated = pdf_path.parent / (docx_path.stem + ".pdf")
    if generated != pdf_path and generated.exists():
        generated.replace(pdf_path)

    if not pdf_path.exists():
        raise RuntimeError(f"PDF не создан: {pdf_path}")

    return pdf_path


def make_documents(row):
    """
    DOCX/PDF создаются из шаблона вида заявления с сохранением печати и
    подписи. Электронный бланк строится из того же набора данных.
    """
    fio = str(row["FIO"]).strip()
    iin = str(row["IIN"]).strip()
    safe = re.sub(r'[\\/:*?"<>|]+', "_", fio)

    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)

    docx_path = DOCX_DIR / f"Заявление_на_выдачу_ИЛ, {safe}, {iin}.docx"
    pdf_path = PDF_DIR / f"Заявление_на_выдачу_ИЛ, {safe}, {iin}.pdf"

    # Электронный бланк.
    if KIND == "reshenie":
        letter_text = build_letter_text_reshenie(row, PROJECT_CFG)
    else:
        letter_text = build_letter_text_opredelenie(row, PROJECT_CFG)

    replacements = {
        "{Sud}": normalize_text(row["Sud"]),
        "{ProjectName}": PROJECT_CFG["project_name"],
        "{ProjectDirector}": PROJECT_CFG["director"],
        "{DefendantsData}": defendant_line(row),
    }
    if KIND == "reshenie":
        replacements["{Judge}"] = normalize_text(row.get("Judge", ""))
        replacements["{Pechat}"] = ""

    fill_docx_ooxml(IL_TEMPLATE, docx_path, replacements, stamp_image=STAMP_IMAGE)

    # Контроль: в сформированном document.xml не должны остаться placeholders.
    with zipfile.ZipFile(docx_path, "r") as z:
        xml_text = z.read("word/document.xml").decode("utf-8", errors="ignore")
    leftovers = [k for k in replacements if k in xml_text]
    if leftovers:
        raise RuntimeError("В DOCX остались незаменённые placeholders: " + ", ".join(leftovers))

    convert_docx_to_pdf(docx_path, pdf_path)

    if not pdf_path.exists() or pdf_path.stat().st_size <= 0:
        raise RuntimeError("Не удалось сформировать PDF заявления.")

    return docx_path, pdf_path, letter_text


# ============================================================
# CRM: ОТБОР СДЕЛОК
# ============================================================
def get_db_connection():
    # На разных машинах установлен разный ODBC-драйвер (17 или 18) — пробуем оба,
    # как в otmeny.py / reestr_gosposhliny.py.
    last_err = None
    for driver in ("ODBC Driver 18 for SQL Server", "ODBC Driver 17 for SQL Server"):
        conn_str = (
            f"DRIVER={{{driver}}};"
            f"SERVER={CRM_DB['server']};"
            f"DATABASE={CRM_DB['database']};"
            f"UID={CRM_DB['username']};"
            f"PWD={CRM_DB['password']};"
            "TrustServerCertificate=yes;"
            "Encrypt=no;"
        )
        try:
            return pyodbc.connect(conn_str, timeout=30)
        except pyodbc.Error as e:
            last_err = e
    raise last_err


def _read_processimport_eids(path):
    if not Path(path).exists():
        return set()
    pi = pd.read_excel(path, usecols=["Уникальный номер сделки"], dtype=str, engine="openpyxl")
    return {_id_str(v) for v in pi["Уникальный номер сделки"].dropna() if _id_str(v)}


def load_source():
    log("Подключение к БД...")
    conn = get_db_connection()
    try:
        df = pd.read_sql(SQL_QUERY, conn)
        log(f"SQL выполнен: {len(df)} строк")
    finally:
        conn.close()
    df["LoanID"] = df["LoanID"].apply(_id_str)
    df["EID"] = df["EID"].apply(_id_str)
    df["IIN"] = df["IIN"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    df["Sud"] = df["Sud"].apply(normalize_text)
    df["FIO"] = df["FIO"].apply(normalize_text)
    df = df.drop_duplicates(subset=["LoanID"]).reset_index(drop=True)
    if SKIP_ALREADY_CREATED:
        if RESULT_XLSX.exists():
            old = pd.read_excel(RESULT_XLSX)
            if {"Status", "LoanID"}.issubset(old.columns):
                status = old["Status"].astype(str)
                done = set(old.loc[status.isin(DONE_STATUSES), "LoanID"].apply(_id_str))
                before = len(df)
                df = df[~df["LoanID"].isin(done)].reset_index(drop=True)
                log(f"По журналу SENT/PENDING пропущено: {before - len(df)}")
                pending = old.loc[status.eq("SENT_PENDING_VERIFY")]
                if not pending.empty:
                    log(f"⚠ В журнале {len(pending)} записей SENT_PENDING_VERIFY — отправка не подтверждена, "
                        f"проверьте их в кабинете вручную (declarationId: "
                        f"{', '.join(_id_str(v) for v in pending['DeclarationID'])})")
        eids = _read_processimport_eids(PROCESS_IMPORT_REGISTRY)
        log(f"EID в постоянном реестре ProcessImport: {len(eids)}")
        if eids:
            before = len(df)
            df = df[~df["EID"].isin(eids)].reset_index(drop=True)
            log(f"По EID в ProcessImport пропущено: {before - len(df)}")
    if MAX_ROWS:
        df = df.head(MAX_ROWS).copy()
    log(f"К обработке: {len(df)}")
    return df


# ============================================================
# API portal-sot.kz — через общий PortalSotHttp
# ============================================================
PORTAL = None  # PortalSotHttp, открывается в main()


class ApiHttpError(RuntimeError):
    def __init__(self, status_code, message):
        super().__init__(message)
        self.status_code = status_code


def api_data(payload):
    if not isinstance(payload, dict):
        return payload
    if payload.get("status") == 0 and "data" in payload:
        return payload["data"]
    if "result" in payload:
        return payload["result"]
    if "organization" in payload:
        return payload["organization"]
    return payload


def api_request(method, path, network_retries=3, **kwargs):
    """Темп, обновление токена и отступ на 429/5xx — внутри PortalSotHttp.
    network_retries — повторы при обрыве сети (для LETTER/sign = 1: повтор
    уже отправленной подписи недопустим)."""
    for attempt in range(1, network_retries + 1):
        try:
            r = PORTAL.request(method, path, **kwargs)
            break
        except (requests.ConnectionError, requests.Timeout):
            if attempt >= network_retries:
                raise
            wait = 3 * attempt
            log(f"[RETRY] Сетевой обрыв. Повтор через {wait} сек...")
            time.sleep(wait)
    if r.status_code in (403, 429):
        raise PortalBlocked(f"Портал ограничил запросы: HTTP {r.status_code}. Массовая обработка "
                            f"остановлена, чтобы не усиливать блокировку.")
    if not r.ok:
        raise ApiHttpError(r.status_code, f"{method} {path} -> HTTP {r.status_code}: {r.text[:1000]}")
    try:
        payload = r.json()
    except Exception:
        return r
    if isinstance(payload, dict):
        st = payload.get("status")
        # У ответов ГБД status — словарь, это не код ошибки.
        if not isinstance(st, dict) and st not in (None, 0):
            raise RuntimeError(f"API error: {payload}")
    return payload


# ============================================================
# СПРАВОЧНИК СУДОВ PORTAL-SOT.KZ
# ============================================================
def _court_norm(v):
    s = normalize_text(v).casefold().replace("ё", "е")
    return re.sub(r"\s+", " ", s).strip()


def _court_base(v):
    return re.sub(r"\s*\([^)]*\)\s*$", "", _court_norm(v)).strip()


def _find_court_cache():
    """Справочник общий с подачей исков (на диске любой компании); запасной —
    копия в scripts/data/."""
    roots = [ROOT] + [r for r in COMPANY_ROOTS.values() if r != ROOT]
    candidates = [Path(r) / COURTS_REL for r in roots] + [_SCRIPT_DIR / "data" / COURTS_FILE_NAME]
    for p in candidates:
        try:
            if p.exists():
                return p
        except OSError:
            pass
    return None


def _load_courts_xlsx(path):
    x = pd.read_excel(path, sheet_name=0, dtype=str).fillna("")
    required = {"Регион", "CODE суда", "Наименование суда"}
    if not required.issubset(set(x.columns)):
        raise RuntimeError(f"Неверный формат справочника судов: {path}")
    rows = []
    for _, r in x.iterrows():
        code = normalize_text(r["CODE суда"])
        name = normalize_text(r["Наименование суда"])
        region = normalize_text(r["Регион"])
        if code and name:
            rows.append(({"VALUE_RU": region}, {"CODE": code, "VALUE_RU": name}))
    return rows


def _download_courts():
    log("[СУД] Локальный справочник не найден. Получаю справочник portal-sot...")
    districts = api_data(api_request("GET", "/api/public/dictionary/district/search?CODE%21=23"))
    all_rows = []
    for i, d in enumerate(districts or [], 1):
        log(f"[СУД] {i}/{len(districts)}: {normalize_text(d.get('VALUE_RU'))}")
        try:
            courts = api_data(api_request("GET", f"/api/public/dictionary/court/search?DISTRICT_ID={d['ID']}"))
        except PortalBlocked:
            raise
        except Exception as e:
            log(f"[СУД]   предупреждение: {type(e).__name__}: {e}")
            continue
        for c in courts or []:
            all_rows.append((d, c))
    if not all_rows:
        raise RuntimeError("Не удалось получить справочник судов portal-sot")
    return all_rows


COURT_ALIASES = {
    # CRM name -> exact current portal-sot court name
    "Межрайонный суд города Туркестана":
        "Межрайонный суд города Туркестана Туркестанской области (Гражданские, уголовные дела)",
    "Бурабайский районный суд Акмолинской области (Общая юрисдикция)":
        "Бурабайский районный суд с (Общая юрисдикция)",
    "Жанаозенский городской суд (Общая юрисдикция)":
        "Жанаозенский городской суд Мангистауской области (Общая юрисдикция)",
}

_COURT_ROWS = None


def _court_rows():
    global _COURT_ROWS
    if _COURT_ROWS is None:
        cache = _find_court_cache()
        if cache:
            _COURT_ROWS = _load_courts_xlsx(cache)
            log(f"[СУД] Справочник: {cache} — судов: {len(_COURT_ROWS)}")
        else:
            _COURT_ROWS = _download_courts()
    return _COURT_ROWS


def resolve_court_portal(region_name, court_name):
    original_court_name = court_name
    court_name = COURT_ALIASES.get(str(court_name).strip(), court_name)
    if court_name != original_court_name:
        log(f"[СУД] Алиас CRM → portal-sot: {court_name}")
    rows = _court_rows()

    target = _court_norm(court_name)
    target_base = _court_base(court_name)
    exact = [(d, c) for d, c in rows if _court_norm(c.get("VALUE_RU")) == target]
    base = [(d, c) for d, c in rows if _court_base(c.get("VALUE_RU")) == target_base]
    matches = exact if exact else base

    uniq = {}
    for d, c in matches:
        uniq[str(c.get("CODE"))] = (d, c)
    matches = list(uniq.values())

    if len(matches) > 1:
        nr = _court_norm(region_name)
        narrowed = [(d, c) for d, c in matches
                    if nr and (nr in _court_norm(d.get("VALUE_RU")) or _court_norm(d.get("VALUE_RU")) in nr)]
        if len(narrowed) == 1:
            matches = narrowed

    if len(matches) != 1:
        variants = [(c.get("CODE"), c.get("VALUE_RU"), d.get("VALUE_RU")) for d, c in matches][:20]
        raise RuntimeError(f"Суд не найден однозначно: {court_name!r}; совпадения={variants}")

    return matches[0]


# ============================================================
# LETTER API: CREATE -> PDF -> BLANK -> NCALayer -> SIGN
# ============================================================
def create_letter_draft():
    data = api_data(api_request("POST", "/api/secure/declaration/LETTER/create", json={}))
    if not isinstance(data, dict) or not data.get("id"):
        raise RuntimeError(f"LETTER/create не вернул declarationId: {data}")
    return data


def upload_main_pdf(declaration_id, pdf_path):
    endpoint = f"/api/secure/declaration/attachment/saveAttachment/{declaration_id}"
    mime = mimetypes.guess_type(pdf_path.name)[0] or "application/pdf"
    # bytes, а не открытый файл: PortalSotHttp может повторить запрос (401/5xx)
    files = {"file": (pdf_path.name, pdf_path.read_bytes(), mime)}
    data = {"type": "MAIN_DECLARATION_FILE", "filename": pdf_path.name}
    return api_data(api_request("POST", endpoint, files=files, data=data))


def save_letter_blank(declaration_id, court_code, summary):
    payload = {
        "courtCode": str(court_code),
        "summary": summary,
        "title": "Письмо",
        "declarationId": int(declaration_id),
        "language": "ru",
    }
    return api_data(api_request("POST", "/api/secure/declaration/LETTER/blank", json=payload))


def get_letter(declaration_id):
    return api_data(api_request("GET", f"/api/secure/declaration/LETTER/{int(declaration_id)}"))


def extract_request_uid(blank_data):
    xml = (blank_data or {}).get("xmlForSign", "") if isinstance(blank_data, dict) else ""
    m = re.search(r"<f1>([^<]+)</f1>", xml)
    return m.group(1).strip() if m else ""


def save_result_row(item):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame([item], columns=RESULT_COLUMNS)
    if RESULT_XLSX.exists():
        try:
            prev = pd.read_excel(RESULT_XLSX)
            df = pd.concat([prev, df], ignore_index=True)
        except Exception:
            pass
    df.to_excel(RESULT_XLSX, index=False)


def ncalayer_sign_xml(xml_to_sign, timeout=180):
    """
    Подписывает XML через NCALayer (kz.gov.pki.knca.basics / sign, как в ноутбуке).
    Окно NCALayer (пароль ЭЦП → «Открыть» → «Подписать») подтверждает pywinauto —
    та же механика, что при входе (portal_sot_auth._handle_ncalayer_dialog).
    """
    import websocket

    ws = websocket.create_connection(
        portal_sot_auth.NCALAYER_WS_URL,
        timeout=15,
        origin=BASE_URL,
        sslopt={"cert_reqs": ssl.CERT_NONE},
    )
    try:
        ws.recv()  # приветствие NCALayer

        request_obj = {
            "module": "kz.gov.pki.knca.basics",
            "method": "sign",
            "args": {
                "allowedStorages": ["PKCS12"],
                "format": "xml",
                "data": xml_to_sign,
                "signingParams": {
                    "decode": False,
                    "encapsulate": True,
                    "digested": False,
                    "tsaProfile": {},
                },
                "signerParams": {
                    "extKeyUsageOids": [],
                    "chain": [],
                },
                "locale": "ru",
            },
        }

        nca_auto_error = []

        def _auto_confirm_ncalayer_sign():
            try:
                portal_sot_auth._handle_ncalayer_dialog(PORTAL_CFG["eds_password"], log=log)
            except Exception as exc:
                nca_auto_error.append(exc)

        # Сначала запускаем ожидание окна, затем отправляем команду sign.
        threading.Thread(target=_auto_confirm_ncalayer_sign, name="NCALayerAutoConfirm", daemon=True).start()
        ws.send(json.dumps(request_obj, ensure_ascii=False))

        ws.settimeout(5)
        deadline = time.time() + timeout
        last_payload = None

        while time.time() < deadline:
            if nca_auto_error:
                raise RuntimeError("Автоподтверждение окна NCALayer не выполнено: " + repr(nca_auto_error[0]))
            try:
                raw = ws.recv()
            except websocket.WebSocketTimeoutException:
                continue
            if raw == "--heartbeat--":
                try:
                    ws.send("--heartbeat--")
                except Exception:
                    pass
                continue

            last_payload = raw
            try:
                obj = json.loads(raw)
            except Exception:
                continue

            if obj.get("status") is False:
                raise RuntimeError("NCALayer/sign: " + json.dumps(obj, ensure_ascii=False)[:4000])

            # NCALayer 1.4 фактически может вернуть {"body": {"result": ["<signed xml>"]}}.
            # Поддерживаем строки, списки строк и вложенные словари.
            candidates = []

            def collect_signed_candidates(value):
                if isinstance(value, str):
                    candidates.append(value)
                elif isinstance(value, list):
                    for item in value:
                        collect_signed_candidates(item)
                elif isinstance(value, dict):
                    for k in ("result", "data", "signedData", "signedXml", "body", "responseObject"):
                        if k in value:
                            collect_signed_candidates(value.get(k))

            collect_signed_candidates(obj)

            for candidate in candidates:
                if "<ds:Signature" in candidate or "<Signature" in candidate:
                    return candidate

            # Не печатаем полный XML/сертификат в ошибку.
            if obj.get("status") is True or str(obj.get("code")) == "200":
                raise RuntimeError("NCALayer сообщил успех, но в ответе не найден XML с ds:Signature")

        raise RuntimeError(f"NCALayer не завершил подпись за {timeout} сек. "
                           f"Последний ответ: {str(last_payload)[:300]!r}")
    finally:
        try:
            ws.close()
        except Exception:
            pass


class SentUnverified(RuntimeError):
    """LETTER/sign ушёл, но ни его ответ, ни статус письма получить не удалось."""


def sign_letter(declaration_id, signed_xml):
    # HAR portal-sot 06.10.2026: POST /LETTER/sign принимает Base64 UTF-8 подписанного XML.
    signed_b64 = base64.b64encode(signed_xml.encode("utf-8")).decode("ascii")
    payload = {"declarationId": int(declaration_id), "signedXml": signed_b64}
    # Без повторов (ни на 5xx, ни на обрыв): одна подпись — один POST.
    return api_data(api_request("POST", "/api/secure/declaration/LETTER/sign",
                                json=payload, backoff=(), network_retries=1))


def letter_request_status(declaration_id, attempts=4, delay=3.0):
    """requestStatus письма или None, если портал так и не ответил."""
    for n in range(1, attempts + 1):
        try:
            data = get_letter(declaration_id)
            if isinstance(data, dict):
                return str(data.get("requestStatus") or "")
        except PortalBlocked:
            raise
        except Exception as e:
            log(f"Проверка статуса письма {n}/{attempts}: {type(e).__name__}: {e}")
        if n < attempts:
            time.sleep(delay)
    return None


# ============================================================
# ProcessImport для Дельты
# ============================================================
def ensure_process_import_registry():
    """Постоянная история всех EID. Никогда не очищается автоматически."""
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    if not PROCESS_IMPORT_REGISTRY.exists():
        shutil.copy2(PROCESS_IMPORT_TEMPLATE, PROCESS_IMPORT_REGISTRY)
        log(f"✓ Создан постоянный {PROCESS_IMPORT_REGISTRY}")
    return PROCESS_IMPORT_REGISTRY


def start_new_process_import():
    """Новый уникальный ProcessImport только для текущего запуска."""
    global PROCESS_IMPORT_OUTPUT
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    PROCESS_IMPORT_OUTPUT = OUT_DIR / f"{KCFG['pi_current']}_{stamp}.xlsx"
    shutil.copy2(PROCESS_IMPORT_TEMPLATE, PROCESS_IMPORT_OUTPUT)
    return PROCESS_IMPORT_OUTPUT


def _append_process_row(path, eid, request_uid, submitted_at=None):
    from openpyxl import load_workbook
    eid = _id_str(eid)
    uid = str(request_uid or "").strip()
    if not eid or not uid:
        raise RuntimeError("ProcessImport: пустой EID или талон")

    wb = load_workbook(path)
    ws = wb.active
    headers = {str(ws.cell(1, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    missing = [h for h in PROCESS_IMPORT_REQUIRED_HEADERS if h not in headers]
    if missing:
        wb.close()
        raise RuntimeError(f"ProcessImport: нет колонок {missing}")

    ce = headers["Уникальный номер сделки"]
    cd = headers[PROCESS_IMPORT_DATE_HEADER]
    cu = headers["Номер талона письма ЕПСО"]
    ct = headers["Тип процесса"]
    cs = headers["Статус процесса"]

    existing = None
    for r in range(2, ws.max_row + 1):
        if _id_str(ws.cell(r, ce).value) == eid:
            existing = r
            break

    target = existing or next(
        (r for r in range(2, ws.max_row + 1) if ws.cell(r, ce).value in (None, "")),
        ws.max_row + 1,
    )
    old_uid = str(ws.cell(target, cu).value or "").strip()
    if existing and old_uid and old_uid != uid:
        wb.close()
        raise RuntimeError(f"ProcessImport: конфликт EID={eid}: {old_uid} != {uid}")

    d = (submitted_at or datetime.now()).date()
    ws.cell(target, ce).value = eid
    if not ws.cell(target, cd).value:
        ws.cell(target, cd).value = d
    ws.cell(target, cd).number_format = "dd.mm.yyyy"
    if not old_uid:
        ws.cell(target, cu).value = uid
    ws.cell(target, ct).value = PROCESS_TYPE_DEFAULT
    ws.cell(target, cs).value = PROCESS_STATUS_DEFAULT

    wb.save(path)
    wb.close()
    return d


def append_process_import(eid, request_uid, submitted_at=None):
    """
    После успешной отправки:
      1) строка идёт в ProcessImport текущего запуска;
      2) та же строка идёт в постоянный ProcessImport_registry.xlsx.
    """
    d = _append_process_row(PROCESS_IMPORT_OUTPUT, eid, request_uid, submitted_at)
    _append_process_row(ensure_process_import_registry(), eid, request_uid, submitted_at)
    log(f"✓ ProcessImport: EID={_id_str(eid)} | {d:%d.%m.%Y} | {request_uid}")


def publish_process_import():
    """В папку автоимпорта уходит ТОЛЬКО ProcessImport текущего запуска (registry — нет)."""
    PROCESS_IMPORT_NETWORK.parent.mkdir(parents=True, exist_ok=True)
    tmp = PROCESS_IMPORT_NETWORK.with_name("ProcessImport.__tmp__.xlsx")
    shutil.copy2(PROCESS_IMPORT_OUTPUT, tmp)
    try:
        os.replace(tmp, PROCESS_IMPORT_NETWORK)
    except PermissionError:
        try:
            if tmp.exists():
                tmp.unlink()
        except Exception:
            pass
        raise RuntimeError(
            f"{PROCESS_IMPORT_NETWORK} открыт в Excel. Закройте его и скопируйте туда вручную "
            f"{PROCESS_IMPORT_OUTPUT.name} из результатов задачи (под именем ProcessImport.xlsx)."
        )
    log(f"✓ ProcessImport текущего запуска опубликован: {PROCESS_IMPORT_NETWORK}")


# ============================================================
# ОДНА СДЕЛКА
# ============================================================
def run_one_il(row):
    log("=" * 100)
    log(f"EID={row['EID']} | ИИН={row['IIN']} | {row['FIO']}")
    log(f"Суд CRM: {row['Sud']}")

    # Незавершённые SAVED_FOR_SIGNING/SIGN_ERROR прошлых запусков не продолжаем:
    # для каждой такой сделки создаётся новый LETTER с нуля.
    district, court = resolve_court_portal(row.get("Region", ""), row["Sud"])
    log(f"Суд portal-sot: {court['CODE']} | {court['VALUE_RU']}")

    docx_path, pdf_path, letter_text = make_documents(row)
    log(f"PDF: {pdf_path}")

    decl = create_letter_draft()
    did = int(decl["id"])
    log(f"declarationId: {did}")

    upload_main_pdf(did, pdf_path)
    log("✓ PDF загружен как MAIN_DECLARATION_FILE")

    blank = save_letter_blank(did, court["CODE"], letter_text)
    xml_to_sign = (blank or {}).get("xmlForSign", "") if isinstance(blank, dict) else ""
    uid = extract_request_uid(blank)
    if not xml_to_sign:
        raise RuntimeError("LETTER/blank не вернул xmlForSign")

    log("✓ Электронный бланк сохранён")
    log(f"Талон: {uid or '(не распознан из xmlForSign)'}")

    # Фиксируем созданный черновик в журнале для аудита.
    draft_item = {
        "LoanID": str(row["LoanID"]),
        "EID": str(row["EID"]),
        "IIN": str(row["IIN"]),
        "FIO": row["FIO"],
        "CourtCRM": row["Sud"],
        "CourtCode": str(court["CODE"]),
        "CourtPortal": court["VALUE_RU"],
        "DeclarationID": did,
        "RequestUID": uid,
        "PDF": str(pdf_path),
        "Status": "SAVED_FOR_SIGNING",
        "Error": "",
    }
    save_result_row(draft_item)

    if not SIGN_AND_SEND:
        log("Режим «без подписи»: черновик оставлен в кабинете, не подписан и не отправлен")
        return draft_item

    signed_xml = ncalayer_sign_xml(xml_to_sign)
    log("✓ NCALayer вернул XML с ЭЦП")
    try:
        sign_letter(did, signed_xml)
    except (requests.ConnectionError, requests.Timeout, ApiHttpError) as exc:
        if isinstance(exc, ApiHttpError) and exc.status_code < 500:
            raise   # портал явно отказал — письмо не отправлено
        # Обрыв сети или 5xx: ушла подпись или нет — неизвестно, узнаём у портала.
        # (GET LETTER нестабилен, бывает 415 — тогда остаётся SENT_PENDING_VERIFY.)
        status = letter_request_status(did)
        if status is None:
            pending = dict(draft_item, Status="SENT_PENDING_VERIFY", Error=repr(exc))
            save_result_row(pending)
            raise SentUnverified(
                f"LETTER/sign отправлен, но статус письма {did} не подтверждён ({exc}); "
                f"повторная подача заблокирована — проверьте в кабинете вручную"
            )
        if status != "SENDED":
            raise
        log(f"LETTER/sign вернул ошибку ({exc}), но портал показывает SENDED — считаю отправленным")
    log("✓ ПОДПИСАНО И ОТПРАВЛЕНО")

    # Письмо уже в суде: сбой записи в журнал/ProcessImport не должен превращать
    # сделку в «ошибку», иначе следующий запуск подаст её повторно.
    item = dict(draft_item, Status="SENT")
    manual = f"Внесите вручную: EID={row['EID']}, талон {uid}, declarationId {did}"
    try:
        save_result_row(item)
    except Exception as exc:
        log(f"⚠⚠ ОТПРАВЛЕНО, но не записано в журнал: {type(exc).__name__}: {exc}. {manual}")
    try:
        append_process_import(row["EID"], uid)
    except Exception as exc:
        log(f"⚠⚠ ОТПРАВЛЕНО, но не записано в ProcessImport: {type(exc).__name__}: {exc}. {manual}")
        try:
            save_result_row(dict(item, Status="SENT_PROCESSIMPORT_PENDING", Error=repr(exc)))
        except Exception:
            pass
    return item


# ============================================================
# ЗАПУСК
# ============================================================
def _portal_login():
    return login_chain(COMPANY_ID, PORTAL_CFG, log=log, dump_dir=str(OUT_DIR))


def main():
    global PORTAL

    log("=" * 90)
    log(f"portal-sot.kz — заявления на выдачу ИЛ {KCFG['title']}")
    log(f"Компания: {COMPANY_ID} | фильтр F209: {DB_COMPANY_FILTER}"
        + (f" + кредитор F234 LIKE '{CREDITOR_LIKE}'" if CREDITOR_LIKE else "")
        + f" | истец: {PROJECT_CFG['project_name']}")
    log("Режим: " + ("АВТОПОДПИСЬ ЭЦП И ОТПРАВКА" if SIGN_AND_SEND else "ТОЛЬКО ЧЕРНОВИКИ, без подписи"))
    log(f"Шаблон: {IL_TEMPLATE.name} | {STAMP_NOTE}")
    log(f"Журнал и реестр: {STATE_DIR}")
    log("=" * 90)

    for p in (IL_TEMPLATE, PROCESS_IMPORT_TEMPLATE):
        if not p.exists():
            log(f"❌ Не найден шаблон: {p}")
            return 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ensure_process_import_registry()
    start_new_process_import()

    source = load_source()
    success, errors = [], []
    blocked = False

    if source.empty:
        log("Нет сделок для обработки.")
    else:
        try:
            with PortalSotHttp(COMPANY_ID, _portal_login, log=log) as portal:
                PORTAL = portal

                # Проверка авторизации
                org = api_data(api_request("GET", f"/api/secure/gbdul/v2/byBin/{ORG_BIN}"))
                log(f"Авторизация portal-sot.kz: OK — "
                    f"{org.get('fullNameRu') if isinstance(org, dict) else ORG_BIN}")

                if SIGN_AND_SEND:
                    # NCALayer должен работать с сертификатом ЭТОЙ компании (вход мог пройти
                    # по сохранённому токену, без NCALayer).
                    portal_sot_auth.ensure_ncalayer_running(COMPANY_ID, NCALAYER_PATH, log=log)

                for pos, (_, row) in enumerate(source.iterrows(), 1):
                    log(f"[{pos}/{len(source)}] LoanID={row['LoanID']} | EID={row['EID']}")
                    try:
                        item = run_one_il(row)
                        success.append(item)
                        log(f"✓ [{pos}/{len(source)}] {item['Status']} | {row['IIN']} | {item.get('RequestUID', '')}")
                    except Exception as e:
                        if not isinstance(e, SentUnverified):
                            # Следующий запуск создаст новый LETTER, если EID не попал в ProcessImport.
                            save_result_row({
                                "LoanID": str(row.get("LoanID", "")), "EID": str(row.get("EID", "")),
                                "IIN": str(row.get("IIN", "")), "FIO": row.get("FIO", ""),
                                "CourtCRM": row.get("Sud", ""), "CourtCode": "", "CourtPortal": "",
                                "DeclarationID": "", "RequestUID": "", "PDF": "",
                                "Status": "SIGN_ERROR", "Error": repr(e),
                            })
                        errors.append(row.get("EID", ""))
                        log(f"✗ [{pos}/{len(source)}] ОШИБКА: {type(e).__name__}: {e}")
                        if isinstance(e, PortalBlocked):
                            log("СТОП: портал ограничил запросы. Остальные сделки не обрабатываются.")
                            blocked = True
                            break

                    if pos < len(source) and PAUSE_BETWEEN_ROWS:
                        time.sleep(PAUSE_BETWEEN_ROWS)
        except PortalBlocked as e:
            log(f"СТОП: {e}")
            blocked = True

    sent = [i for i in success if i["Status"] == "SENT"]
    log("=" * 100)
    log(f"Отправлено: {len(sent)} | черновиков без подписи: {len(success) - len(sent)} | ошибок: {len(errors)}")

    if RESULT_XLSX.exists():
        try:
            shutil.copy2(RESULT_XLSX, OUT_DIR / RESULT_XLSX.name)
        except Exception as e:
            log(f"Не удалось скопировать журнал в результаты задачи: {e}")
    log(f"Журнал: {RESULT_XLSX}")

    if sent:
        publish_process_import()
    else:
        log("В этом запуске ничего не отправлено — ProcessImport.xlsx в автоимпорт не публикуется")

    return 1 if (blocked or (errors and not success)) else 0


if __name__ == "__main__":
    try:
        code = main()
    except Exception:
        traceback.print_exc()
        code = 1
    sys.exit(code)
