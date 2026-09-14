# app.py
import ast
import os
import uuid
import shutil
import subprocess
import threading
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Any, Optional
from database import init_db, get_db, User, Company
from fastapi import FastAPI, Request, Depends, Form, HTTPException, UploadFile, File, Cookie
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, Response, JSONResponse
from jinja2 import Environment, FileSystemLoader, select_autoescape
from sqlalchemy.orm import Session

from scripts_registry import SCRIPTS
from database import (
    init_db, get_db, User, Company, SessionLocal,
    Schedule, get_all_schedules, get_active_schedules, get_schedule,
    create_schedule, update_schedule, delete_schedule, toggle_schedule,
    create_job_run, update_job_run_status, get_job_runs, get_job_run_script_keys, count_job_runs,
    get_user_companies, set_user_companies,
    JobRun,
    Mailing,
    get_mailings, get_mailing, create_mailing, update_mailing,
    delete_mailing, toggle_mailing, get_scheduled_mailings,
    get_mailing_reports, get_mailing_report, create_mailing_report,
    update_mailing_report, delete_mailing_report,
    get_chsi_contacts, get_chsi_contact, create_chsi_contact,
    update_chsi_contact, delete_chsi_contact, bulk_add_chsi_contacts,
)
from mailing_common import PLACEHOLDERS
from auth import (
    create_session, get_current_user, get_admin_user,
    verify_password, SESSIONS, user_can_access_company,
    get_user_allowed_company_ids,
    get_active_company_id, set_active_company_id,
)
import aisoip_web

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data" / "runs"
DATA_DIR.mkdir(parents=True, exist_ok=True)

# Папка для загружаемых пользователем файлов
UPLOADS_DIR = BASE_DIR / "uploads"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI()

# Инициализация БД
init_db()
aisoip_web.init_aisoip_db()

env = Environment(
    loader=FileSystemLoader(str(BASE_DIR)),
    autoescape=select_autoescape(["html", "xml"]),
)

JOBS: Dict[str, Dict[str, Any]] = {}
QUEUE = []
LOCK = threading.Lock()
MAX_PARALLEL_JOBS = 3
RUNNING_COUNT = 0
PROCESSES: Dict[str, subprocess.Popen] = {}  # job_id -> запущенный процесс (для остановки из интерфейса)


def kill_process_tree(pid: int):
    """Убивает процесс и всех его потомков (например, Chrome/chromedriver,
    запущенные скриптом через Selenium) — обычный Popen.terminate() убивает
    только сам python.exe и оставляет дочерние процессы висеть."""
    try:
        subprocess.run(
            ["taskkill", "/F", "/T", "/PID", str(pid)],
            capture_output=True, timeout=15,
        )
    except Exception as e:
        print(f"KILL PROCESS TREE ERROR (pid={pid}):", e)

def runner_loop():
    global RUNNING_COUNT
    while True:
        with LOCK:
            if RUNNING_COUNT < MAX_PARALLEL_JOBS and QUEUE:
                job_id = QUEUE.pop(0)
                RUNNING_COUNT += 1
                job = JOBS[job_id]
                job["status"] = "running"
                job["started_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                update_job_run_status(job_id, "running", started_at=datetime.now())
                threading.Thread(target=execute_job, args=(job_id,), daemon=True).start()
        import time
        time.sleep(0.3)

def render(html: str, **ctx):
    template = env.from_string(html)
    return template.render(**ctx)


# ─────────────────────────────────────────────────────────────
#  Папка-корень для скриптов переименования / конвертации.
#
#  Сотрудники запускают эти скрипты с разных компьютеров, где сетевая
#  шара смонтирована по-разному (Z:\, \\SRVAPP\..., \\192.168.1.x\...).
#  Скрипт же выполняется на сервере, поэтому путь должен быть таким,
#  каким его видит СЕРВЕР приложения. Задаётся один раз на сервере —
#  в переменной окружения RENAME_ROOT или строкой
#      RENAME_ROOT = r'...'
#  в scripts/config.py (этот файл свой на каждой машине, см. .gitignore).
#  В форме сотрудник только выбирает подпапку из выпадающего списка.
# ─────────────────────────────────────────────────────────────
def _rename_root() -> Optional[Path]:
    env_val = os.environ.get("RENAME_ROOT")
    if env_val:
        return Path(env_val.strip().strip('"'))
    cfg = BASE_DIR / "scripts" / "config.py"
    if cfg.exists():
        try:
            tree = ast.parse(cfg.read_text(encoding="utf-8"))
            for node in tree.body:
                if isinstance(node, ast.Assign) and any(
                    isinstance(t, ast.Name) and t.id == "RENAME_ROOT" for t in node.targets
                ):
                    val = ast.literal_eval(node.value)
                    return Path(val) if val else None
        except Exception:
            return None
    return None


def list_rename_folders() -> list:
    """Непосредственные подпапки RENAME_ROOT, новые — первыми."""
    root = _rename_root()
    if not root:
        return []
    try:
        subs = [p for p in root.iterdir() if p.is_dir()]
        subs.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    except OSError:
        return []
    return [{"value": p.name, "label": p.name} for p in subs]


def resolve_rename_folder(raw: str) -> Optional[str]:
    """Превращает выбор из формы в абсолютный путь, который получит скрипт.

    Принимает: имя подпапки (или вложенный относительный путь) внутри
    RENAME_ROOT, либо — как запасной вариант — уже готовый абсолютный путь.
    Любой относительный путь резолвится строго внутри RENAME_ROOT (защита
    от '..' и подстановки чужого пути)."""
    raw = (raw or "").strip().strip('"')
    if not raw:
        return None
    p = Path(raw)
    if p.is_absolute():
        return str(p) if p.is_dir() else None
    root = _rename_root()
    if not root:
        return None
    cand = (root / raw).resolve()
    try:
        cand.relative_to(root.resolve())
    except ValueError:
        return None
    return str(cand) if cand.is_dir() else None


_scheduler_ran_today = {}  # {(schedule_id, "YYYY-MM-DD"): True}

def enqueue_scheduled_job(script_key: str, params: dict, company_id: Optional[int] = None, company_name: Optional[str] = None):
    job_id = uuid.uuid4().hex[:12]
    job_dir = make_job_dir(job_id)
    job = {
        "id": job_id,
        "script_key": script_key,
        "params": params,
        "status": "queued",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "started_at": None,
        "finished_at": None,
        "job_dir": str(job_dir),
        "result_zip": None,
        "created_ts": datetime.now().timestamp(),
        "username": "⏰ Автозапуск",
        "company_name": company_name or "Все компании",
        "company_id": company_id,
        "source": "scheduler",
    }
    JOBS[job_id] = job
    create_job_run(
        job_id, script_key, source="scheduler", username="scheduler",
        company_id=company_id, company_name=company_name,
    )
    with LOCK:
        QUEUE.append(job_id)
    return job_id

def scheduler_loop():
    while True:
        try:
            now = datetime.now()
            today_key = now.strftime("%Y-%m-%d")
            for s in get_active_schedules():
                if s.hour == now.hour and s.minute == now.minute:
                    if now.weekday() not in s.weekday_list():
                        continue
                    key = (s.id, today_key)
                    if _scheduler_ran_today.get(key):
                        continue
                    _scheduler_ran_today[key] = True
                    if s.script_key in SCRIPTS:
                        company_name = None
                        if s.company_id:
                            db = SessionLocal()
                            try:
                                c = db.query(Company).filter(Company.id == s.company_id).first()
                                company_name = c.name if c else None
                            finally:
                                db.close()
                        enqueue_scheduled_job(s.script_key, s.params(), company_id=s.company_id, company_name=company_name)

            # ── Рассылки по отчётам со своим расписанием ──
            for m in get_scheduled_mailings():
                if m.schedule_hour == now.hour and m.schedule_minute == now.minute:
                    if now.weekday() not in m.weekday_list():
                        continue
                    key = ("mailing", m.id, today_key)
                    if _scheduler_ran_today.get(key):
                        continue
                    _scheduler_ran_today[key] = True
                    company_name = None
                    if m.company_id:
                        db = SessionLocal()
                        try:
                            c = db.query(Company).filter(Company.id == m.company_id).first()
                            company_name = c.name if c else None
                        finally:
                            db.close()
                    enqueue_scheduled_job(
                        "mailing_send",
                        {"mailing_id": str(m.id), "mode": "real",
                         "company_id": ",".join(str(x) for x in m.company_id_list())},
                        company_id=m.company_id, company_name=company_name,
                    )
        except Exception as e:
            print("SCHEDULER ERROR:", e)
        import time
        time.sleep(20)

threading.Thread(target=scheduler_loop, daemon=True).start()
# ═══════════════════════════════════════════════════════════════
# HTML ШАБЛОНЫ
# ═══════════════════════════════════════════════════════════════

LOGIN_HTML = """
<!doctype html>
<html>
<head>
    <meta charset="utf-8">
    <title>Вход в OmegaRunner</title>
    <style>
        body {
            font-family: Arial, sans-serif;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            height: 100vh;
            margin: 0;
            display: flex;
            justify-content: center;
            align-items: center;
        }
        .login-container {
            background: white;
            padding: 40px;
            border-radius: 10px;
            box-shadow: 0 10px 30px rgba(0,0,0,0.2);
            width: 350px;
        }
        h2 {
            text-align: center;
            color: #333;
            margin-bottom: 30px;
        }
        .form-group {
            margin-bottom: 20px;
        }
        label {
            display: block;
            margin-bottom: 5px;
            color: #555;
            font-weight: bold;
        }
        input {
            width: 100%;
            padding: 10px;
            border: 1px solid #ddd;
            border-radius: 5px;
            box-sizing: border-box;
            font-size: 14px;
        }
        button {
            width: 100%;
            padding: 12px;
            background: #667eea;
            color: white;
            border: none;
            border-radius: 5px;
            font-size: 16px;
            cursor: pointer;
            transition: background 0.3s;
        }
        button:hover {
            background: #5a67d8;
        }
        .error {
            color: red;
            text-align: center;
            margin-top: 10px;
        }
        .info {
            text-align: center;
            margin-top: 20px;
            padding-top: 20px;
            border-top: 1px solid #eee;
            color: #666;
            font-size: 12px;
        }
    </style>
</head>
<body>
    <div class="login-container">
        <h2>🔐 Вход в Runner</h2>
        <form method="post" action="/login">
            <div class="form-group">
                <label>Логин:</label>
                <input type="text" name="username" required autofocus>
            </div>
            <div class="form-group">
                <label>Пароль:</label>
                <input type="password" name="password" required>
            </div>
            <button type="submit">Войти</button>
        </form>
        {% if error %}
        <div class="error">{{ error }}</div>
        {% endif %}
        <div class="info">
            Тестовый доступ: admin / admin
        </div>
    </div>
</body>
</html>
"""

@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, error: Optional[str] = None):
    return render(LOGIN_HTML, error=error)

@app.post("/login")
async def login_submit(
    username: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(User.username == username).first()
    
    if not user or not verify_password(password, user.password_hash):
        return HTMLResponse(render(LOGIN_HTML, error="Неверный логин или пароль"))
    
    if not user.is_active:
        return HTMLResponse(render(LOGIN_HTML, error="Пользователь заблокирован"))
    
    session_id = create_session(user.id, user.username, user.role, user.company_id)
    
    response = RedirectResponse("/", status_code=303)
    response.set_cookie("session", session_id, httponly=True)
    return response

@app.get("/logout")
def logout():
    response = RedirectResponse(url="/login", status_code=303)
    response.delete_cookie("session_id")
    return response

INDEX_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Runner</title>
<style>
body { font-family: Arial; max-width: 1400px; margin: 20px auto; padding: 0 20px; background: #f5f7fa; }
.header { display: flex; justify-content: space-between; align-items: center; background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); margin-bottom: 20px; }
.user-info { text-align: right; }
.company-badge { display: inline-block; padding: 4px 12px; background: #28a745; color: white; border-radius: 12px; font-size: 12px; margin-left: 8px; }
.logout-btn { padding: 8px 16px; background: #dc3545; color: white; text-decoration: none; border-radius: 4px; margin-left: 10px; }
.admin-btn { padding: 8px 16px; background: #28a745; color: white; text-decoration: none; border-radius: 4px; margin-right: 10px; }

.company-filter { margin: 20px 0; background: white; padding: 15px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
.company-tab { display: inline-block; padding: 10px 20px; margin-right: 5px; background: #e9ecef; border-radius: 6px; cursor: pointer; text-decoration: none; color: #333; }
.company-tab.active { background: #007bff; color: white; }

.scripts-container { background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); margin-bottom: 30px; }
.category-header { background: #f8f9fa; padding: 12px 15px; margin: 20px -20px 15px -20px; border-left: 4px solid #007bff; font-weight: bold; font-size: 16px; color: #333; }
.category-header:first-child { margin-top: -20px; }

.scripts-table { width: 100%; border-collapse: collapse; }
.scripts-table th { background: #495057; color: white; padding: 12px; text-align: left; font-weight: 500; }
.scripts-table td { padding: 12px; border-bottom: 1px solid #e9ecef; }
.scripts-table tr:hover { background: #f8f9fa; }

.script-name { font-weight: 500; color: #333; }
.script-description { font-size: 13px; color: #6c757d; margin-top: 4px; }

.badge { display: inline-block; padding: 4px 10px; border-radius: 12px; font-size: 11px; font-weight: 500; }
.badge-auto { background: #28a745; color: white; }
.badge-manual { background: #ffc107; color: #000; }
.badge-daily { background: #17a2b8; color: white; }
.badge-weekly { background: #6610f2; color: white; }
.badge-monthly { background: #e83e8c; color: white; }
.badge-ondemand { background: #6c757d; color: white; }

.run-btn { padding: 8px 16px; background: #007bff; color: white; text-decoration: none; border-radius: 4px; font-size: 13px; }
.run-btn:hover { background: #0056b3; }

.history-container { background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
.history-table { width: 100%; border-collapse: collapse; margin-top: 15px; }
.history-table th { background: #007bff; color: white; padding: 12px; text-align: left; }
.history-table td { padding: 10px; border-bottom: 1px solid #e9ecef; font-size: 14px; }
.history-table tr:hover { background: #f8f9fa; }
.status-success { color: #28a745; font-weight: 500; }
.status-error { color: #dc3545; font-weight: 500; }
.status-running { color: #ffc107; font-weight: 500; }
.status-queued { color: #6c757d; font-weight: 500; }
</style>
</head>
<body>

<div class="header">
  <div>
    <h2 style="margin: 0;">📊 Runner</h2>
    <p style="margin: 5px 0 0 0; color: #6c757d; font-size: 14px;">Система автоматизации процессов</p>
  </div>
  <div class="user-info">
    <strong>{{ user.full_name }}</strong> 
    ({{ user.role }})
    {% if user.company %}
      <span class="company-badge">{{ user.company.name }}</span>
    {% elif user.role == 'manager' %}
      <span class="company-badge" style="background:#6c757d;">{{ companies | map(attribute='name') | join(', ') if companies else 'Нет компаний' }}</span>
    {% else %}
      <span class="company-badge" style="background:#6c757d;">Все компании</span>
    {% endif %}
    <a href="/aisoip" class="admin-btn" style="background:#0b4f6c;">🔎 АИС ОИП поиск</a>
    {% if user.role == 'admin' %}<a href="/admin" class="admin-btn">⚙️ Админка</a>{% endif %}
    <a href="/logout" class="logout-btn">Выход</a>
  </div>
</div>

{% if user.role in ('admin', 'manager') %}
<div class="company-filter">
  <strong style="margin-right: 15px;">Компания:</strong>
  <a href="/?company=all" class="company-tab {% if current_company == 'all' %}active{% endif %}">Все компании</a>
  {% for c in companies %}
    <a href="/?company={{ c.id }}" class="company-tab {% if current_company == c.id|string %}active{% endif %}">{{ c.name }}</a>
  {% endfor %}
</div>
{% if active_company %}
  <p style="margin:8px 0 0 0; color:#28a745; font-size:13px;">
    ✅ Скрипты ниже будут запущены для компании: <strong>{{ active_company.name }}</strong>
  </p>
{% else %}
  <p style="margin:8px 0 0 0; color:#dc3545; font-size:13px;">
    ⚠ Компания не выбрана — выберите её вкладкой выше, иначе скрипты, привязанные к компании, запустить не получится.
  </p>
{% endif %}
{% endif %}

<div class="scripts-container">
  <h3 style="margin-top: 0;">Доступные скрипты</h3>
  
  <div class="category-header">📥 Парсинг данных</div>
  <table class="scripts-table">
    <thead>
      <tr>
        <th style="width: 35%;">Наименование</th>
        <th style="width: 25%;">Периодичность</th>
        <th style="width: 15%;">Запуск</th>
        <th style="width: 25%;">Действие</th>
      </tr>
    </thead>
    <tbody>
    {% for s in scripts %}
      {% if s.key in ['ogranichenie', 'otmeny', 'status' , 'sud', 'sudFIO', 'sud_download_mediation', 'reestr_gosposhliny'] %}
      <tr>
        <td>
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td>
          {% if s.key == 'aisoip_status' or s.key == 'otmeny' %}
            <span class="badge badge-daily">Ежедневно</span>

          {% elif s.key == 'sud' %}
            <span class="badge badge-weekly">2 раза в неделю</span>

          {% elif s.key == 'sudFIO' %}
            <span class="badge badge-ondemand">При покупке портфеля</span>

          {% else %}
            <span class="badge badge-ondemand">При необходимости</span>
          {% endif %}

        </td>
        <td>
          {% if s.key == 'ogranichenie' or s.key == 'otmeny' %}
            <span class="badge badge-auto">Автоматически</span>
          {% else %}
            <span class="badge badge-manual">Ручной</span>
          {% endif %}
        </td>
        <td><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>
  
  <div class="category-header">📤 Автоподача заявлений/исков</div>
  <table class="scripts-table">
    <tbody>
    {% for s in scripts %}
      {% if s.key in ['sbor','poiskvsk', 'podacha_iska', 'podacha_iska_v2', 'chsi_podacha', 'reestr_gp_import'] %}
      <tr>
        <td style="width: 35%;">
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td style="width: 25%;"><span class="badge badge-ondemand">При необходимости</span></td>
        <td style="width: 15%;"><span class="badge badge-manual">Ручной</span></td>
        <td style="width: 25%;"><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>
  
  <div class="category-header">🔄 Конвертация документов</div>
  <table class="scripts-table">
    <tbody>
    {% for s in scripts %}
      {% if s.key in ['convert_msg'] %}
      <tr>
        <td style="width: 35%;">
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td style="width: 25%;"><span class="badge badge-ondemand">При покупке портфеля</span></td>
        <td style="width: 15%;"><span class="badge badge-manual">Ручной</span></td>
        <td style="width: 25%;"><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>

  <div class="category-header">💳 Импорт платежей</div>
  <table class="scripts-table">
    <tbody>
    {% for s in scripts %}
      {% if s.key in ['paymentimporter'] %}
      <tr>
        <td style="width: 35%;">
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td style="width: 25%;"><span class="badge badge-ondemand">При необходимости</span></td>
        <td style="width: 15%;"><span class="badge badge-manual">Ручной</span></td>
        <td style="width: 25%;"><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>
  

  <div class="category-header">📈 Отчёты ФРСП</div>
  <table class="scripts-table">
    <tbody>
    {% for s in scripts %}
      {% if s.key in ['frsp_omega', 'frsp_kpi', 'frsp_orion'] %}
      <tr>
        <td style="width: 35%;">
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td style="width: 25%;"><span class="badge badge-ondemand">При необходимости</span></td>
        <td style="width: 15%;"><span class="badge badge-manual">Ручной</span></td>
        <td style="width: 25%;"><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>
  
  <div class="category-header">🏷️ Переименование документов</div>
  <table class="scripts-table">
    <tbody>
    {% for s in scripts %}
      {% if 'rename' in s.key %}
      <tr>
        <td style="width: 35%;">
          <div class="script-name">{{ s.title }}</div>
          <div class="script-description">{{ s.description }}</div>
        </td>
        <td style="width: 25%;"><span class="badge badge-ondemand">При покупке портфеля</span></td>
        <td style="width: 15%;"><span class="badge badge-manual">Ручной</span></td>
        <td style="width: 25%;"><a href="/run/{{ s.key }}" class="run-btn">▶ Запустить</a></td>
      </tr>
      {% endif %}
    {% endfor %}
    </tbody>
  </table>
</div>

<div class="history-container">
  <h3 style="margin-top: 0;">История запусков</h3>
  <table class="history-table">
    <tr>
      <th>Время</th>
      <th>Скрипт</th>
      <th>Пользователь</th>
      {% if user.role in ('admin', 'manager') %}<th>Компания</th>{% endif %}
      <th>Статус</th>
      <th>Логи</th>
      <th>Результат</th>
      <th>Действия</th>
    </tr>
    {% for j in jobs %}
    <tr>
      <td>{{ j.created_at }}</td>
      <td>{{ j.script_key }}</td>
      <td>{{ j.username }}</td>
      {% if user.role in ('admin', 'manager') %}<td>{{ j.company_name }}</td>{% endif %}
      <td>
        {% if j.status == 'success' %}
          <span class="status-success">✅ Успешно</span>
        {% elif j.status == 'error' %}
          <span class="status-error">❌ Ошибка</span>
        {% elif j.status == 'running' %}
          <span class="status-running">⏳ Выполняется</span>
        {% elif j.status == 'scheduled' %}
          <span class="status-queued">🕒 Запланировано на {{ j.scheduled_display }}</span>
        {% elif j.status == 'cancelled' %}
          <span class="status-queued">🚫 Отменено</span>
        {% else %}
          <span class="status-queued">🕐 В очереди</span>
        {% endif %}
      </td>
      <td><a href="/logs/{{ j.id }}">Открыть</a></td>
      <td>{% if j.result_zip %}<a href="/download/{{ j.id }}">Скачать</a>{% else %}—{% endif %}</td>
      <td>
        {% if j.status in ['running', 'queued', 'scheduled'] and (user.role == 'admin' or j.username == user.username) %}
          <a href="/stop/{{ j.id }}" onclick="return confirm('Остановить задачу?')" style="color:#dc3545;">⛔ Остановить</a>
        {% else %}—{% endif %}
      </td>
    </tr>
    {% endfor %}
  </table>
</div>

</body></html>
"""

AISOIP_HTML = """
<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>АИС ОИП — поиск</title>
<style>
:root{--ink:#0c4964;--line:#d8dee2;--soft:#f3f4f5;--muted:#72777b;--blue:#0b4f6c;--green:#4c7c58;--shadow:0 2px 7px rgba(0,0,0,.11)}
*{box-sizing:border-box} body{margin:0;background:#f5f7fa;color:#293238;font-family:Arial,"Segoe UI",sans-serif;font-size:14px}
.header{display:flex;justify-content:space-between;align-items:center;background:white;padding:16px 24px;box-shadow:0 2px 4px rgba(0,0,0,.1);margin-bottom:20px}
.header a.back{color:#0b4f6c;text-decoration:none;font-weight:700;margin-right:18px}
.user-info{font-size:13px;color:#555}
.logout-btn{padding:6px 14px;background:#dc3545;color:white;text-decoration:none;border-radius:4px;margin-left:10px;font-size:12px}
.page{padding:0 24px 60px;max-width:1720px;margin:auto}
.company-pick{background:white;max-width:520px;margin:60px auto;padding:30px;border-radius:8px;box-shadow:var(--shadow);text-align:center}
.company-pick select{height:40px;width:100%;margin:14px 0;font-size:15px}
.company-pick button{height:40px;padding:0 22px;background:#0b4f6c;color:white;border:none;border-radius:4px;cursor:pointer;font-weight:700}
.company-switch{margin-bottom:14px;font-size:13px;color:#58656b}
.company-switch select{margin-left:8px;height:30px}
.filters{display:grid;grid-template-columns:350px 350px 1fr;gap:12px 38px;align-items:end;max-width:1190px}.dates{grid-column:1/3;display:flex;align-items:center;gap:14px;margin-bottom:20px}.dates .label{font-size:18px;margin-right:2px}.datebox,.field,.select{height:40px;border:1px solid #adb4b8;border-radius:3px;background:#fff;padding:0 12px;font-size:15px;color:#32393d}.datebox{width:205px}.field{width:100%}.floating{position:relative}.floating label{position:absolute;top:-8px;left:11px;background:#f5f7fa;padding:0 3px;font-size:11px;color:#8a9195}.select{border:none;border-bottom:1px solid #9fa6aa;border-radius:0;color:#6d7478;padding-left:0}.actions{display:flex;gap:18px;margin-top:26px}.btn{border:1px solid #687278;background:white;color:#28343a;font-weight:700;letter-spacing:.8px;border-radius:4px;height:40px;padding:0 18px;cursor:pointer}.btn.primary{background:#0b4f6c;color:white;border-color:#0b4f6c;box-shadow:var(--shadow)}.btn.sync{margin-left:auto;background:#0b4f6c;color:#fff;border-color:#0b4f6c}.btn.import{background:white;color:#0b4f6c;border-color:#0b4f6c}.btn:hover{filter:brightness(.98)}
.chips{display:flex;gap:8px;flex-wrap:wrap;margin:26px 0 20px}.chip{background:#e5e7e8;border-radius:18px;padding:8px 34px 8px 14px;position:relative;font-size:12px;line-height:1.25}.chip b{display:block;font-size:10px;color:#555f63;margin-bottom:3px}.chip .x{position:absolute;right:8px;top:8px;width:15px;height:15px;border-radius:50%;background:#6f7477;color:white;text-align:center;line-height:15px;font-size:11px;cursor:pointer}
.results-head{display:flex;justify-content:space-between;align-items:center;margin:6px 0 12px}.found{font-size:13px;color:#606b70}.export{border:none;background:none;color:#2977a5;cursor:pointer;font-weight:700}.cards{display:flex;flex-direction:column;gap:13px}.card{border:1px solid #d7dcdf;border-radius:5px;box-shadow:0 1px 4px rgba(0,0,0,.10);display:grid;grid-template-columns:1.15fr 1.15fr 1.12fr 1.1fr;min-height:170px;background:white}.cell{padding:15px 18px;border-right:1px solid #d5dadd}.cell:last-child{border-right:none}.label2{font-size:13px;margin-bottom:8px}.value{color:#1d5d7d;font-size:14px;margin-bottom:16px;line-height:1.25}.linklike{color:#2e91c9;text-decoration:underline}.status{color:#4e875a}.sum{color:#1d5d7d}.muted{color:#758087}.empty{padding:50px;text-align:center;color:#8b9498;border:1px dashed #ccd2d5;border-radius:5px;background:white}.toast{position:fixed;right:26px;bottom:26px;background:#25343b;color:#fff;padding:14px 18px;border-radius:5px;box-shadow:var(--shadow);display:none;max-width:440px}.spinner{display:none;margin-left:8px}.syncbox{margin-top:12px;display:none;max-width:760px;border:1px solid #d7dcdf;background:#f8fafb;border-radius:4px;padding:10px 12px;color:#58656b;font-size:12px}.syncrow{display:flex;justify-content:space-between;gap:14px;align-items:center}.progress{height:5px;background:#e3e8ea;border-radius:5px;overflow:hidden;margin-top:8px}.progress>div{height:100%;background:#0b4f6c;width:0;transition:width .25s}.syncerr{color:#a33}
.stats{font-size:12px;color:#8a9295;margin-bottom:14px}
@media(max-width:950px){.filters{grid-template-columns:1fr}.dates{grid-column:1;flex-wrap:wrap}.card{grid-template-columns:1fr}.cell{border-right:none;border-bottom:1px solid #d5dadd}.cell:last-child{border-bottom:none}.actions{flex-wrap:wrap}}
</style>
</head>
<body>
<div class="header">
  <div><a href="/" class="back">← Runner</a><strong>АИС ОИП — внутренняя база</strong></div>
  <div class="user-info">
    {{ user.full_name }} ({{ user.role }})
    <a href="/logout" class="logout-btn">Выход</a>
  </div>
</div>
<div class="page">
{% if not selected_company_id %}
  <div class="company-pick">
    <h3>Выберите компанию</h3>
    <p class="muted">Поиск и синхронизация АИС ОИП ведутся отдельно по каждой компании.</p>
    <select id="companyPick">
      {% for c in companies %}<option value="{{ c.id }}">{{ c.name }}</option>{% endfor %}
    </select><br>
    <button onclick="location.href='/aisoip?company_id='+document.getElementById('companyPick').value">Открыть</button>
  </div>
{% else %}
  {% if user.role in ('admin', 'manager') %}
  <div class="company-switch">Компания:
    <select onchange="location.href='/aisoip?company_id='+this.value">
      {% for c in companies %}<option value="{{ c.id }}" {% if c.id == selected_company_id %}selected{% endif %}>{{ c.name }}</option>{% endfor %}
    </select>
  </div>
  {% endif %}
  <div class="stats" id="stats"></div>
  <div class="filters" id="filters">
    <div class="dates"><span class="label">Дата возбуждения</span><span>с</span><input class="datebox" id="date_from" type="date"><span>по</span><input class="datebox" id="date_to" type="date"><span class="muted">(необязательно)</span></div>
    <div class="floating"><label>ИИН должника</label><input class="field" id="iin" autocomplete="off" placeholder=""></div>
    <select class="select" id="executor"><option value="">Судебный исполнитель</option></select><div></div>
    <div class="floating"><label>БИН должника</label><input class="field" id="bin"></div>
    <select class="select" id="status"><option value="">Статус</option></select><div></div>
    <div class="floating"><label>ФИО/Наименование ЮЛ</label><input class="field" id="debtor"></div>
    <input class="select" id="issuer" placeholder="Регион / орган, выдавший документ"><div></div>
    <div class="floating"><label>Номер исполнительного документа</label><input class="field" id="document_no"></div>
    <div class="floating"><label>Номер исполнительного производства</label><input class="field" id="production_no"></div><div></div>
  </div>
  <div class="actions"><button class="btn primary" id="searchBtn">ПОИСК <span class="spinner" id="spinner">…</span></button><button class="btn" id="resetBtn">СБРОСИТЬ</button><button class="btn sync" id="syncBtn">↻ ОБНОВИТЬ С АИС ОИП</button><button class="btn import" id="importBtn">⇩ ИМПОРТ EXCEL</button><input type="file" id="fileInput" accept=".xlsx" hidden></div><div class="syncbox" id="syncbox"><div class="syncrow"><b id="syncstage">Подготовка…</b><span id="synccount"></span></div><div id="syncmsg"></div><div class="progress"><div id="syncprogress"></div></div></div>
  <div class="chips" id="chips"></div>
  <div class="results-head"><div class="found" id="found">Введите ИИН и нажмите «ПОИСК»</div><button class="export" id="exportBtn">▣ ЭКСПОРТ РЕЗУЛЬТАТОВ</button></div>
  <div class="cards" id="cards"><div class="empty">Поиск выполняется по внутренней базе без ограничения периода в 3 месяца.</div></div>
{% endif %}
</div>
<div class="toast" id="toast"></div>
<script>
const COMPANY_ID = {{ selected_company_id or 'null' }};
const ids=['date_from','date_to','iin','bin','debtor','document_no','production_no','executor','status','issuer'];
const names={date_from:'Дата возбуждения с',date_to:'Дата возбуждения по',iin:'ИИН должника',bin:'БИН должника',debtor:'ФИО/Наименование ЮЛ',document_no:'Номер исполнительного документа',production_no:'Исполнительное производство',executor:'Судебный исполнитель',status:'Статус',issuer:'Регион / орган'};
const $=id=>document.getElementById(id);
function toast(t){$('toast').textContent=t;$('toast').style.display='block';setTimeout(()=>$('toast').style.display='none',4200)}
function esc(s){return String(s??'').replace(/[&<>'"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;',"'":'&#39;','"':'&quot;'}[c]))}
function fmtDate(s){if(!s)return '—'; const [y,m,d]=s.split('-'); return d&&m&&y?`${d}.${m}.${y}`:s}
function fmtMoney(n){if(n===null||n===undefined||n==='')return '—';return new Intl.NumberFormat('ru-RU',{minimumFractionDigits:0,maximumFractionDigits:2}).format(Number(n))}
function params(){const p=new URLSearchParams();ids.forEach(id=>{const v=$(id).value.trim();if(v)p.set(id,v)});if(COMPANY_ID)p.set('company_id',COMPANY_ID);return p}
function renderChips(){const c=$('chips');c.innerHTML='';ids.forEach(id=>{const v=$(id).value.trim();if(!v)return;const el=document.createElement('div');el.className='chip';el.innerHTML=`<b>${esc(names[id])}</b>${esc(id.startsWith('date_')?fmtDate(v):v)}<span class="x">×</span>`;el.querySelector('.x').onclick=()=>{$(id).value='';renderChips();search()};c.appendChild(el)})}
async function loadOptions(){if(!COMPANY_ID)return;const r=await fetch('/api/aisoip/options?company_id='+COMPANY_ID);const d=await r.json();d.executors.forEach(v=>$('executor').insertAdjacentHTML('beforeend',`<option>${esc(v)}</option>`));d.statuses.forEach(v=>$('status').insertAdjacentHTML('beforeend',`<option>${esc(v)}</option>`))}
async function loadStats(){if(!COMPANY_ID)return;const r=await fetch('/api/aisoip/stats?company_id='+COMPANY_ID);const d=await r.json();$('stats').textContent=`В базе: ${d.total.toLocaleString('ru-RU')} производств · ${d.debtors.toLocaleString('ru-RU')} ИИН/БИН`}
async function search(){if(!COMPANY_ID)return;renderChips();$('spinner').style.display='inline';try{const r=await fetch('/api/aisoip/search?'+params().toString());const d=await r.json();$('found').textContent=`Найдено: ${d.count.toLocaleString('ru-RU')}`;const c=$('cards');c.innerHTML='';if(!d.items.length){c.innerHTML='<div class="empty">По заданным параметрам исполнительные производства не найдены.</div>';return}d.items.forEach(x=>{const card=document.createElement('div');card.className='card';card.innerHTML=`
<div class="cell"><div class="label2">Исполнительный документ:</div><div class="value">${esc(x.document_no||'—')} &nbsp; от ${fmtDate(x.document_date)}</div><div class="label2">Орган, выдавший исполнительный документ:</div><div class="value">${esc(x.issuer||'—')}</div></div>
<div class="cell"><div class="label2">Исполнительное производство:</div><div class="value linklike">${esc(x.production_no)}</div><div class="label2">Дата возбуждения:</div><div class="value">${fmtDate(x.start_date)}</div></div>
<div class="cell"><div class="label2">Судебный исполнитель:</div><div class="value">${esc(x.executor||'—')}</div><div class="label2">Статус:</div><div class="value status">${esc(x.status||'—')}</div></div>
<div class="cell"><div class="label2">Должник:</div><div class="value">${esc(x.debtor||'—')}</div><div class="label2">ИИН/БИН:</div><div class="value">${esc(x.iin||'—')}</div><div class="label2">Тип взыскания:</div><div class="value">${esc(x.collection_type||'—')}</div><div class="label2">Сумма взыскания (тенге):</div><div class="value sum">${fmtMoney(x.amount)}</div></div>`;c.appendChild(card)})}catch(e){toast('Ошибка поиска: '+e.message)}finally{$('spinner').style.display='none'}}
if(COMPANY_ID){
$('searchBtn').onclick=search;
$('resetBtn').onclick=()=>{ids.forEach(id=>$(id).value='');renderChips();$('found').textContent='Введите ИИН и нажмите «ПОИСК»';$('cards').innerHTML='<div class="empty">Поиск выполняется по внутренней базе без ограничения периода в 3 месяца.</div>'};
ids.forEach(id=>$(id).addEventListener('keydown',e=>{if(e.key==='Enter')search()}));
let syncTimer=null;
async function pollSync(){
  try{
    const r=await fetch('/api/aisoip/sync/status?company_id='+COMPANY_ID); const d=await r.json();
    $('syncbox').style.display=(d.running||d.started_at)?'block':'none';
    $('syncstage').textContent=d.stage||'Обновление';
    $('syncmsg').textContent=d.error ? ('Ошибка: '+d.error) : (d.message||'');
    $('syncmsg').className=d.error?'syncerr':'';
    const pct=d.pages_total?Math.round((d.pages_done/d.pages_total)*100):(d.total_elements?5:0);
    $('syncprogress').style.width=pct+'%';
    $('synccount').textContent=`${d.pages_done||0}/${d.pages_total||0} частей · строк: ${(d.rows_imported||0).toLocaleString('ru-RU')}${d.total_elements?' из '+d.total_elements.toLocaleString('ru-RU'):''}`;
    $('syncBtn').disabled=!!d.running;
    $('syncBtn').textContent=d.running?'↻ ОБНОВЛЕНИЕ…':'↻ ОБНОВИТЬ С АИС ОИП';
    if(d.running){syncTimer=setTimeout(pollSync,1500)}
    else if(d.started_at){await loadStats(); await loadOptionsFresh(); if(!d.error) toast('Данные АИС ОИП обновлены')}
  }catch(e){toast('Ошибка статуса обновления: '+e.message)}
}
async function loadOptionsFresh(){
  const ex=$('executor'), st=$('status'); const exv=ex.value, stv=st.value;
  ex.innerHTML='<option value="">Судебный исполнитель</option>'; st.innerHTML='<option value="">Статус</option>';
  const r=await fetch('/api/aisoip/options?company_id='+COMPANY_ID); const d=await r.json();
  d.executors.forEach(v=>ex.insertAdjacentHTML('beforeend',`<option>${esc(v)}</option>`));
  d.statuses.forEach(v=>st.insertAdjacentHTML('beforeend',`<option>${esc(v)}</option>`));
  ex.value=exv; st.value=stv;
}
$('syncBtn').onclick=async()=>{
  if(!confirm('Обновить внутреннюю базу данными из АИС ОИП?\\n\\nБудет выполнена API-выгрузка АИС ОИП за весь период с 01.01.2023 по сегодня.\\n\\nОкно Chrome не откроется — браузер работает скрыто (headless).'))return;
  $('syncBtn').disabled=true; $('syncbox').style.display='block'; $('syncstage').textContent='Запуск…'; $('syncmsg').textContent='';
  try{
    const r=await fetch('/api/aisoip/sync/start?company_id='+COMPANY_ID,{method:'POST'}); const d=await r.json();
    if(!r.ok&&!d.ok) throw new Error(d.message||'Не удалось запустить обновление');
    toast(d.message||'Обновление запущено'); if(syncTimer)clearTimeout(syncTimer); pollSync();
  }catch(e){$('syncBtn').disabled=false;toast('Ошибка запуска: '+e.message)}
};
$('importBtn').onclick=()=>$('fileInput').click();
$('fileInput').onchange=async()=>{const f=$('fileInput').files[0];if(!f)return;const fd=new FormData();fd.append('file',f);fd.append('company_id',COMPANY_ID);$('importBtn').disabled=true;$('importBtn').textContent='ИМПОРТ…';try{const r=await fetch('/api/aisoip/import',{method:'POST',body:fd});const d=await r.json();if(!d.ok)throw new Error(d.error);toast(`Импорт завершён. Обработано строк: ${d.processed.toLocaleString('ru-RU')}`);await loadStats();location.reload()}catch(e){toast('Ошибка импорта: '+e.message)}finally{$('importBtn').disabled=false;$('importBtn').textContent='⇩ ИМПОРТ EXCEL';$('fileInput').value=''}};
$('exportBtn').onclick=()=>{location.href='/api/aisoip/export?'+params().toString()};
loadOptions();loadStats();pollSync();
}
</script>
</body></html>
"""

USER_FORM_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Пользователь</title>
<style>
body { font-family: Arial; max-width: 800px; margin: 40px auto; }
.form-group { margin: 15px 0; }
label { display: block; margin-bottom: 5px; font-weight: bold; }
input, select { padding: 10px; width: 100%; box-sizing: border-box; }
button { padding: 12px 24px; background: #007bff; color: white; border: none; border-radius: 4px; cursor: pointer; }
button:hover { background: #0056b3; }

.scripts-section { margin: 20px 0; padding: 20px; background: #f8f9fa; border-radius: 6px; }
.script-checkbox { margin: 10px 0; }
.script-checkbox label { font-weight: normal; display: flex; align-items: center; }
.script-checkbox input { width: auto; margin-right: 10px; }
.select-all-btn { padding: 8px 16px; background: #28a745; color: white; border: none; border-radius: 4px; cursor: pointer; margin-bottom: 10px; }
.deselect-all-btn { padding: 8px 16px; background: #6c757d; color: white; border: none; border-radius: 4px; cursor: pointer; margin-bottom: 10px; margin-left: 10px; }
</style>
</head>
<body>
<a href="/admin">← Назад</a>
<h2>{% if user %}Редактировать{% else %}Добавить{% endif %} пользователя</h2>
<form method="post">
  <div class="form-group">
    <label>Логин:</label>
    <input name="username" value="{{ user.username if user else '' }}" required>
  </div>
  <div class="form-group">
    <label>ФИО:</label>
    <input name="full_name" value="{{ user.full_name if user else '' }}" required>
  </div>
  <div class="form-group">
    <label>Пароль {% if user %}(оставьте пустым чтобы не менять){% endif %}:</label>
    <input name="password" type="password" {% if not user %}required{% endif %}>
  </div>
  <div class="form-group">
    <label>Роль:</label>
    <select name="role" id="role" onchange="toggleFields()">
      <option value="operator" {% if user and user.role=='operator' %}selected{% endif %}>Оператор</option>
      <option value="manager" {% if user and user.role=='manager' %}selected{% endif %}>Менеджер</option>
      <option value="admin" {% if user and user.role=='admin' %}selected{% endif %}>Администратор</option>
    </select>
  </div>
  <div class="form-group" id="company-group">
    <label>Компания:</label>
    <select name="company_id">
      <option value="">-- Выберите компанию --</option>
      {% for c in companies %}
        <option value="{{ c.id }}" {% if user and user.company_id == c.id %}selected{% endif %}>{{ c.name }}</option>
      {% endfor %}
    </select>
  </div>

  <div class="scripts-section" id="companies-group">
    <h3>Компании</h3>
    <p style="color: #666; font-size: 14px;">Менеджер может переключаться между отмеченными компаниями</p>
    {% for c in companies %}
      <div class="script-checkbox">
        <label>
          <input type="checkbox" name="companies" value="{{ c.id }}"
                 {% if user and c.id in user_company_ids %}checked{% endif %}>
          {{ c.name }}
        </label>
      </div>
    {% endfor %}
  </div>

  <div class="scripts-section" id="scripts-section">
    <h3>Разрешённые скрипты</h3>
    <p style="color: #666; font-size: 14px;">Выберите какие скрипты будут доступны пользователю</p>
    
    <button type="button" class="select-all-btn" onclick="selectAll()">✅ Выбрать все</button>
    <button type="button" class="deselect-all-btn" onclick="deselectAll()">❌ Снять все</button>
    
    {% for script_key, script in scripts.items() %}
      <div class="script-checkbox">
        <label>
          <input type="checkbox" name="scripts" value="{{ script_key }}" 
                 {% if user and script_key in user_scripts %}checked{% endif %}>
          <strong>{{ script.title }}</strong> — {{ script.description }}
        </label>
      </div>
    {% endfor %}
  </div>
  
  <div class="form-group">
    <label>
      <input name="is_active" type="checkbox" {% if not user or user.is_active %}checked{% endif %}> Активен
    </label>
  </div>
  <button type="submit">💾 Сохранить</button>
</form>

<script>
function toggleFields() {
  const role = document.getElementById('role').value;
  const companyGroup = document.getElementById('company-group');
  const companiesGroup = document.getElementById('companies-group');
  const scriptsSection = document.getElementById('scripts-section');

  companyGroup.style.display = (role === 'operator') ? 'block' : 'none';
  companiesGroup.style.display = (role === 'manager') ? 'block' : 'none';
  scriptsSection.style.display = (role === 'admin') ? 'none' : 'block';
}

function selectAll() {
  document.querySelectorAll('input[name="scripts"]').forEach(cb => cb.checked = true);
}

function deselectAll() {
  document.querySelectorAll('input[name="scripts"]').forEach(cb => cb.checked = false);
}

toggleFields();
</script>
</body></html>
"""

ADMIN_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Админка</title>
<style>
body { font-family: Arial; max-width: 1200px; margin: 20px auto; }
h2 { border-bottom: 2px solid #28a745; padding-bottom: 10px; }
table { width: 100%; border-collapse: collapse; margin: 20px 0; }
th, td { padding: 10px; text-align: left; border: 1px solid #ddd; }
th { background: #28a745; color: white; }
.btn { padding: 6px 12px; text-decoration: none; border-radius: 4px; color: white; display: inline-block; margin: 2px; }
.btn-edit { background: #007bff; }
.btn-delete { background: #dc3545; }
.btn-add { background: #28a745; padding: 10px 20px; }
.company-badge { display: inline-block; padding: 4px 10px; background: #6c757d; color: white; border-radius: 8px; font-size: 11px; }
.scripts-count { color: #666; font-size: 12px; }
</style>
</head>
<body>
<a href="/">← Назад</a>
<h2>⚙️ Управление пользователями</h2>

<a href="/admin/user/new" class="btn btn-add">➕ Добавить пользователя</a>
<a href="/admin/schedule" class="btn-add" style="text-decoration:none;">⏰ Расписание автозапуска</a>
<a href="/admin/mailings" class="btn-add" style="text-decoration:none; background:#e67e22;">📨 Рассылки по отчётам</a>
<a href="/monitoring" class="btn-add" style="text-decoration:none; background:#17a2b8;">📊 Мониторинг</a>
<table>
  <tr><th>Логин</th><th>ФИО</th><th>Компания</th><th>Роль</th><th>Скрипты</th><th>Статус</th><th>Действия</th></tr>
  {% for u in users %}
  <tr>
    <td>{{ u.username }}</td>
    <td>{{ u.full_name }}</td>
    <td>
      {% if u.company %}
        <span class="company-badge">{{ u.company.name }}</span>
      {% else %}
        <span class="company-badge">Все</span>
      {% endif %}
    </td>
    <td>{{ u.role }}</td>
    <td>
      {% if u.role == 'admin' %}
        <span class="scripts-count">Все ({{ total_scripts }})</span>
      {% else %}
        <span class="scripts-count">{{ user_scripts_count.get(u.id, 0) }} из {{ total_scripts }}</span>
      {% endif %}
    </td>
    <td>{% if u.is_active %}Активен{% else %}Заблокирован{% endif %}</td>
    <td>
      <a href="/admin/user/{{ u.id }}/edit" class="btn btn-edit">Изменить</a>
      {% if u.username != 'admin' %}
      <a href="/admin/user/{{ u.id }}/delete" class="btn btn-delete" onclick="return confirm('Удалить пользователя?')">Удалить</a>
      {% endif %}
    </td>
  </tr>
  {% endfor %}
</table>
</body></html>
"""

# ─────────────────────────────────────────────────────────────
# RUN_HTML — поддержка type="file" в параметрах
# Форма использует enctype="multipart/form-data" только если
# среди параметров скрипта есть хотя бы один файловый.
# ─────────────────────────────────────────────────────────────
RUN_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Запуск</title>
<style>
body { font-family: Arial; max-width: 800px; margin: 40px auto; padding: 0 20px; }
h2 { border-bottom: 2px solid #007bff; padding-bottom: 10px; }
.form-group { margin: 20px 0; }
label { display: block; margin-bottom: 8px; font-weight: bold; }
input[type=text], input[type=date], input[type=password] {
  padding: 10px; width: 100%; box-sizing: border-box;
  border: 1px solid #ddd; border-radius: 4px;
}
.file-upload-area {
  border: 2px dashed #007bff; border-radius: 8px;
  padding: 30px; text-align: center; background: #f0f7ff;
  cursor: pointer; transition: background 0.2s;
}
.file-upload-area:hover { background: #dceeff; }
.file-upload-area input[type=file] { display: none; }
.file-upload-label { color: #007bff; font-size: 15px; cursor: pointer; }
.file-name { margin-top: 10px; color: #28a745; font-size: 13px; font-weight: bold; }
button {
  padding: 12px 24px; background: #007bff; color: white;
  border: none; border-radius: 4px; cursor: pointer; font-size: 16px;
}
button:hover { background: #0056b3; }
.back-link { display: inline-block; margin-bottom: 20px; color: #007bff; text-decoration: none; }
.back-link:hover { text-decoration: underline; }
</style>
</head>
<body>
  <a href="/" class="back-link">← Назад</a>
  <h2>▶ Запуск: {{ script.title }}</h2>
  <p>{{ script.description }}</p>

  {# Если есть файловый параметр — нужен multipart #}
  {% set has_file = script.params | selectattr('type', 'eq', 'file') | list | length > 0 %}

  <form method="post" {% if has_file %}enctype="multipart/form-data"{% endif %}>
    {% for p in script.params %}
      <div class="form-group">
        <label>{{ p.label }}</label>

        {% if p.type == 'file' %}
          {# Красивая зона drag-and-drop #}
          <div class="file-upload-area" onclick="document.getElementById('file_{{ p.name }}').click()">
            <input
              type="file"
              id="file_{{ p.name }}"
              name="{{ p.name }}"
              accept="{{ p.accept if p.accept is defined else '*' }}"
              onchange="showFileName(this, 'fname_{{ p.name }}')"
            >
            <div class="file-upload-label">
              📂 Нажмите чтобы выбрать файл<br>
              <small style="color:#666;">{{ p.accept if p.accept is defined else 'Любой файл' }}</small>
            </div>
            <div class="file-name" id="fname_{{ p.name }}"></div>
          </div>

        {% elif p.type == 'date' %}
          <input name="{{ p.name }}" type="date" value="{{ today }}">

        {% elif p.type == 'password' %}
          <input name="{{ p.name }}" type="password">

        {% elif p.name == 'company_id' %}
          {# Компания больше не выбирается здесь вручную — она определяется
             автоматически: из вкладки "Компания:" в шапке сайта (admin/manager)
             или из привязки пользователя (operator). #}
          {% if active_company %}
            <input type="hidden" name="company_id" value="{{ active_company.id }}">
            <div style="padding:10px 12px;background:#eef7ff;border:1px solid #bcdfff;border-radius:4px;">
              🏢 {{ active_company.name }}
              <small style="color:#666;">(определена автоматически)</small>
            </div>
          {% else %}
            <div style="padding:10px 12px;background:#fff3cd;border:1px solid #ffe08a;border-radius:4px;color:#7a5b00;">
              ⚠ Компания не определена автоматически. Если вы администратор или
              менеджер — вернитесь на <a href="/">главную</a> и выберите её во
              вкладках "Компания:" над списком скриптов. Если это оператор без
              привязанной компании — обратитесь к администратору.
            </div>
          {% endif %}

        {% elif p.type == 'select' %}
          <select name="{{ p.name }}">
            {% for opt in p.options %}
              <option value="{{ opt.value }}" {% if p.default is defined and opt.value == p.default %}selected{% endif %}>{{ opt.label }}</option>
            {% endfor %}
          </select>

        {% elif p.type == 'folder' %}
          {% if folder_options %}
            <select name="{{ p.name }}">
              <option value="">— выберите папку —</option>
              {% for opt in folder_options %}
                <option value="{{ opt.value }}">{{ opt.label }}</option>
              {% endfor %}
            </select>
            <small style="display:block;color:#666;margin-top:6px;">
              Папки на сервере{% if rename_root %} ({{ rename_root }}){% endif %}, новые — сверху.
              Нет нужной? <a href="javascript:location.reload()">обновить список</a>.
            </small>
          {% else %}
            <div style="padding:10px 12px;background:#fff3cd;border:1px solid #ffe08a;border-radius:4px;color:#7a5b00;">
              ⚠ Список папок с сервера недоступен{% if rename_root %} (корень: {{ rename_root }}){% endif %}.
              Задайте <code>RENAME_ROOT</code> в <code>scripts/config.py</code> на сервере
              или впишите полный путь вручную ниже.
            </div>
          {% endif %}
          <input name="{{ p.name }}__manual" type="text" style="margin-top:8px;"
                 placeholder="…или вручную: имя подпапки, либо полный путь как его видит сервер">

        {% else %}
          <input name="{{ p.name }}" type="text">
        {% endif %}

      </div>
    {% endfor %}
    {% if script.allow_scheduled_start %}
      <div class="form-group">
        <label>⏰ Время запуска (необязательно — если не указать, запустится сразу):</label>
        <input type="datetime-local" name="scheduled_at" min="{{ min_datetime }}">
        <small style="color:#666;">Оставьте пустым, чтобы запустить прямо сейчас</small>
      </div>
    {% endif %}
    <button type="submit" {% if needs_company and not active_company %}disabled{% endif %}>🚀 Запустить</button>
  </form>

<script>
function showFileName(input, labelId) {
  const label = document.getElementById(labelId);
  if (input.files && input.files.length > 0) {
    label.textContent = '✅ ' + input.files[0].name;
  } else {
    label.textContent = '';
  }
}
</script>
</body></html>
"""

LOGS_HTML = """
<!doctype html><html>
<head>
  <meta charset="utf-8">
  <title>Логи</title>
  <style>
    body { font-family: Arial; max-width: 1000px; margin: 24px auto; }
    pre { background: #1e1e1e; color: #d4d4d4; padding: 16px; border-radius: 6px; white-space: pre-wrap; font-size: 13px; }
    .status-success { color: green; font-weight: bold; }
    .status-error { color: red; font-weight: bold; }
    .status-running { color: orange; font-weight: bold; }
    .status-cancelled { color: #6c757d; font-weight: bold; }
    .back-link { display: inline-block; margin-bottom: 20px; color: #007bff; text-decoration: none; }
  </style>
  {% if job.status == 'running' %}
    <meta http-equiv="refresh" content="3">
  {% endif %}
</head>
<body>
  <a href="/" class="back-link">← Назад</a>
  <h2>📄 Логи: {{ job.id }}</h2>
  <p>
    Скрипт: <b>{{ job.script_key }}</b> |
    Статус: <span class="status-{{ job.status }}">{{ job.status }}</span>
    {% if job.status == 'running' %}
      &nbsp; <small>(обновление каждые 3 сек)</small>
    {% endif %}
  </p>
  {% if job.started_at %}
    <p>🕐 Начало: {{ job.started_at }} | 🏁 Конец: {{ job.finished_at or '...' }}</p>
  {% endif %}
  {% if job.status in ['running', 'queued', 'scheduled'] %}
    <p>
      <a href="/stop/{{ job.id }}" onclick="return confirm('Остановить задачу? Процесс будет принудительно завершён.')"
         style="padding:8px 16px; background:#dc3545; color:white; text-decoration:none; border-radius:4px; display:inline-block;">
        ⛔ Остановить
      </a>
    </p>
  {% endif %}
  <pre>{{ log_text }}</pre>
  {% if job.result_zip %}
    <a href="/download/{{ job.id }}" style="padding:10px 16px; background:green; color:white; text-decoration:none; border-radius:4px;">
      📦 Скачать результат (zip)
    </a>
  {% endif %}
</body></html>
"""

# ═══════════════════════════════════════════════════════════════
# РОУТЫ
# ═══════════════════════════════════════════════════════════════

from database import init_db, get_db, User, Company, get_user_scripts, set_user_scripts
from scripts_registry import SCRIPTS


def visible_scripts() -> dict:
    """SCRIPTS без служебных (hidden) — для списка на главной и чек-листа прав."""
    return {k: v for k, v in SCRIPTS.items() if not getattr(v, "hidden", False)}


def _resolve_active_company(current_user: User, session_id: Optional[str], requested: Optional[str], db: Session):
    """
    Определяет компанию, с которой автоматически запустятся скрипты.

    - operator: всегда его единственная привязанная компания — выбирать нечего.
    - admin/manager: компания, выбранная в шапке сайта ("Компания:" над списком
      скриптов). Выбор запоминается в сессии (см. auth.set_active_company_id),
      поэтому его не нужно повторять на странице запуска каждого скрипта —
      именно эта компания и уходит в --company_id при старте скрипта.

    `requested` — значение ?company= с текущего запроса (если пользователь только
    что кликнул по вкладке в шапке); "all" сбрасывает выбор. Если оно не задано,
    используется то, что уже сохранено в сессии.

    Возвращает (company_id: int|None, company: Company|None).
    """
    if current_user.role == "operator":
        company_id = current_user.company_id
    else:
        allowed_company_ids = get_user_allowed_company_ids(current_user)  # None = все компании

        if requested == "all":
            set_active_company_id(session_id, None)
        elif requested is not None:
            try:
                requested_int = int(requested)
            except ValueError:
                requested_int = None
            if requested_int is not None and (allowed_company_ids is None or requested_int in allowed_company_ids):
                set_active_company_id(session_id, requested_int)

        company_id = get_active_company_id(session_id)
        # Доступ к ранее выбранной компании могли отозвать — тогда сбрасываем.
        if company_id is not None and allowed_company_ids is not None and company_id not in allowed_company_ids:
            company_id = None
            set_active_company_id(session_id, None)

    company = db.query(Company).filter(Company.id == company_id).first() if company_id else None
    return company_id, company


@app.get("/", response_class=HTMLResponse)
def index(
    company: Optional[str] = None,
    session_id: Optional[str] = Cookie(None, alias="session"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    current_company = company or "all"

    # None = все компании (admin); [] / [id, ...] = ограниченный список (manager/operator)
    allowed_company_ids = get_user_allowed_company_ids(current_user)
    companies = (
        db.query(Company).all() if allowed_company_ids is None
        else db.query(Company).filter(Company.id.in_(allowed_company_ids)).all()
    )

    # Клик по вкладке "Компания:" в шапке одновременно фильтрует список задач
    # ниже (через current_company) и запоминается как компания, с которой
    # автоматически запустятся скрипты (см. _resolve_active_company).
    active_company_id, active_company = _resolve_active_company(current_user, session_id, company, db)

    if current_user.role == "admin":
        available_scripts = list(visible_scripts().values())
    else:
        allowed_keys = get_user_scripts(current_user.id)
        available_scripts = [s for s in visible_scripts().values() if s.key in allowed_keys]

    # Автозапуски (по расписанию) сюда не попадают — для них отдельная
    # страница /monitoring с фильтрами по периоду/скрипту.
    all_jobs = sorted(
        (j for j in JOBS.values() if j.get("source") != "scheduler"),
        key=lambda x: x["created_ts"], reverse=True,
    )

    if allowed_company_ids is None:
        # admin — весь список задач, опционально сужаем по выбранной вкладке
        if current_company != "all":
            filtered_jobs = [j for j in all_jobs if j.get("company_id") == int(current_company)]
        else:
            filtered_jobs = all_jobs
    elif current_user.role == "manager":
        # manager — только свои компании, опционально сужаем по выбранной вкладке
        if current_company != "all" and int(current_company) in allowed_company_ids:
            filtered_jobs = [j for j in all_jobs if j.get("company_id") == int(current_company)]
        else:
            filtered_jobs = [j for j in all_jobs if j.get("company_id") in allowed_company_ids]
    else:
        filtered_jobs = [j for j in all_jobs if j.get("company_id") == current_user.company_id]

    return render(
        INDEX_HTML,
        scripts=available_scripts,
        jobs=filtered_jobs,
        user=current_user,
        companies=companies,
        current_company=current_company,
        active_company=active_company,
    )


# ═══════════════════════════════════════════════════════════════
# АИС ОИП — встроенный поиск/синхронизация (aisoip_web.py)
# ═══════════════════════════════════════════════════════════════

def _resolve_aisoip_company_id(current_user: User, requested: Optional[str]) -> Optional[int]:
    """У оператора всегда своя единственная компания (company_id из его аккаунта).
    Админ и менеджер (у которых компаний может быть несколько или "все") выбирают
    компанию через ?company_id=."""
    if current_user.role in ("admin", "manager"):
        if requested:
            try:
                return int(requested)
            except ValueError:
                return None
        return None
    return current_user.company_id


def _require_aisoip_company(current_user: User, requested: Optional[str]) -> Optional[int]:
    resolved = _resolve_aisoip_company_id(current_user, requested)
    if resolved is None:
        return None
    if not user_can_access_company(current_user, resolved):
        return None
    return resolved


@app.get("/aisoip", response_class=HTMLResponse)
def aisoip_page(
    company_id: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    resolved = _resolve_aisoip_company_id(current_user, company_id)
    if resolved is not None and not user_can_access_company(current_user, resolved):
        return HTMLResponse("⛔ Нет доступа к этой компании", status_code=403)

    allowed_company_ids = get_user_allowed_company_ids(current_user)
    if allowed_company_ids is None:
        companies = db.query(Company).all()
    elif current_user.role == "manager":
        companies = db.query(Company).filter(Company.id.in_(allowed_company_ids)).all()
    else:
        companies = []

    return render(
        AISOIP_HTML,
        user=current_user,
        companies=companies,
        selected_company_id=resolved,
    )


@app.get("/api/aisoip/stats")
def aisoip_stats(company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return {"total": 0, "debtors": 0, "updated": None, "statuses": []}
    return aisoip_web.get_stats(resolved)


@app.get("/api/aisoip/options")
def aisoip_options(company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return {"executors": [], "statuses": []}
    return aisoip_web.get_options(resolved)


@app.get("/api/aisoip/search")
def aisoip_search(
    request: Request,
    company_id: Optional[str] = None,
    limit: int = 100,
    current_user: User = Depends(get_current_user),
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return {"count": 0, "items": []}
    qs = dict(request.query_params)
    limit = min(max(limit, 1), 500)
    count, items = aisoip_web.search(resolved, qs, limit)
    return {"count": count, "items": items}


@app.get("/api/aisoip/export")
def aisoip_export(
    request: Request,
    company_id: Optional[str] = None,
    current_user: User = Depends(get_current_user),
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        raise HTTPException(status_code=403, detail="Компания не выбрана или недоступна")

    qs = dict(request.query_params)
    rows = aisoip_web.search_all(resolved, qs)

    import csv
    import io as _io

    out = _io.StringIO()
    out.write('﻿')
    w = csv.writer(out, delimiter=';')
    w.writerow([
        "Номер исполнительного документа", "Номер исполнительного производства", "Судебный исполнитель",
        "Дата возбуждения", "Дата документа", "Орган", "Тип взыскания", "Должник", "ИИН/БИН", "Статус", "Сумма",
    ])
    for r in rows:
        w.writerow([
            r["document_no"], r["production_no"], r["executor"], r["start_date"], r["document_date"],
            r["issuer"], r["collection_type"], r["debtor"], r["iin"], r["status"], r["amount"],
        ])
    data = out.getvalue().encode("utf-8")
    return Response(
        content=data,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=ais_oip_search.csv"},
    )


@app.post("/api/aisoip/import")
async def aisoip_import(
    company_id: Optional[str] = Form(None),
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return {"ok": False, "error": "Компания не выбрана или недоступна"}
    if not (file.filename or "").lower().endswith(".xlsx"):
        return {"ok": False, "error": "Поддерживается только .xlsx"}

    import os as _os
    import tempfile

    contents = await file.read()
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    _os.close(fd)
    try:
        Path(tmp).write_bytes(contents)
        total = aisoip_web.import_xlsx(tmp, file.filename, resolved)
    except Exception as e:
        return {"ok": False, "error": str(e)}
    finally:
        try:
            _os.remove(tmp)
        except OSError:
            pass
    return {"ok": True, "processed": total}


@app.get("/api/aisoip/sync/status")
def aisoip_sync_status(company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return {"running": False, "stage": "Компания не выбрана"}
    return aisoip_web.get_sync_manager(resolved).state()


@app.post("/api/aisoip/sync/start")
def aisoip_sync_start(company_id: Optional[str] = None, current_user: User = Depends(get_current_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    resolved = _require_aisoip_company(current_user, company_id)
    if resolved is None:
        return JSONResponse({"ok": False, "message": "Компания не выбрана или недоступна"}, status_code=400)
    if resolved not in aisoip_web.COMPANY_AISOIP_CREDENTIALS:
        return JSONResponse(
            {"ok": False, "message": f"Нет учётных данных АИС ОИП для компании {resolved}"},
            status_code=400,
        )
    ok, message = aisoip_web.get_sync_manager(resolved).start()
    return JSONResponse({"ok": ok, "message": message}, status_code=200 if ok else 409)


@app.get("/run/{script_key}", response_class=HTMLResponse)
def run_form(
    script_key: str,
    session_id: Optional[str] = Cookie(None, alias="session"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    if script_key not in SCRIPTS:
        return HTMLResponse("Скрипт не найден", status_code=404)

    if not current_user.has_script_access(script_key):
        return HTMLResponse("⛔ У вас нет доступа к этому скрипту", status_code=403)

    from datetime import date
    today = date.today().strftime("%Y-%m-%d")
    min_datetime = datetime.now().strftime("%Y-%m-%dT%H:%M")

    script = SCRIPTS[script_key]
    needs_company = any(p.get("name") == "company_id" for p in script.params)
    # Компания больше не выбирается на этой форме — она определяется автоматически
    # из выбора в шапке сайта (admin/manager) или из привязки пользователя (operator).
    _, active_company = _resolve_active_company(current_user, session_id, None, db)

    has_folder = any(p.get("type") == "folder" for p in script.params)
    folder_options = list_rename_folders() if has_folder else []
    rename_root = _rename_root() if has_folder else None

    return render(
        RUN_HTML,
        script=script,
        today=today,
        min_datetime=min_datetime,
        needs_company=needs_company,
        active_company=active_company,
        folder_options=folder_options,
        rename_root=str(rename_root) if rename_root else None,
    )


@app.post("/run/{script_key}", include_in_schema=False)
async def run_submit(
    script_key: str,
    request: Request,
    session_id: Optional[str] = Cookie(None, alias="session"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    if script_key not in SCRIPTS:
        return HTMLResponse("Not found", status_code=404)

    if not current_user.has_script_access(script_key):
        return HTMLResponse("⛔ У вас нет доступа к этому скрипту", status_code=403)

    script = SCRIPTS[script_key]
    form = await request.form()
    params = {}

    for p in script.params:
        if p["name"] == "company_id":
            continue  # решается ниже отдельно — из формы не берётся
        field = form.get(p["name"])
        if p["type"] == "file" and field is not None:
            upload: UploadFile = field
            if upload.filename:
                save_to = p.get("save_to", f"uploads/{uuid.uuid4().hex[:8]}_{upload.filename}")
                dest = BASE_DIR / save_to
                dest.parent.mkdir(parents=True, exist_ok=True)
                contents = await upload.read()
                dest.write_bytes(contents)
                params[p["name"]] = str(dest)
        elif p["type"] == "folder":
            # Из выпадающего списка или из поля ручного ввода (оно приоритетнее)
            raw = (form.get(p["name"] + "__manual") or form.get(p["name"]) or "").strip()
            resolved = resolve_rename_folder(raw)
            if resolved is None:
                return HTMLResponse(
                    "⛔ Папка не найдена или путь недопустим: "
                    f"<b>{raw or '(пусто)'}</b><br><br>"
                    "Выберите папку из списка. Если списка нет — проверьте RENAME_ROOT "
                    "в scripts/config.py на сервере.<br>"
                    '<a href="javascript:history.back()">← назад</a>',
                    status_code=400,
                )
            params[p["name"]] = resolved
        else:
            if field is not None:
                params[p["name"]] = str(field)

    # company_id больше не приходит из формы (её там и нет — см. RUN_HTML) и не
    # зависит от того, что мог бы подделать клиент в POST-запросе. Для скриптов,
    # которым она нужна, она ВСЕГДА определяется здесь на сервере: у operator —
    # это его единственная привязанная компания, у admin/manager — та, что выбрана
    # вкладкой "Компания:" в шапке сайта (и хранится в их сессии).
    needs_company = any(p.get("name") == "company_id" for p in script.params)
    active_company_id, active_company = (None, None)
    if needs_company:
        active_company_id, active_company = _resolve_active_company(current_user, session_id, None, db)
        if active_company_id is None:
            from datetime import date
            today = date.today().strftime("%Y-%m-%d")
            min_datetime = datetime.now().strftime("%Y-%m-%dT%H:%M")
            return HTMLResponse(
                render(
                    RUN_HTML,
                    script=script,
                    today=today,
                    min_datetime=min_datetime,
                    needs_company=True,
                    active_company=None,
                ),
                status_code=400,
            )
        params["company_id"] = str(active_company_id)

    # ─── Обработка отложенного запуска ───
    scheduled_at_raw = form.get("scheduled_at")
    scheduled_ts = None
    scheduled_display = None

    if getattr(script, "allow_scheduled_start", False) and scheduled_at_raw:
        try:
            scheduled_dt = datetime.strptime(scheduled_at_raw, "%Y-%m-%dT%H:%M")
            if scheduled_dt > datetime.now():
                scheduled_ts = scheduled_dt.timestamp()
                scheduled_display = scheduled_dt.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass  # некорректный формат — просто игнорируем, запустим сразу

    job_id = uuid.uuid4().hex[:12]
    job_dir = make_job_dir(job_id)

    # Компания задачи для отображения/фильтрации в интерфейсе — берём из реально
    # определённого company_id (см. _resolve_active_company выше), а не только
    # из current_user.company_id: у admin/manager он всегда None, но сам запуск
    # мог быть для конкретной компании, выбранной в шапке сайта.
    job_company_id = getattr(current_user, "company_id", None)
    job_company_name = getattr(getattr(current_user, "company", None), "name", None)
    resolved_company_id = params.get("company_id")
    if resolved_company_id:
        try:
            job_company_id = int(resolved_company_id)
            company_row = db.query(Company).filter(Company.id == job_company_id).first()
            job_company_name = company_row.name if company_row else job_company_name
        except ValueError:
            pass

    job = {
        "id": job_id,
        "script_key": script_key,
        "params": params,
        "status": "scheduled" if scheduled_ts else "queued",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "started_at": None,
        "finished_at": None,
        "job_dir": str(job_dir),
        "result_zip": None,
        "created_ts": datetime.now().timestamp(),
        "username": getattr(current_user, "username", "—"),
        "company_name": job_company_name or "—",
        "company_id": job_company_id,
        "source": "manual",
        "scheduled_ts": scheduled_ts,
        "scheduled_display": scheduled_display,
    }
    JOBS[job_id] = job
    create_job_run(
        job_id, script_key, source="manual", username=getattr(current_user, "username", "—"),
        company_id=job_company_id,
        company_name=job_company_name,
    )

    if scheduled_ts:
        # НЕ добавляем в QUEUE сразу — ждём наступления времени
        pass
    else:
        with LOCK:
            QUEUE.append(job_id)

    return RedirectResponse("/", status_code=303)


@app.get("/admin", response_class=HTMLResponse)
def admin_panel(
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    
    users = db.query(User).all()
    
    user_scripts_count = {}
    for u in users:
        if u.role != "admin":
            user_scripts_count[u.id] = len(get_user_scripts(u.id))
    
    return render(
        ADMIN_HTML,
        users=users,
        total_scripts=len(visible_scripts()),
        user_scripts_count=user_scripts_count
    )

@app.get("/admin/user/new", response_class=HTMLResponse)
def add_user_form(
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    
    companies = db.query(Company).all()
    return render(
        USER_FORM_HTML,
        user=None,
        companies=companies,
        scripts=visible_scripts(),
        user_scripts=[],
        user_company_ids=[]
    )

@app.post("/admin/user/new")
async def add_user_submit(
    request: Request,
    username: str = Form(...),
    full_name: str = Form(...),
    password: str = Form(...),
    role: str = Form(...),
    company_id: Optional[str] = Form(None),
    is_active: Optional[str] = Form(None),
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    from auth import hash_password

    form = await request.form()
    selected_scripts = form.getlist("scripts")
    selected_companies = [int(x) for x in form.getlist("companies") if x]

    # Одиночная компания (FK) — только у operator. У manager компании живут
    # в user_companies (много), у admin компании не нужны вовсе.
    final_company_id = int(company_id) if (role == "operator" and company_id) else None

    new_user = User(
        username=username,
        full_name=full_name,
        password_hash=hash_password(password),
        role=role,
        company_id=final_company_id,
        is_active=bool(is_active)
    )
    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    if role != "admin":
        set_user_scripts(new_user.id, selected_scripts)
    set_user_companies(new_user.id, selected_companies if role == "manager" else [])

    return RedirectResponse("/admin", status_code=303)

@app.get("/admin/user/{user_id}/edit", response_class=HTMLResponse)
def edit_user_form(
    user_id: int,
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    
    user = db.query(User).filter(User.id == user_id).first()
    companies = db.query(Company).all()
    user_scripts = get_user_scripts(user_id)
    user_company_ids = get_user_companies(user_id)

    return render(
        USER_FORM_HTML,
        user=user,
        companies=companies,
        scripts=visible_scripts(),
        user_scripts=user_scripts,
        user_company_ids=user_company_ids
    )

@app.post("/admin/user/{user_id}/edit")
async def edit_user_submit(
    user_id: int,
    request: Request,
    username: str = Form(...),
    full_name: str = Form(...),
    password: Optional[str] = Form(None),
    role: str = Form(...),
    company_id: Optional[str] = Form(None),
    is_active: Optional[str] = Form(None),
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    
    from auth import hash_password

    form = await request.form()
    selected_scripts = form.getlist("scripts")
    selected_companies = [int(x) for x in form.getlist("companies") if x]

    user = db.query(User).filter(User.id == user_id).first()
    user.username = username
    user.full_name = full_name
    user.role = role
    user.company_id = int(company_id) if (role == "operator" and company_id) else None
    user.is_active = bool(is_active)

    if password:
        user.password_hash = hash_password(password)

    db.commit()
    set_user_companies(user_id, selected_companies if role == "manager" else [])
    
    if role != "admin":
        set_user_scripts(user_id, selected_scripts)
    
    return RedirectResponse("/admin", status_code=303)

@app.get("/admin/user/{user_id}/delete")
def delete_user(
    user_id: int,
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db)
):
    user = db.query(User).filter(User.id == user_id).first()
    if user.username != "admin":
        db.delete(user)
        db.commit()
    return RedirectResponse("/admin", status_code=303)

# ═══════════════════════════════════════════════════════════════
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ═══════════════════════════════════════════════════════════════

def make_job_dir(job_id: str) -> Path:
    d = DATA_DIR / job_id
    (d / "out").mkdir(parents=True, exist_ok=True)
    (d / "downloads").mkdir(parents=True, exist_ok=True)
    (d / "chrome_profile").mkdir(parents=True, exist_ok=True)
    return d

def zip_out(job_dir: Path) -> Optional[Path]:
    out_dir = job_dir / "out"
    if not out_dir.exists():
        return None
    files = [f for f in out_dir.rglob("*") if f.is_file()]
    if not files:
        return None
    # Один файл — возвращаем его напрямую, без архива
    if len(files) == 1:
        return files[0]
    # Несколько файлов — архивируем
    zip_path = job_dir / "result"
    shutil.make_archive(str(zip_path), "zip", root_dir=str(out_dir))
    return Path(str(zip_path) + ".zip")


# ═══════════════════════════════════════════════════════════════
# ВОССТАНОВЛЕНИЕ ИСТОРИИ ЗАДАЧ ПОСЛЕ ПЕРЕЗАПУСКА
# ═══════════════════════════════════════════════════════════════
#
# JOBS — обычный словарь в памяти процесса, поэтому при каждом рестарте
# приложения (деплой, падение, ручной перезапуск) он обнуляется: страница
# "История запусков" и /logs/<id> переставали показывать уже прошедшие
# задачи, хотя сами файлы data/runs/<id>/logs.txt никуда не пропадали —
# их просто было нечем открыть, т.к. /logs/{job_id} ищет job в JOBS.get().
# Восстанавливаем JOBS из таблицы job_runs (она пишется в create_job_run/
# update_job_run_status при каждом запуске и переживает рестарт), поэтому
# история и логи остаются доступны из интерфейса и после перезапуска.
JOB_HISTORY_RETENTION_HOURS = 24 * 30  # держим в памяти историю за 30 дней


def hydrate_jobs_from_db(retention_hours: int = JOB_HISTORY_RETENTION_HOURS):
    db = SessionLocal()
    try:
        cutoff = datetime.now() - timedelta(hours=retention_hours)
        rows = (
            db.query(JobRun)
            .filter(JobRun.created_at >= cutoff)
            .order_by(JobRun.created_at.desc())
            .all()
        )

        restored = 0
        for jr in rows:
            job_dir = DATA_DIR / jr.id
            status = jr.status
            finished_at = jr.finished_at

            # Задача не могла пережить перезапуск процесса как реально
            # выполняющаяся — если в БД она осталась running/queued/
            # scheduled, значит сервис остановился посреди неё.
            if status in ("running", "queued", "scheduled"):
                status = "error"
                finished_at = finished_at or datetime.now()
                update_job_run_status(
                    jr.id, "error",
                    finished_at=finished_at,
                    error_text="Прервано перезапуском сервиса",
                )

            result_zip = None
            if status == "success":
                z = zip_out(job_dir)
                result_zip = str(z) if z else None

            JOBS[jr.id] = {
                "id": jr.id,
                "script_key": jr.script_key,
                "params": {},
                "status": status,
                "created_at": jr.created_at.strftime("%Y-%m-%d %H:%M:%S") if jr.created_at else None,
                "started_at": jr.started_at.strftime("%Y-%m-%d %H:%M:%S") if jr.started_at else None,
                "finished_at": finished_at.strftime("%Y-%m-%d %H:%M:%S") if finished_at else None,
                "job_dir": str(job_dir),
                "result_zip": result_zip,
                "created_ts": jr.created_at.timestamp() if jr.created_at else 0,
                "username": jr.username or "—",
                "company_name": jr.company_name or "—",
                "company_id": jr.company_id,
                "source": jr.source or "manual",
                "scheduled_ts": None,
                "scheduled_display": None,
            }
            restored += 1

        print(f"История задач: восстановлено {restored} записей из БД (за последние {retention_hours} ч.)")
    finally:
        db.close()


hydrate_jobs_from_db()


# ═══════════════════════════════════════════════════════════════
# RUNNER
# ═══════════════════════════════════════════════════════════════

def delayed_dispatcher_loop():
    while True:
        try:
            now_ts = datetime.now().timestamp()
            for job_id, job in list(JOBS.items()):
                if job["status"] == "scheduled" and job.get("scheduled_ts"):
                    if now_ts >= job["scheduled_ts"]:
                        job["status"] = "queued"
                        with LOCK:
                            QUEUE.append(job_id)
        except Exception as e:
            print("DELAYED DISPATCHER ERROR:", e)
        import time
        time.sleep(5)  # проверка каждые 5 секунд — достаточно точно

threading.Thread(target=delayed_dispatcher_loop, daemon=True).start()

def execute_job(job_id: str):
    global RUNNING_COUNT
    job = JOBS[job_id]
    script = SCRIPTS[job["script_key"]]
    job_dir = Path(job["job_dir"])
    log_path = job_dir / "logs.txt"

    cmd = list(script.command)
    cmd += ["--workdir", str(job_dir)]
    for k, v in job["params"].items():
        if v:
            cmd += [f"--{k}", str(v)]

    error_text = None
    try:
        with open(log_path, "w", encoding="utf-8") as lf:
            lf.write(f"COMMAND: {' '.join(cmd)}\n")
            lf.write(f"STARTED: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
            lf.flush()

            p = subprocess.Popen(
                cmd, stdout=lf, stderr=lf, cwd=str(BASE_DIR),
                text=True, encoding="utf-8", errors="replace",
            )
            PROCESSES[job_id] = p
            rc = p.wait()

        if job.get("status") == "cancelled":
            # Остановлен вручную через интерфейс — не перезаписываем статус
            # по коду возврата (killed-процесс обычно возвращает ненулевой rc).
            job["result_zip"] = None
        elif rc == 0:
            job["status"] = "success"
            z = zip_out(job_dir)
            job["result_zip"] = str(z) if z else None
        else:
            job["status"] = "error"
            job["result_zip"] = None
            error_text = f"exit code {rc}"

    except Exception as e:
        if job.get("status") != "cancelled":
            job["status"] = "error"
            error_text = repr(e)
        with open(log_path, "a", encoding="utf-8") as lf:
            lf.write(f"\nEXCEPTION:\n{repr(e)}\n")

    finally:
        PROCESSES.pop(job_id, None)
        job["finished_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        update_job_run_status(
            job_id, job["status"],
            finished_at=datetime.now(),
            error_text=error_text,
        )
        with LOCK:
            RUNNING_COUNT -= 1

threading.Thread(target=runner_loop, daemon=True).start()

# ═══════════════════════════════════════════════════════════════
# ЛОГИ И СКАЧИВАНИЕ
# ═══════════════════════════════════════════════════════════════

@app.get("/logs/{job_id}", response_class=HTMLResponse)
def logs(job_id: str):
    job = JOBS.get(job_id)
    if not job:
        return HTMLResponse("Not found", status_code=404)
    log_path = Path(job["job_dir"]) / "logs.txt"
    log_text = log_path.read_text(encoding="utf-8", errors="replace") if log_path.exists() else "(логов пока нет)"
    return render(LOGS_HTML, job=job, log_text=log_text)


@app.get("/stop/{job_id}")
def stop_job(
    job_id: str,
    request: Request,
    current_user: User = Depends(get_current_user),
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    job = JOBS.get(job_id)
    if not job:
        return HTMLResponse("Задача не найдена", status_code=404)

    is_admin = getattr(current_user, "role", None) == "admin"
    is_owner = job.get("username") == getattr(current_user, "username", None)
    if not is_admin and not is_owner:
        return HTMLResponse("⛔ У вас нет доступа к этой задаче", status_code=403)

    if job["status"] in ("queued", "scheduled"):
        with LOCK:
            if job_id in QUEUE:
                QUEUE.remove(job_id)
        job["status"] = "cancelled"
        job["finished_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        update_job_run_status(job_id, "cancelled", finished_at=datetime.now())

    elif job["status"] == "running":
        job["status"] = "cancelled"  # execute_job() увидит это и не перезапишет статус
        p = PROCESSES.get(job_id)
        if p is not None:
            kill_process_tree(p.pid)

    referer = request.headers.get("referer") or "/"
    return RedirectResponse(referer, status_code=303)

@app.get("/download/{job_id}")
def download(job_id: str):
    job = JOBS.get(job_id)
    if not job or not job.get("result_zip"):
        return HTMLResponse("Нет результата", status_code=404)
    zp = Path(job["result_zip"])
    # Отдаём с оригинальным именем файла
    return FileResponse(str(zp), filename=zp.name)

@app.get("/status")
def status():
    return {
        "running_count": RUNNING_COUNT,
        "max_parallel_jobs": MAX_PARALLEL_JOBS,
        "queue_length": len(QUEUE),
        "total_jobs": len(JOBS),
    }



SCHEDULE_LIST_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Расписание</title>
<style>
body { font-family: Arial; max-width: 1100px; margin: 20px auto; }
table { width: 100%; border-collapse: collapse; margin: 20px 0; }
th, td { padding: 10px; text-align: left; border: 1px solid #ddd; }
th { background: #6610f2; color: white; }
.btn { padding: 6px 12px; text-decoration: none; border-radius: 4px; color: white; display: inline-block; margin: 2px; }
.btn-edit { background: #007bff; }
.btn-delete { background: #dc3545; }
.btn-toggle { background: #ffc107; color: #000; }
.btn-add { background: #28a745; padding: 10px 20px; }
.status-on { color: #28a745; font-weight: bold; }
.status-off { color: #999; }
</style>
</head>
<body>
<a href="/admin">← Назад в админку</a> | <a href="/monitoring">📊 Мониторинг автозапусков</a>
<h2>⏰ Расписание автозапуска</h2>
<a href="/admin/schedule/new" class="btn btn-add">➕ Добавить расписание</a>
<table>
  <tr><th>Скрипт</th><th>Компания</th><th>Время</th><th>Дни недели</th><th>Статус</th><th>Действия</th></tr>
  {% for s in schedules %}
  <tr>
    <td>{{ scripts[s.script_key].title if s.script_key in scripts else s.script_key }}</td>
    <td>{{ s.company_name }}</td>
    <td>{{ '%02d:%02d' % (s.hour, s.minute) }}</td>
    <td>{{ s.weekdays_display }}</td>
    <td>{% if s.is_active %}<span class="status-on">Включено</span>{% else %}<span class="status-off">Выключено</span>{% endif %}</td>
    <td>
      <a href="/admin/schedule/{{ s.id }}/edit" class="btn btn-edit">Изменить</a>
      <a href="/admin/schedule/{{ s.id }}/toggle" class="btn btn-toggle">{% if s.is_active %}Выключить{% else %}Включить{% endif %}</a>
      <a href="/admin/schedule/{{ s.id }}/delete" class="btn btn-delete" onclick="return confirm('Удалить расписание?')">Удалить</a>
    </td>
  </tr>
  {% endfor %}
</table>
</body></html>
"""

SCHEDULE_FORM_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Расписание</title>
<style>
body { font-family: Arial; max-width: 600px; margin: 40px auto; }
.form-group { margin: 15px 0; }
label { display: block; margin-bottom: 5px; font-weight: bold; }
select, input { padding: 8px; width: 100%; box-sizing: border-box; }
.weekdays { display: flex; gap: 10px; }
.weekdays label { font-weight: normal; display: flex; align-items: center; gap: 4px; }
button { padding: 12px 24px; background: #6610f2; color: white; border: none; border-radius: 4px; cursor: pointer; }
</style>
</head>
<body>
<a href="/admin/schedule">← Назад</a>
<h2>{% if schedule %}Изменить{% else %}Новое{% endif %} расписание</h2>
<form method="post">
  <div class="form-group">
    <label>Скрипт:</label>
    <select name="script_key" required>
      {% for key, s in scripts.items() %}
        <option value="{{ key }}" {% if schedule and schedule.script_key == key %}selected{% endif %}>{{ s.title }}</option>
      {% endfor %}
    </select>
  </div>
  <div class="form-group">
    <label>Время запуска:</label>
    <input type="time" name="time" value="{{ '%02d:%02d' % (schedule.hour, schedule.minute) if schedule else '09:00' }}" required>
  </div>
  <div class="form-group">
    <label>Компания:</label>
    <select name="company_id">
      <option value="">Без привязки (по умолчанию скрипта)</option>
      {% for c in companies %}
        <option value="{{ c.id }}" {% if schedule and schedule.company_id == c.id %}selected{% endif %}>{{ c.name }}</option>
      {% endfor %}
    </select>
  </div>
  <div class="form-group">
    <label>Дни недели:</label>
    <div class="weekdays">
      {% set names = ['Пн','Вт','Ср','Чт','Пт','Сб','Вс'] %}
      {% for i in range(7) %}
        <label>
          <input type="checkbox" name="weekdays" value="{{ i }}"
            {% if not schedule or i in schedule.weekday_list() %}checked{% endif %}>
          {{ names[i] }}
        </label>
      {% endfor %}
    </div>
  </div>
  <button type="submit">💾 Сохранить</button>
</form>
</body></html>
"""

@app.get("/admin/schedule", response_class=HTMLResponse)
def schedule_list(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    names = ['Пн','Вт','Ср','Чт','Пт','Сб','Вс']
    company_map = {c.id: c.name for c in db.query(Company).all()}
    schedules = get_all_schedules()
    for s in schedules:
        wl = s.weekday_list()
        s.weekdays_display = "Ежедневно" if len(wl) == 7 else ", ".join(names[i] for i in wl)
        s.company_name = company_map.get(s.company_id, "Без привязки")
    return render(SCHEDULE_LIST_HTML, schedules=schedules, scripts=SCRIPTS)

@app.get("/admin/schedule/new", response_class=HTMLResponse)
def schedule_new_form(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    companies = db.query(Company).all()
    return render(SCHEDULE_FORM_HTML, schedule=None, scripts=visible_scripts(), companies=companies)

@app.post("/admin/schedule/new")
async def schedule_new_submit(
    request: Request,
    script_key: str = Form(...),
    time: str = Form(...),
    company_id: Optional[str] = Form(None),
    current_user: User = Depends(get_admin_user)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    weekdays = form.getlist("weekdays")
    hour, minute = map(int, time.split(":"))
    create_schedule(
        script_key, hour, minute, ",".join(weekdays) or "0,1,2,3,4,5,6",
        company_id=int(company_id) if company_id else None,
    )
    return RedirectResponse("/admin/schedule", status_code=303)

@app.get("/admin/schedule/{schedule_id}/edit", response_class=HTMLResponse)
def schedule_edit_form(schedule_id: int, current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    s = get_schedule(schedule_id)
    companies = db.query(Company).all()
    return render(SCHEDULE_FORM_HTML, schedule=s, scripts=visible_scripts(), companies=companies)

@app.post("/admin/schedule/{schedule_id}/edit")
async def schedule_edit_submit(
    schedule_id: int,
    request: Request,
    script_key: str = Form(...),
    time: str = Form(...),
    company_id: Optional[str] = Form(None),
    current_user: User = Depends(get_admin_user)
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    weekdays = form.getlist("weekdays")
    hour, minute = map(int, time.split(":"))
    update_schedule(
        schedule_id, script_key, hour, minute, ",".join(weekdays) or "0,1,2,3,4,5,6",
        company_id=int(company_id) if company_id else None,
    )
    return RedirectResponse("/admin/schedule", status_code=303)

@app.get("/admin/schedule/{schedule_id}/toggle")
def schedule_toggle_route(schedule_id: int, current_user: User = Depends(get_admin_user)):
    toggle_schedule(schedule_id)
    return RedirectResponse("/admin/schedule", status_code=303)

@app.get("/admin/schedule/{schedule_id}/delete")
def schedule_delete_route(schedule_id: int, current_user: User = Depends(get_admin_user)):
    delete_schedule(schedule_id)
    return RedirectResponse("/admin/schedule", status_code=303)


# ═══════════════════════════════════════════════════════════════
# РАССЫЛКИ ПО ОТЧЁТАМ (ЧСИ)
# ═══════════════════════════════════════════════════════════════

MAILING_DEFAULT_SUBJECT = (
    "Отчет по прекращенным производствам / "
    "Тоқтатылған атқарушылық іс жүргізулер бойынша есеп"
)

MAILING_DEFAULT_BODY = """<html><body style="font-family:Arial,sans-serif;font-size:14px;line-height:1.6">
<p>Здравствуйте, {{ФИО_ЧСИ}}!</p>
<p>Во вложении список займов ({{кол_во_записей}} шт.), по которым прекращено производство
на основании пункта 7 статьи 47 Закона «Об исполнительном производстве и статусе судебных
исполнителей», однако денежные средства нам не перечислены.</p>
<p>Просим перечислить удержанные средства в нашу Компанию для изменения статуса кредита
на «Погашен».</p>
<p><b>{{компания}}</b><br>{{дата}}</p>
</body></html>"""

_MAIL_CSS = """
body { font-family: Arial; max-width: 1100px; margin: 20px auto; padding: 0 16px; }
h2 { border-bottom: 2px solid #e67e22; padding-bottom: 10px; }
table { width: 100%; border-collapse: collapse; margin: 16px 0; }
th, td { padding: 9px; text-align: left; border: 1px solid #ddd; font-size: 14px; }
th { background: #e67e22; color: white; }
.btn { padding: 6px 12px; text-decoration: none; border-radius: 4px; color: white; display: inline-block; margin: 2px; border: none; cursor: pointer; font-size: 13px; }
.b-run { background: #007bff; } .b-test { background: #6f42c1; } .b-send { background: #dc3545; }
.b-edit { background: #007bff; } .b-del { background: #6c757d; } .b-add { background: #28a745; padding: 9px 16px; }
.form-group { margin: 14px 0; }
label { display: block; margin-bottom: 5px; font-weight: bold; }
input[type=text], input[type=email], input[type=time], select, textarea {
  padding: 9px; width: 100%; box-sizing: border-box; border: 1px solid #ccc; border-radius: 4px; font-size: 14px;
}
textarea { font-family: Consolas, monospace; }
.hint { color: #666; font-size: 12px; margin-top: 4px; }
.weekdays { display: flex; gap: 12px; flex-wrap: wrap; }
.weekdays label { font-weight: normal; display: flex; align-items: center; gap: 4px; }
.tabs a { display: inline-block; padding: 7px 14px; margin: 2px; background: #eee; border-radius: 5px; text-decoration: none; color: #333; }
.tabs a.active { background: #e67e22; color: white; }
.warn { background: #fff3cd; border: 1px solid #ffe08a; padding: 12px; border-radius: 6px; color: #7a5b00; }
"""

MAILINGS_LIST_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Рассылки по отчётам</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin">← Назад в админку</a>
<h2>📨 Рассылки по отчётам</h2>
<p>
  <a href="/admin/mailings/new" class="btn b-add">➕ Новая рассылка</a>
  <a href="/admin/mailings/reports" class="btn b-add" style="background:#17a2b8;">📋 Отчёты (SQL)</a>
  <a href="/admin/mailings/contacts" class="btn b-add" style="background:#6610f2;">📇 Справочник ЧСИ</a>
</p>
<table>
  <tr><th>Название</th><th>Компания</th><th>Источник</th><th>Колонка-ЧСИ</th><th>Расписание</th><th>Статус</th><th>Действия</th></tr>
  {% for m in mailings %}
  <tr>
    <td>{{ m.name }}</td>
    <td>{{ mailing_company_names.get(m.id, '—') }}</td>
    <td>{% if m.sql_text and m.sql_text.strip() %}свой SQL{% elif m.report_id %}{{ report_map.get(m.report_id, 'отчёт') }}{% else %}<span style="color:#dc3545;">не задан</span>{% endif %}</td>
    <td>{{ m.group_column }}</td>
    <td>{% if m.schedule_enabled %}{{ '%02d:%02d' % (m.schedule_hour or 0, m.schedule_minute or 0) }}{% else %}—{% endif %}</td>
    <td>{% if m.is_active %}<span style="color:#28a745;">вкл</span>{% else %}<span style="color:#6c757d;">выкл</span>{% endif %}</td>
    <td style="white-space:nowrap;">
      <form method="post" action="/admin/mailings/{{ m.id }}/test" style="display:inline;">
        <button class="btn b-test" onclick="return confirm('Отправить ТЕСТОВОЕ письмо на {{ m.test_email or 'адрес из рассылки' }}?')">Тест</button>
      </form>
      <a href="/admin/mailings/{{ m.id }}/send" class="btn b-send">Отправить</a>
      <a href="/admin/mailings/{{ m.id }}/edit" class="btn b-edit">Изменить</a>
      <a href="/admin/mailings/{{ m.id }}/toggle" class="btn b-del">{% if m.is_active %}Выкл{% else %}Вкл{% endif %}</a>
      <a href="/admin/mailings/{{ m.id }}/delete" class="btn b-del" onclick="return confirm('Удалить рассылку?')">✕</a>
    </td>
  </tr>
  {% endfor %}
  {% if not mailings %}<tr><td colspan="7" style="color:#666;">Пока нет ни одной рассылки.</td></tr>{% endif %}
</table>
</body></html>
"""

MAILING_FORM_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Рассылка</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin/mailings">← Назад к рассылкам</a>
<h2>{% if mailing %}Изменить рассылку{% else %}Новая рассылка{% endif %}</h2>
<form method="post">
  <div class="form-group">
    <label>Название</label>
    <input type="text" name="name" required value="{{ mailing.name if mailing else '' }}">
  </div>
  <div class="form-group">
    <label>Компани{{ 'я' if companies|length == 1 else 'и' }} (можно выбрать несколько — рассылка
      прогонится по очереди, свой SQL/SMTP/справочник ЧСИ для каждой)</label>
    {% set selected_ids = mailing.company_id_list() if mailing else [] %}
    {% for c in companies %}
      <label style="display:inline-block;width:auto;margin-right:16px;font-weight:normal;">
        <input type="checkbox" name="company_ids" value="{{ c.id }}"
          {% if c.id in selected_ids %}checked{% endif %} style="width:auto;"> {{ c.name }}
      </label>
    {% endfor %}
  </div>
  <div class="form-group">
    <label>Готовый отчёт (SQL)</label>
    <select name="report_id">
      <option value="">— не использовать —</option>
      {% for r in reports %}
        <option value="{{ r.id }}" {% if mailing and mailing.report_id == r.id %}selected{% endif %}>{{ r.name }}{% if r.company_id %} ({{ company_map.get(r.company_id, '') }}){% endif %}</option>
      {% endfor %}
    </select>
    <div class="hint">Управление отчётами — на странице «Отчёты (SQL)».</div>
  </div>
  <div class="form-group">
    <label>…или свой SQL-запрос</label>
    <textarea name="sql_text" rows="8" placeholder="SELECT ... WHERE ... l.F209 = {company_filter}">{{ mailing.sql_text if mailing and mailing.sql_text else '' }}</textarea>
    <div class="hint">Если это поле заполнено — используется оно (отчёт игнорируется). Только SELECT.
      Подстановка <code>{company_filter}</code> заменяется на имя компании в CRM.</div>
  </div>
  <div class="form-group">
    <label>Колонка результата, по которой группировать (= ЧСИ)</label>
    <input type="text" name="group_column" required value="{{ mailing.group_column if mailing else '' }}" placeholder="напр. ЧСИ или ФИО ЧСИ">
  </div>
  <div class="form-group">
    <label>Тема письма</label>
    <input type="text" name="subject" required value="{{ mailing.subject if mailing else default_subject }}">
  </div>
  <div class="form-group">
    <label>Текст письма (HTML)</label>
    <textarea name="body_html" rows="14" required>{{ mailing.body_html if mailing else default_body }}</textarea>
    <div class="hint">Доступные подстановки:
      {% for ph, desc in placeholders.items() %}<code>{{ ph }}</code> — {{ desc }}{% if not loop.last %}; {% endif %}{% endfor %}</div>
  </div>
  <div class="form-group">
    <label><input type="checkbox" name="attach_enabled" value="1" {% if not mailing or mailing.attach_enabled %}checked{% endif %} style="width:auto;"> Прикладывать xlsx со строками этого ЧСИ</label>
  </div>
  <div class="form-group">
    <label>Тестовый email (для кнопки «Тест»)</label>
    <input type="text" name="test_email" value="{{ mailing.test_email or '' if mailing else '' }}">
  </div>
  <div class="form-group">
    <label>Скрытая копия (BCC), адреса через запятую</label>
    <input type="text" name="bcc" value="{{ mailing.bcc or '' if mailing else '' }}">
    <div class="hint">Объединяется с BCC, заданным в config.MAILING_SMTP для компании.</div>
  </div>
  <div class="form-group">
    <label><input type="checkbox" name="is_active" value="1" {% if not mailing or mailing.is_active %}checked{% endif %} style="width:auto;"> Рассылка активна</label>
  </div>

  <fieldset style="border:1px solid #ddd; border-radius:6px; padding:12px;">
    <legend>Автозапуск по расписанию</legend>
    <div class="form-group">
      <label><input type="checkbox" name="schedule_enabled" value="1" {% if mailing and mailing.schedule_enabled %}checked{% endif %} style="width:auto;"> Включить расписание (режим «реальная отправка»)</label>
    </div>
    <div class="form-group">
      <label>Время</label>
      <input type="time" name="sched_time" value="{{ '%02d:%02d' % (mailing.schedule_hour or 9, mailing.schedule_minute or 0) if mailing else '09:00' }}">
    </div>
    <div class="form-group">
      <label>Дни недели</label>
      <div class="weekdays">
        {% set names = ['Пн','Вт','Ср','Чт','Пт','Сб','Вс'] %}
        {% set wl = mailing.weekday_list() if mailing else [0,1,2,3,4,5,6] %}
        {% for i in range(7) %}
          <label><input type="checkbox" name="weekdays" value="{{ i }}" {% if i in wl %}checked{% endif %}> {{ names[i] }}</label>
        {% endfor %}
      </div>
    </div>
  </fieldset>

  <p><button type="submit" class="btn b-add" style="font-size:15px;">💾 Сохранить</button></p>
</form>
</body></html>
"""

MAILING_CONFIRM_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Отправка рассылки</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin/mailings">← Назад к рассылкам</a>
<h2>Реальная отправка: {{ mailing.name }}</h2>
<div class="warn">
  <p>Письма уйдут <b>каждому ЧСИ</b> на его адрес из справочника. Действие необратимо.</p>
  <ul>
    <li>Компани{{ 'я' if contacts_by_company|length == 1 else 'и' }}: <b>{{ company_name }}</b></li>
    <li>Источник: {% if mailing.sql_text and mailing.sql_text.strip() %}свой SQL{% else %}{{ report_name }}{% endif %}</li>
    <li>Активных ЧСИ в справочнике: <b>{{ contacts_count }}</b>
      {% if contacts_by_company|length > 1 %}
      <ul>{% for cname, n in contacts_by_company.items() %}<li>{{ cname }}: {{ n }}</li>{% endfor %}</ul>
      {% endif %}
    </li>
    <li>Вложение: {{ 'да' if mailing.attach_enabled else 'нет' }}</li>
    <li>BCC: {{ mailing.bcc or '—' }}</li>
  </ul>
</div>
<form method="post" style="margin-top:16px;">
  <button type="submit" class="btn b-send" style="font-size:15px;">Да, отправить всем</button>
  <a href="/admin/mailings" class="btn b-del" style="font-size:15px;">Отмена</a>
</form>
</body></html>
"""

MAILING_REPORTS_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Отчёты (SQL)</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin/mailings">← Назад к рассылкам</a>
<h2>📋 Отчёты (SQL)</h2>
<p><a href="/admin/mailings/reports/new" class="btn b-add">➕ Новый отчёт</a></p>
<table>
  <tr><th>Название</th><th>Компания</th><th>SQL (начало)</th><th>Действия</th></tr>
  {% for r in reports %}
  <tr>
    <td>{{ r.name }}</td>
    <td>{{ company_map.get(r.company_id, 'общий') }}</td>
    <td><code>{{ (r.sql_text or '')[:80] }}…</code></td>
    <td style="white-space:nowrap;">
      <a href="/admin/mailings/reports/{{ r.id }}/edit" class="btn b-edit">Изменить</a>
      <a href="/admin/mailings/reports/{{ r.id }}/delete" class="btn b-del" onclick="return confirm('Удалить отчёт?')">✕</a>
    </td>
  </tr>
  {% endfor %}
  {% if not reports %}<tr><td colspan="4" style="color:#666;">Пока нет отчётов.</td></tr>{% endif %}
</table>
</body></html>
"""

MAILING_REPORT_FORM_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Отчёт (SQL)</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin/mailings/reports">← Назад к отчётам</a>
<h2>{% if report %}Изменить отчёт{% else %}Новый отчёт{% endif %}</h2>
<form method="post">
  <div class="form-group">
    <label>Название</label>
    <input type="text" name="name" required value="{{ report.name if report else '' }}">
  </div>
  <div class="form-group">
    <label>Компания</label>
    <select name="company_id">
      <option value="">— общий (для всех) —</option>
      {% for c in companies %}
        <option value="{{ c.id }}" {% if report and report.company_id == c.id %}selected{% endif %}>{{ c.name }}</option>
      {% endfor %}
    </select>
  </div>
  <div class="form-group">
    <label>SQL-запрос (только SELECT)</label>
    <textarea name="sql_text" rows="16" required>{{ report.sql_text if report else '' }}</textarea>
    <div class="hint">Подстановка <code>{company_filter}</code> в WHERE заменяется на имя компании
      в CRM (поле l.F209) при запуске рассылки.</div>
  </div>
  <p><button type="submit" class="btn b-add" style="font-size:15px;">💾 Сохранить</button></p>
</form>
</body></html>
"""

CHSI_CONTACTS_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Справочник ЧСИ</title>
<style>""" + _MAIL_CSS + """</style></head><body>
<a href="/admin/mailings">← Назад к рассылкам</a>
<h2>📇 Справочник ЧСИ</h2>
<div class="tabs">
  {% for c in companies %}
    <a href="/admin/mailings/contacts?company_id={{ c.id }}" class="{% if c.id == company_id %}active{% endif %}">{{ c.name }}</a>
  {% endfor %}
</div>

{% if company_id %}
<h3 style="margin-top:18px;">Добавить ЧСИ</h3>
<form method="post" action="/admin/mailings/contacts/add">
  <input type="hidden" name="company_id" value="{{ company_id }}">
  <div style="display:flex; gap:10px; flex-wrap:wrap; align-items:end;">
    <div style="flex:2;"><label>ФИО ЧСИ</label><input type="text" name="fio" required></div>
    <div style="flex:2;"><label>Email</label><input type="text" name="email" required></div>
    <div><button type="submit" class="btn b-add">Добавить</button></div>
  </div>
</form>

<h3 style="margin-top:18px;">Загрузить из xlsx</h3>
<form method="post" action="/admin/mailings/contacts/import" enctype="multipart/form-data">
  <input type="hidden" name="company_id" value="{{ company_id }}">
  <div style="display:flex; gap:10px; align-items:end;">
    <div><input type="file" name="file" accept=".xlsx,.xls" required></div>
    <div><button type="submit" class="btn b-add">Импорт</button></div>
  </div>
  <div class="hint">Колонка A — ФИО, колонка B — email. Первая строка (заголовок) пропускается.</div>
</form>
{% if imported is not none %}<p style="color:#28a745;">Добавлено записей: {{ imported }}</p>{% endif %}

<h3 style="margin-top:18px;">Контакты ({{ contacts|length }})</h3>
{% for k in contacts %}
<form method="post" action="/admin/mailings/contacts/{{ k.id }}/edit"
      style="display:flex; gap:10px; align-items:center; padding:6px 0; border-bottom:1px solid #eee; flex-wrap:wrap;">
  <input type="hidden" name="company_id" value="{{ company_id }}">
  <input type="text" name="fio" value="{{ k.fio }}" style="flex:2; min-width:180px;">
  <input type="text" name="email" value="{{ k.email }}" style="flex:2; min-width:180px;">
  <label style="font-weight:normal; white-space:nowrap;">
    <input type="checkbox" name="is_active" value="1" {% if k.is_active %}checked{% endif %} style="width:auto;"> активен
  </label>
  <button type="submit" class="btn b-edit">Сохранить</button>
  <a href="/admin/mailings/contacts/{{ k.id }}/delete?company_id={{ company_id }}" class="btn b-del" onclick="return confirm('Удалить?')">✕</a>
</form>
{% endfor %}
{% if not contacts %}<p style="color:#666;">Пока нет контактов для этой компании.</p>{% endif %}
{% else %}
<p style="margin-top:18px; color:#666;">Выберите компанию.</p>
{% endif %}
</body></html>
"""


def _mailing_company_map(db):
    return {c.id: c.name for c in db.query(Company).all()}


def _mailing_form_common(db):
    companies = db.query(Company).order_by(Company.id).all()
    return companies, {c.id: c.name for c in companies}


@app.get("/admin/mailings", response_class=HTMLResponse)
def mailings_list(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    mailings = get_mailings()
    report_map = {r.id: r.name for r in get_mailing_reports()}
    cmap = _mailing_company_map(db)
    mailing_company_names = {
        m.id: ", ".join(cmap.get(cid, f"#{cid}") for cid in m.company_id_list())
        for m in mailings
    }
    return render(
        MAILINGS_LIST_HTML,
        mailings=mailings,
        company_map=cmap,
        mailing_company_names=mailing_company_names,
        report_map=report_map,
    )


@app.get("/admin/mailings/new", response_class=HTMLResponse)
def mailing_new_form(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    companies, cmap = _mailing_form_common(db)
    return render(
        MAILING_FORM_HTML,
        mailing=None, companies=companies, company_map=cmap,
        reports=get_mailing_reports(), placeholders=PLACEHOLDERS,
        default_subject=MAILING_DEFAULT_SUBJECT, default_body=MAILING_DEFAULT_BODY,
    )


def _mailing_form_values(form):
    t = (form.get("sched_time") or "09:00").strip()
    try:
        sh, sm = [int(x) for x in t.split(":")]
    except Exception:
        sh, sm = 9, 0
    report_id = form.get("report_id")
    company_ids = [int(x) for x in form.getlist("company_ids") if x.strip().isdigit()]
    return dict(
        name=(form.get("name") or "").strip(),
        company_id=company_ids[0] if company_ids else None,
        extra_company_ids=",".join(str(x) for x in company_ids[1:]) or None,
        report_id=int(report_id) if report_id else None,
        sql_text=(form.get("sql_text") or "").strip() or None,
        group_column=(form.get("group_column") or "").strip(),
        subject=(form.get("subject") or "").strip(),
        body_html=form.get("body_html") or "",
        attach_enabled=bool(form.get("attach_enabled")),
        test_email=(form.get("test_email") or "").strip() or None,
        bcc=(form.get("bcc") or "").strip() or None,
        is_active=bool(form.get("is_active")),
        schedule_enabled=bool(form.get("schedule_enabled")),
        schedule_hour=sh,
        schedule_minute=sm,
        schedule_weekdays=",".join(form.getlist("weekdays")) or "0,1,2,3,4,5,6",
    )


@app.post("/admin/mailings/new")
async def mailing_new_submit(request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    vals = _mailing_form_values(form)
    if not vals["company_id"]:
        return HTMLResponse("⛔ Не выбрана компания. <a href='javascript:history.back()'>назад</a>", status_code=400)
    create_mailing(**vals)
    return RedirectResponse("/admin/mailings", status_code=303)


@app.get("/admin/mailings/{mailing_id}/edit", response_class=HTMLResponse)
def mailing_edit_form(mailing_id: int, current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    m = get_mailing(mailing_id)
    if not m:
        return HTMLResponse("Рассылка не найдена", status_code=404)
    companies, cmap = _mailing_form_common(db)
    return render(
        MAILING_FORM_HTML,
        mailing=m, companies=companies, company_map=cmap,
        reports=get_mailing_reports(), placeholders=PLACEHOLDERS,
        default_subject=MAILING_DEFAULT_SUBJECT, default_body=MAILING_DEFAULT_BODY,
    )


@app.post("/admin/mailings/{mailing_id}/edit")
async def mailing_edit_submit(mailing_id: int, request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    vals = _mailing_form_values(form)
    if not vals["company_id"]:
        return HTMLResponse("⛔ Не выбрана компания. <a href='javascript:history.back()'>назад</a>", status_code=400)
    update_mailing(mailing_id, **vals)
    return RedirectResponse("/admin/mailings", status_code=303)


@app.get("/admin/mailings/{mailing_id}/delete")
def mailing_delete(mailing_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    delete_mailing(mailing_id)
    return RedirectResponse("/admin/mailings", status_code=303)


@app.get("/admin/mailings/{mailing_id}/toggle")
def mailing_toggle(mailing_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    toggle_mailing(mailing_id)
    return RedirectResponse("/admin/mailings", status_code=303)


def _enqueue_mailing_job(mailing, mode: str, username: str) -> str:
    job_id = uuid.uuid4().hex[:12]
    job_dir = make_job_dir(job_id)
    ids = mailing.company_id_list()
    db = SessionLocal()
    try:
        names = [c.name for c in db.query(Company).filter(Company.id.in_(ids)).all()]
        company_name = ", ".join(names) if names else None
    finally:
        db.close()
    params = {"mailing_id": str(mailing.id), "mode": mode, "company_id": ",".join(str(x) for x in ids)}
    JOBS[job_id] = {
        "id": job_id, "script_key": "mailing_send", "params": params,
        "status": "queued", "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "started_at": None, "finished_at": None, "job_dir": str(job_dir),
        "result_zip": None, "created_ts": datetime.now().timestamp(),
        "username": username, "company_name": company_name or "—",
        "company_id": mailing.company_id, "source": "manual",
        "scheduled_ts": None, "scheduled_display": None,
    }
    create_job_run(job_id, "mailing_send", source="manual", username=username,
                   company_id=mailing.company_id, company_name=company_name)
    with LOCK:
        QUEUE.append(job_id)
    return job_id


@app.post("/admin/mailings/{mailing_id}/test")
def mailing_test(mailing_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    m = get_mailing(mailing_id)
    if not m:
        return HTMLResponse("Рассылка не найдена", status_code=404)
    _enqueue_mailing_job(m, "test", getattr(current_user, "username", "—"))
    return RedirectResponse("/", status_code=303)


@app.get("/admin/mailings/{mailing_id}/send", response_class=HTMLResponse)
def mailing_send_confirm(mailing_id: int, current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    m = get_mailing(mailing_id)
    if not m:
        return HTMLResponse("Рассылка не найдена", status_code=404)
    ids = m.company_id_list()
    companies = db.query(Company).filter(Company.id.in_(ids)).all()
    cname_by_id = {c.id: c.name for c in companies}
    contacts_by_company = {
        cname_by_id.get(cid, f"#{cid}"): len(get_chsi_contacts(cid, only_active=True))
        for cid in ids
    }
    report_name = "—"
    if m.report_id:
        r = get_mailing_report(m.report_id)
        report_name = r.name if r else "отчёт"
    return render(
        MAILING_CONFIRM_HTML,
        mailing=m, company_name=", ".join(cname_by_id.get(cid, f"#{cid}") for cid in ids),
        contacts_by_company=contacts_by_company,
        contacts_count=sum(contacts_by_company.values()), report_name=report_name,
    )


@app.post("/admin/mailings/{mailing_id}/send")
def mailing_send_run(mailing_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    m = get_mailing(mailing_id)
    if not m:
        return HTMLResponse("Рассылка не найдена", status_code=404)
    _enqueue_mailing_job(m, "real", getattr(current_user, "username", "—"))
    return RedirectResponse("/", status_code=303)


# ── Отчёты (SQL) ─────────────────────────────────────────────

@app.get("/admin/mailings/reports", response_class=HTMLResponse)
def mailing_reports_list(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    return render(
        MAILING_REPORTS_HTML,
        reports=get_mailing_reports(),
        company_map=_mailing_company_map(db),
    )


@app.get("/admin/mailings/reports/new", response_class=HTMLResponse)
def mailing_report_new_form(current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    companies, _ = _mailing_form_common(db)
    return render(MAILING_REPORT_FORM_HTML, report=None, companies=companies)


@app.post("/admin/mailings/reports/new")
async def mailing_report_new_submit(request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    cid = form.get("company_id")
    create_mailing_report(
        name=(form.get("name") or "").strip(),
        sql_text=form.get("sql_text") or "",
        company_id=int(cid) if cid else None,
    )
    return RedirectResponse("/admin/mailings/reports", status_code=303)


@app.get("/admin/mailings/reports/{report_id}/edit", response_class=HTMLResponse)
def mailing_report_edit_form(report_id: int, current_user: User = Depends(get_admin_user), db: Session = Depends(get_db)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    r = get_mailing_report(report_id)
    if not r:
        return HTMLResponse("Отчёт не найден", status_code=404)
    companies, _ = _mailing_form_common(db)
    return render(MAILING_REPORT_FORM_HTML, report=r, companies=companies)


@app.post("/admin/mailings/reports/{report_id}/edit")
async def mailing_report_edit_submit(report_id: int, request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    cid = form.get("company_id")
    update_mailing_report(
        report_id,
        name=(form.get("name") or "").strip(),
        sql_text=form.get("sql_text") or "",
        company_id=int(cid) if cid else None,
    )
    return RedirectResponse("/admin/mailings/reports", status_code=303)


@app.get("/admin/mailings/reports/{report_id}/delete")
def mailing_report_delete(report_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    delete_mailing_report(report_id)
    return RedirectResponse("/admin/mailings/reports", status_code=303)


# ── Справочник ЧСИ ───────────────────────────────────────────

@app.get("/admin/mailings/contacts", response_class=HTMLResponse)
def chsi_contacts_page(
    company_id: Optional[int] = None,
    imported: Optional[int] = None,
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db),
):
    if isinstance(current_user, RedirectResponse):
        return current_user
    companies = db.query(Company).order_by(Company.id).all()
    contacts = get_chsi_contacts(company_id) if company_id else []
    return render(
        CHSI_CONTACTS_HTML,
        companies=companies, company_id=company_id,
        contacts=contacts, imported=imported,
    )


@app.post("/admin/mailings/contacts/add")
async def chsi_contact_add(request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    cid = int(form.get("company_id"))
    create_chsi_contact(cid, (form.get("fio") or "").strip(), (form.get("email") or "").strip())
    return RedirectResponse(f"/admin/mailings/contacts?company_id={cid}", status_code=303)


@app.post("/admin/mailings/contacts/{contact_id}/edit")
async def chsi_contact_edit(contact_id: int, request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    form = await request.form()
    cid = int(form.get("company_id"))
    update_chsi_contact(
        contact_id,
        fio=(form.get("fio") or "").strip(),
        email=(form.get("email") or "").strip(),
        is_active=bool(form.get("is_active")),
    )
    return RedirectResponse(f"/admin/mailings/contacts?company_id={cid}", status_code=303)


@app.get("/admin/mailings/contacts/{contact_id}/delete")
def chsi_contact_delete(contact_id: int, company_id: int, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    delete_chsi_contact(contact_id)
    return RedirectResponse(f"/admin/mailings/contacts?company_id={company_id}", status_code=303)


@app.post("/admin/mailings/contacts/import")
async def chsi_contacts_import(request: Request, current_user: User = Depends(get_admin_user)):
    if isinstance(current_user, RedirectResponse):
        return current_user
    from openpyxl import load_workbook
    form = await request.form()
    cid = int(form.get("company_id"))
    upload = form.get("file")
    added = 0
    if upload is not None and getattr(upload, "filename", ""):
        dest = UPLOADS_DIR / f"chsi_import_{uuid.uuid4().hex[:8]}.xlsx"
        dest.write_bytes(await upload.read())
        try:
            wb = load_workbook(dest, data_only=True, read_only=True)
            ws = wb.active
            rows = []
            email_col = 1  # по умолчанию колонка B, как в старом узком формате
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if not row:
                    continue
                if i == 0:
                    # ищем колонку "Email" по заголовку — файлы бывают шире двух колонок
                    # (напр. ФИО/Область/Адрес/Телефон/Email)
                    header_hit = next(
                        (j for j, v in enumerate(row) if v and "email" in str(v).strip().lower()),
                        None,
                    )
                    if header_hit is not None:
                        email_col = header_hit
                        continue  # это точно строка-заголовок
                    email0 = row[1] if len(row) > 1 else None
                    if email0 and "@" in str(email0):
                        pass  # не заголовок, первая строка — уже данные
                    else:
                        continue  # строка-заголовок без явного "Email"
                fio = row[0] if len(row) > 0 else None
                email = row[email_col] if len(row) > email_col else None
                rows.append((str(fio or "").strip(), str(email or "").strip()))
            wb.close()
            added = bulk_add_chsi_contacts(cid, rows)
        finally:
            try:
                dest.unlink()
            except OSError:
                pass
    return RedirectResponse(f"/admin/mailings/contacts?company_id={cid}&imported={added}", status_code=303)


MONITORING_HTML = """
<!doctype html><html><head><meta charset="utf-8"><title>Мониторинг</title>
<style>
body { font-family: Arial; max-width: 1200px; margin: 20px auto; }
table { width: 100%; border-collapse: collapse; margin: 20px 0; }
th, td { padding: 10px; border: 1px solid #ddd; text-align: left; }
th { background: #17a2b8; color: white; }
.status-success { color: #28a745; font-weight: bold; }
.status-error { color: #dc3545; font-weight: bold; }
.status-running { color: #ffc107; font-weight: bold; }
.status-queued { color: #6c757d; }
.summary { display: flex; gap: 15px; margin: 20px 0; }
.card { flex: 1; background: white; padding: 15px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); text-align: center; }
.card .num { font-size: 28px; font-weight: bold; }
.filters { margin: 15px 0; }
.filters a { margin-right: 10px; padding: 6px 12px; background: #e9ecef; border-radius: 4px; text-decoration: none; color: #333; }
.filters a.active { background: #17a2b8; color: white; }
.company-badge { display: inline-block; padding: 3px 10px; background: #e7f1ff; color: #0b4f6c; border-radius: 12px; font-size: 12px; font-weight: 500; }
.script-stats { width: 100%; border-collapse: collapse; margin: 10px 0 25px; }
.script-stats th, .script-stats td { padding: 6px 10px; border: 1px solid #ddd; text-align: right; font-size: 13px; }
.script-stats th:first-child, .script-stats td:first-child { text-align: left; }
.script-stats th { background: #f1f3f5; }
.script-stats tr.active-row { background: #e7f1ff; }
.date-form { display: inline-flex; gap: 6px; align-items: center; margin-right: 10px; }
.date-form input[type=date] { padding: 4px 6px; }
.date-form button { padding: 6px 12px; background: #17a2b8; color: white; border: none; border-radius: 4px; cursor: pointer; }
select.script-select { padding: 6px 10px; border-radius: 4px; border: 1px solid #ced4da; }
</style>
{% if auto_refresh %}<meta http-equiv="refresh" content="15">{% endif %}
</head>
<body>
<a href="/admin">← Назад в админку</a> | <a href="/admin/schedule">⏰ Расписание</a>
<h2>📊 Мониторинг автоматических запусков</h2>

<div class="summary">
  <div class="card"><div class="num" style="color:#28a745;">{{ success_count }}</div>Успешно ({{ period_label }})</div>
  <div class="card"><div class="num" style="color:#dc3545;">{{ error_count }}</div>Ошибок ({{ period_label }})</div>
  <div class="card"><div class="num" style="color:#ffc107;">{{ running_count }}</div>Выполняется</div>
</div>

<div class="filters">
  <a href="{{ mk_url(period='today') }}" class="{% if period=='today' %}active{% endif %}">За день</a>
  <a href="{{ mk_url(period='week') }}" class="{% if period=='week' %}active{% endif %}">За неделю</a>
  <a href="{{ mk_url(period='month') }}" class="{% if period=='month' %}active{% endif %}">За месяц</a>
  <a href="{{ mk_url(period='all') }}" class="{% if period=='all' %}active{% endif %}">Всё время</a>
</div>
<div class="filters">
  <form class="date-form" method="get" action="/monitoring">
    <input type="hidden" name="status" value="{{ status }}">
    <input type="hidden" name="company" value="{{ current_company }}">
    <input type="hidden" name="script" value="{{ current_script }}">
    <input type="hidden" name="period" value="custom">
    с <input type="date" name="date_from" value="{{ date_from or '' }}">
    по <input type="date" name="date_to" value="{{ date_to or '' }}">
    <button type="submit">Показать</button>
  </form>
</div>
<div class="filters">
  <a href="{{ mk_url(status='all') }}" class="{% if status=='all' %}active{% endif %}">Все статусы</a>
  <a href="{{ mk_url(status='error') }}" class="{% if status=='error' %}active{% endif %}">Только ошибки</a>
</div>
<div class="filters">
  <a href="{{ mk_url(company='all') }}" class="{% if current_company=='all' %}active{% endif %}">Все компании</a>
  {% for c in companies %}
    <a href="{{ mk_url(company=c.id) }}" class="{% if current_company==c.id|string %}active{% endif %}">{{ c.name }}</a>
  {% endfor %}
</div>
<div class="filters">
  <form method="get" action="/monitoring" style="display:inline;">
    <input type="hidden" name="status" value="{{ status }}">
    <input type="hidden" name="company" value="{{ current_company }}">
    <input type="hidden" name="period" value="{{ period }}">
    <input type="hidden" name="date_from" value="{{ date_from or '' }}">
    <input type="hidden" name="date_to" value="{{ date_to or '' }}">
    <select class="script-select" name="script" onchange="this.form.submit()">
      <option value="all" {% if current_script=='all' %}selected{% endif %}>Все скрипты</option>
      {% for sk in script_keys %}
        <option value="{{ sk }}" {% if current_script==sk %}selected{% endif %}>{{ scripts[sk].title if sk in scripts else sk }}</option>
      {% endfor %}
    </select>
  </form>
</div>

<h3>По скриптам ({{ period_label }})</h3>
<table class="script-stats">
  <tr><th>Скрипт</th><th>Всего</th><th>✅ Успешно</th><th>❌ Ошибок</th><th>⏳ Выполняется</th></tr>
  {% for sk, st in by_script.items() %}
  <tr class="{% if current_script==sk %}active-row{% endif %}">
    <td><a href="{{ mk_url(script=sk) }}">{{ scripts[sk].title if sk in scripts else sk }}</a></td>
    <td>{{ st.total }}</td>
    <td style="color:#28a745;">{{ st.success }}</td>
    <td style="color:#dc3545;">{{ st.error }}</td>
    <td style="color:#ffc107;">{{ st.running }}</td>
  </tr>
  {% endfor %}
  {% if not by_script %}<tr><td colspan="5" style="text-align:center; color:#666;">Нет запусков за этот период.</td></tr>{% endif %}
</table>

<table>
  <tr><th>Время</th><th>Скрипт</th><th>Компания</th><th>Статус</th><th>Начало</th><th>Конец</th><th>Ошибка</th><th>Логи</th><th>Действия</th></tr>
  {% for j in runs %}
  <tr>
    <td>{{ j.created_at.strftime('%Y-%m-%d %H:%M:%S') if j.created_at else '—' }}</td>
    <td>{{ scripts[j.script_key].title if j.script_key in scripts else j.script_key }}</td>
    <td><span class="company-badge">{{ j.company_name or 'Все компании' }}</span></td>
    <td>
      {% if j.status == 'success' %}<span class="status-success">✅ Успешно</span>
      {% elif j.status == 'error' %}<span class="status-error">❌ Ошибка</span>
      {% elif j.status == 'running' %}<span class="status-running">⏳ Выполняется</span>
      {% elif j.status == 'cancelled' %}<span class="status-queued">🚫 Отменено</span>
      {% else %}<span class="status-queued">🕐 В очереди</span>{% endif %}
    </td>
    <td>{{ j.started_at.strftime('%H:%M:%S') if j.started_at else '—' }}</td>
    <td>{{ j.finished_at.strftime('%H:%M:%S') if j.finished_at else '—' }}</td>
    <td style="max-width:250px; font-size:12px; color:#dc3545;">{{ j.error_text or '' }}</td>
    <td><a href="/logs/{{ j.id }}">Открыть</a></td>
    <td>
      {% if j.status in ['running', 'queued', 'scheduled'] %}
        <a href="/stop/{{ j.id }}" onclick="return confirm('Остановить задачу?')" style="color:#dc3545;">⛔ Остановить</a>
      {% else %}—{% endif %}
    </td>
  </tr>
  {% endfor %}
</table>

<div class="filters" style="display:flex; align-items:center; justify-content:space-between;">
  <span style="color:#666; font-size:13px;">Всего запусков: {{ total_runs }} | страница {{ page }} из {{ total_pages }}</span>
  <span>
    {% if page > 1 %}<a href="{{ mk_url(page=page-1) }}">← Назад</a>{% endif %}
    {% if page < total_pages %}<a href="{{ mk_url(page=page+1) }}">Вперёд →</a>{% endif %}
  </span>
</div>
</body></html>
"""

@app.get("/monitoring", response_class=HTMLResponse)
def monitoring(
    status: str = "all",
    period: str = "today",
    date_from: str = None,
    date_to: str = None,
    company: str = "all",
    script: str = "all",
    page: int = 1,
    current_user: User = Depends(get_admin_user),
    db: Session = Depends(get_db),
):
    if isinstance(current_user, RedirectResponse):
        return current_user

    company_id = int(company) if company != "all" else None
    script_key = script if script != "all" else None
    page = max(1, page)
    PAGE_SIZE = 20

    today = datetime.now().date()
    df = dt_ = None
    period_label = "всё время"
    if period == "today":
        df = today
        period_label = "сегодня"
    elif period == "week":
        df = today - timedelta(days=6)
        period_label = "за 7 дней"
    elif period == "month":
        df = today - timedelta(days=29)
        period_label = "за 30 дней"
    elif period == "custom":
        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date() if date_from else None
        except ValueError:
            df = None
        try:
            dt_ = datetime.strptime(date_to, "%Y-%m-%d").date() if date_to else None
        except ValueError:
            dt_ = None
        if df and dt_:
            period_label = f"{df.strftime('%d.%m.%Y')} — {dt_.strftime('%d.%m.%Y')}"
        elif df:
            period_label = f"с {df.strftime('%d.%m.%Y')}"
        elif dt_:
            period_label = f"по {dt_.strftime('%d.%m.%Y')}"
        else:
            period_label = "всё время"
    # period == "all" — df/dt_ остаются None

    status_filter = "error" if status == "error" else None

    total_runs = count_job_runs(source="scheduler", status=status_filter, date_from=df, date_to=dt_,
                                 script_key=script_key, company_id=company_id)
    total_pages = max(1, -(-total_runs // PAGE_SIZE))  # ceil
    page = min(page, total_pages)
    runs = get_job_runs(source="scheduler", status=status_filter, date_from=df, date_to=dt_,
                         script_key=script_key, limit=PAGE_SIZE, offset=(page - 1) * PAGE_SIZE,
                         company_id=company_id)

    # для сводки/разбивки по скриптам — тот же период+компания, без фильтра статуса/скрипта
    period_runs = get_job_runs(source="scheduler", date_from=df, date_to=dt_,
                                limit=5000, company_id=company_id)
    success_count = len([j for j in period_runs if j.status == "success"])
    error_count = len([j for j in period_runs if j.status == "error"])
    running_count = len([j for j in period_runs if j.status == "running"])

    by_script = {}
    for j in period_runs:
        st = by_script.setdefault(j.script_key, {"total": 0, "success": 0, "error": 0, "running": 0})
        st["total"] += 1
        if j.status in ("success", "error", "running"):
            st[j.status] += 1
    by_script = dict(sorted(by_script.items(), key=lambda kv: -kv[1]["total"]))

    companies = db.query(Company).all()
    script_keys = get_job_run_script_keys()

    def mk_url(**overrides):
        params = {
            "status": status, "period": period, "date_from": date_from or "",
            "date_to": date_to or "", "company": company, "script": script, "page": 1,
        }
        params.update({k: str(v) for k, v in overrides.items()})
        return "/monitoring?" + "&".join(f"{k}={v}" for k, v in params.items())

    return render(
        MONITORING_HTML,
        runs=runs,
        scripts=SCRIPTS,
        status=status,
        period=period,
        period_label=period_label,
        date_from=date_from,
        date_to=date_to,
        companies=companies,
        current_company=company,
        script_keys=script_keys,
        current_script=script,
        by_script=by_script,
        mk_url=mk_url,
        page=page,
        total_pages=total_pages,
        total_runs=total_runs,
        success_count=success_count,
        error_count=error_count,
        running_count=running_count,
        auto_refresh=(running_count > 0),
    )